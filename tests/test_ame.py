"""Pruebas automáticas del Análisis Modal Espectral.

Ejecutar desde la carpeta del proyecto:   python -m pytest -q

Los casos de las hojas (14, 18 y 19) son los resultados hechos a mano que sirven de referencia: si una
modificación futura los cambia, estas pruebas fallan.
"""
import copy
import io
import re
import sys
from pathlib import Path

import numpy as np
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import ame_dimensionado as dim  # noqa: E402
import ame_ntc as ntc  # noqa: E402
import ame_pdf  # noqa: E402
import ame_sismico as ame  # noqa: E402
from ame_marco import rigidez_marco  # noqa: E402
from ame_revision import combinar, irregularidades, rho_cqc, theta_pdelta  # noqa: E402
from ame_secciones import parse_secciones  # noqa: E402

EJ = ame.EJEMPLOS_MODELO
EJ1, EJ2, EJ3, EJ4, EJ5, EJ6, EJ7, EJ8, EJ9, EJ10, EJ11, EJ12, EJ13, EJ14, EJ15, EJ16 = list(EJ.values())
# Los ejemplos 12 y 13 vienen con el acero articulado (3EI/h³, como la solución del examen); EJ12 y EJ13 son su variante empotrada.
EJ12_SOL, EJ13_SOL = EJ12, EJ13
EJ12 = copy.deepcopy(EJ12_SOL)
EJ12['niveles'][2]['artic'] = EJ12['niveles'][3]['artic'] = ''
EJ13 = copy.deepcopy(EJ13_SOL)
for _f in EJ13['columnas']:
    if _f['seccion'].startswith('I:'):
        _f['apoyo'] = 'Empotrado'


def correr(d):
    return ame.calcular(ame.preparar(d))


# ---------------------------------------------------------------- casos de las hojas
def test_hoja14_marco_de_un_crujia():
    r = correr(EJ1)
    p = r['p']
    assert p['K'][0, 0] == pytest.approx(20985.1162, rel=1e-5)
    assert p['K'][1, 1] == pytest.approx(17088.7895, rel=1e-5)
    assert p['K'][2, 2] == pytest.approx(2876.65, rel=1e-5)
    assert np.diag(p['M']) == pytest.approx([12.2324, 12.2324, 9.1743], rel=1e-4)
    assert r['T'] == pytest.approx([0.5434, 0.2679, 0.1200], abs=5e-5)
    assert r['gam'] == pytest.approx([0.5851, 0.3022, 0.1127], abs=1e-4)
    assert r['Sa_modal'][:, 0] == pytest.approx([60.475, 59.00, 59.452], abs=1e-2)
    assert r['F_final'] == pytest.approx([491.54, 632.73, 800.56], abs=0.02)
    t = r['tabla_din']
    assert t['Vu'][-1] == pytest.approx(2.117, abs=1e-3)
    assert t['Mvu'].sum() == pytest.approx(19.941, abs=2e-3)


def test_hoja18_dos_crujias():
    r = correr(EJ2)
    assert np.diag(r['p']['M']) == pytest.approx([43.0431, 42.8128, 31.8807], abs=0.011)
    assert r['p']['K'][1, 1] == pytest.approx(69096.21, rel=1e-6)
    assert r['T'] == pytest.approx([0.6800, 0.1912, 0.1230], abs=5e-5)
    assert r['gam'] == pytest.approx([0.83623, 0.14614, 0.01763], abs=1e-4)
    assert r['Sa_modal'][:, 0] == pytest.approx([87.696, 86.218, 87.11], abs=2e-3)
    assert r['F_final'] == pytest.approx([3203.4, 3884.5, 3164.8], abs=0.2)
    t = r['tabla_din']
    assert t['Fu'] == pytest.approx([3.48, 4.27, 3.52], abs=0.01)
    assert t['Vu'][-1] == pytest.approx(11.28, abs=0.005)
    assert t['Mvu'].sum() == pytest.approx(95.72, abs=0.01)


def datos_hoja19(Q='1', sismo='B'):
    K = np.array([[300, -150, 0, 0], [-150, 225, -75, 0], [0, -75, 150, -75], [0, 0, -75, 75]]) * 1000
    M = np.diag([3, 3, 1.5, 1.5]) * 1000
    fmt = lambda A: '\n'.join(' '.join(map(str, f)) for f in A)
    return dict(K=fmt(K), M=fmt(M), niveles='5\n5\n5\n5', grupo='A', sismo=sismo, k1='0.8', Q=Q, a0='80', c='500',
                Ta='0.4', Tb='1.6', k='0.605', Ts='0.53', factor_Fu='1.1', mult_V='1')


def test_hoja19_ocupacion_inmediata_con_K_y_M_manuales():
    r = ame.calcular(ame.preparar_datos(datos_hoja19()))
    assert r['T'] == pytest.approx([2.1790, 0.9939, 0.5735, 0.5020], abs=5e-5)
    assert r['Sa_modal'][:, 0] == pytest.approx([210.02, 476.19, 476.19, 476.19], abs=0.01)
    assert r['F_final'] / 1000 == pytest.approx([608.39, 763.3, 426.32, 601.01], abs=0.02)
    t = r['tabla_din']
    assert t['Vu'] == pytest.approx([661.11, 1130.1, 1969.7, 2638.9], abs=0.05)
    assert t['Mvu'].sum() == pytest.approx(31999, abs=1)


def test_estado_ocupacion_inmediata_equivale_a_Q_uno():
    base = ame.calcular(ame.preparar_datos(datos_hoja19(Q='1')))
    d = datos_hoja19(Q='2')
    d['estado'] = 'Ocupación inmediata'
    oi = ame.calcular(ame.preparar_datos(d))
    assert oi['F_final'] == pytest.approx(base['F_final'])
    assert oi['Qs'] == pytest.approx([1.0])


# ---------------------------------------------------------------- modelo por ejes
def test_ejemplo3_columna_faltante_circulares_y_cargas_por_crujia():
    p = ame.preparar_modelo(EJ3)
    E, I, Ic = 158000, 25 * 50 ** 3 / 12, np.pi * 50 ** 4 / 64
    esperado = [3 * 12 * E * I / 500 ** 3, 2 * 12 * E * I / 350 ** 3, 3 * 12 * E * Ic / 350 ** 3]
    assert [r['k'] for r in p['info']] == pytest.approx(esperado)
    assert [r['ncol'] for r in p['info']] == [3, 2, 3]
    assert [r['W'] for r in p['info']] == pytest.approx([4.7 * 9, 4.0 * 4 + 5.5 * 5, 3.0 * 4 + 3.8 * 5])


@pytest.mark.parametrize('cambio, texto', [
    (lambda d: d['niveles'][1].update(ejes='1 5'), 'el eje 5 no existe'),
    (lambda d: d['niveles'][1].update(seccion='25x50 30x30 40x40'), 'escribiste 3 secciones'),
    (lambda d: d['niveles'][1].update(artic='2'), 'no tiene columna en ese nivel'),
    (lambda d: d['niveles'][0].update(cargas='1 2 3'), 'escribiste 3 cargas'),
    (lambda d: d['niveles'][2].update(seccion='redonda'), 'no es válida'),
    (lambda d: d['niveles'][2].update(seccion=''), 'falta la sección'),
    (lambda d: d['niveles'][2].update(ejes='abc'), 'no es un eje válido'),
    (lambda d: d['niveles'][0].update(h='0'), 'mayor que cero'),
    (lambda d: d.update(crujias='4 -5'), 'mayores que cero'),
    (lambda d: d.update(E='', fc=''), "Indica E o f'c"),
    (lambda d: d.update(estado='inventado'), 'Estado límite no válido'),
    (lambda d: d.update(combinacion='XYZ'), 'Combinación modal no válida'),
    (lambda d: d.update(vigas='flexibles'), 'falta la sección'),  # sin viga en los niveles
])
def test_errores_de_captura_con_mensaje_claro(cambio, texto):
    d = copy.deepcopy(EJ3)
    cambio(d)
    with pytest.raises(ValueError, match=texto):
        ame.preparar_modelo(d)


def test_variantes_de_escritura_aceptadas():
    d = copy.deepcopy(EJ3)
    d['niveles'][2]['seccion'] = 'ø 50, d50 Ø50'
    d['niveles'][1]['seccion'] = '25 X 50'
    d['niveles'][1]['ejes'] = '1,3'
    p = ame.preparar_modelo(d)
    assert [r['seccion'] for r in p['info']] == ['25x50', '25x50', 'Ø50']


def test_proyecto_en_formato_anterior_se_convierte():
    viejo = dict(niveles=[dict(h='4.5', ncol='2', b='25', hs='50', nart='1', w='2.0', L='6', W=''),
                          dict(h='3.7', ncol='2', b='25', hs='45', nart='0', w='2.0', L='6', W=''),
                          dict(h='5.2', ncol='2', b='20', hs='40', nart='0', w='1.5', L='6', W='')],
                 E='158000', fc='', tipo='Grupo B', Q='4', k1='1.0', a0='153', c='636', Ta='0.6', Tb='1.8',
                 k='0.505', Ts='0.9')
    assert correr(ame.convertir_legacy(viejo))['T'] == pytest.approx(correr(EJ1)['T'])


# ---------------------------------------------------------------- secciones
def test_propiedades_de_secciones():
    rect = parse_secciones('25x50', 's')[0]
    assert rect['Ix'] == pytest.approx(25 * 50 ** 3 / 12) and rect['Iy'] == pytest.approx(50 * 25 ** 3 / 12)
    circ = parse_secciones('Ø60', 's')[0]
    assert circ['Ix'] == pytest.approx(np.pi * 60 ** 4 / 64) and circ['A'] == pytest.approx(np.pi * 900)
    tubo = parse_secciones('Ø60e5', 's')[0]
    assert tubo['Ix'] == pytest.approx(np.pi * (60 ** 4 - 50 ** 4) / 64)
    hueca = parse_secciones('30x40e4', 's')[0]
    assert hueca['Ix'] == pytest.approx((30 * 40 ** 3 - 22 * 32 ** 3) / 12)
    # perfil I armado: fórmula cerrada de la inercia fuerte y débil
    bf, tf, d, tw = 30, 1.6, 40, 1.0
    H = parse_secciones('H:30x1.6x40x1', 's')[0]
    fuerte = bf * d ** 3 / 12 - (bf - tw) * (d - 2 * tf) ** 3 / 12
    assert H['Ix'] == pytest.approx(fuerte)
    assert H['Iy'] < H['Ix'] / 5
    assert H['A'] == pytest.approx(2 * bf * tf + (d - 2 * tf) * tw)
    # sección T: el eje neutro está en el centroide
    T = parse_secciones('T:50x10x60x15', 's')[0]
    A = 50 * 10 + 50 * 15
    yc = (50 * 10 * 55 + 50 * 15 * 25) / A
    esperado = 50 * 10 ** 3 / 12 + 500 * (55 - yc) ** 2 + 15 * 50 ** 3 / 12 + 750 * (25 - yc) ** 2
    assert T['Ix'] == pytest.approx(esperado)
    gen = parse_secciones('I:260000/A:1250/I2:45000', 's')[0]
    assert (gen['Ix'], gen['Iy'], gen['A']) == (260000, 45000, 1250)


def test_inercia_directa_sin_area():
    a = parse_secciones('I:3688', 'x')[0]
    assert (a['Ix'], a['Iy'], a['texto']) == (3688, 3688, 'I:3688') and a['A'] == 1e9
    b = parse_secciones('I=3688/I2:98@90', 'x')[0]
    assert (b['Ix'], b['Iy']) == (98, 3688)
    d = copy.deepcopy(EJ2)
    d['niveles'] = [dict(n, seccion='I:260416.6667') for n in d['niveles']]
    assert ame.preparar_modelo(d)['K'] == pytest.approx(ame.preparar_modelo(EJ2)['K'], rel=1e-6)   # = 25x50
    d['vigas'] = 'flexibles'
    d['niveles'] = [dict(n, viga='I:520833') for n in d['niveles']]
    assert np.all(np.linalg.eigvalsh(ame.preparar_modelo(d)['K']) > 0)                          # viga solo con I


def test_direccion_del_sismo_gira_las_secciones():
    d = copy.deepcopy(EJ2)
    kx = ame.preparar_modelo(d)['info'][0]['k']
    d['direccion'] = 'Y'
    ky = ame.preparar_modelo(d)['info'][0]['k']
    assert ky / kx == pytest.approx((50 * 25 ** 3) / (25 * 50 ** 3))  # (h·b³)/(b·h³) = (b/h)²


def test_materiales():
    d = copy.deepcopy(EJ2)
    d.update(E='', fc='250')
    assert ame.preparar_modelo(d)['E'] == pytest.approx(14000 * np.sqrt(250))
    d['clase'] = 'Clase 2 (8000·√f\'c)'
    assert ame.preparar_modelo(d)['E'] == pytest.approx(8000 * np.sqrt(250))
    d.update(material='Acero', E='')
    assert ame.preparar_modelo(d)['E'] == 2_040_000
    d.update(material='Mampostería', E='')
    with pytest.raises(ValueError, match='mampostería'):
        ame.preparar_modelo(d)


# ---------------------------------------------------------------- vigas flexibles
def test_marco_con_vigas_rigidas_coincide_con_modelo_de_cortante():
    E, I = 158000, 25 * 50 ** 3 / 12
    h = [500, 350, 350]
    cols = [[dict(eje=1, A=1e9, I=I, artic=False), dict(eje=2, A=1e9, I=I, artic=False)]] * 3
    K = rigidez_marco(E, h, [0, 600], cols, [dict(I=1e13, A=1e9)] * 3)
    ks = [2 * 12 * E * I / x ** 3 for x in h]
    Kc = np.array([[ks[0] + ks[1], -ks[1], 0], [-ks[1], ks[1] + ks[2], -ks[2]], [0, -ks[2], ks[2]]])
    assert K == pytest.approx(Kc, rel=1e-5, abs=0.05)


@pytest.mark.parametrize('nivel_art', [0, 1])
def test_articulacion_en_cualquier_nivel_con_vigas_rigidas(nivel_art):
    E, I = 158000, 25 * 50 ** 3 / 12
    h = [500, 350, 350]
    cols = [[dict(eje=1, A=1e9, I=I, artic=False), dict(eje=2, A=1e9, I=I, artic=(i == nivel_art))] for i in range(3)]
    K = rigidez_marco(E, h, [0, 600], cols, [dict(I=1e13, A=1e9)] * 3)
    k = [(12 + (3 if i == nivel_art else 12)) * E * I / h[i] ** 3 for i in range(3)]
    assert K[0, 0] == pytest.approx(k[0] + k[1], rel=1e-5)


def test_vigas_flexibles_dan_marco_mas_flexible_y_matriz_simetrica():
    rig = correr(EJ4 | dict(vigas='rigidas'))
    fle = correr(EJ4)
    assert fle['T'][0] > rig['T'][0]
    K = fle['p']['K']
    assert K == pytest.approx(K.T)
    assert np.all(np.linalg.eigvalsh(K) > 0)  # estable
    assert fle['p']['vigas'] == 'flexibles'


def test_marco_de_acero_corre():
    r = correr(EJ5)
    assert r['p']['E'] == 2_040_000
    assert np.all(np.diff(r['T']) < 0) and r['T'][0] > 0.1


# ---------------------------------------------------------------- combinación modal
def test_combinaciones_modales():
    w = np.array([10.0, 30.0, 60.0])
    v = np.array([[3.0, 2.0, 1.0], [-1.0, 4.0, 2.0]])
    srss = combinar(v, w, 0.05, 'SRSS')
    assert srss == pytest.approx(np.sqrt((v ** 2).sum(axis=1)))
    assert combinar(v, w, 0.05, 'Suma absoluta') == pytest.approx(np.abs(v).sum(axis=1))
    # modos bien separados: CQC ≈ SRSS; con amortiguamiento ~0 son idénticos
    assert combinar(v, w, 0.05, 'CQC') == pytest.approx(srss, rel=0.05)
    assert combinar(v, w, 1e-6, 'CQC') == pytest.approx(srss, rel=1e-6)
    # modos casi iguales: CQC suma casi como valores absolutos
    cercanos = np.array([[1.0, 1.0]])
    assert combinar(cercanos, np.array([10.0, 10.2]), 0.05, 'CQC')[0] > 1.9
    rho = rho_cqc(w, 0.05)
    assert rho == pytest.approx(rho.T) and np.diag(rho) == pytest.approx(1.0)
    assert np.all(rho <= 1.0 + 1e-12)


def test_srss_cqc_y_suma_absoluta_ordenados_en_un_caso_real():
    d = copy.deepcopy(EJ2)
    resultados = {}
    for comb in ('SRSS', 'CQC', 'Suma absoluta'):
        d['combinacion'] = comb
        resultados[comb] = correr(d)['V_din']
    assert resultados['Suma absoluta'] >= resultados['SRSS'] - 1e-9
    assert resultados['Suma absoluta'] >= resultados['CQC'] - 1e-9


# ---------------------------------------------------------------- estados límite y revisiones
def test_limitacion_de_danos_usa_Sa_por_Ks():
    d = copy.deepcopy(EJ2)
    d.update(estado='Limitación de daños', Ks='0.25')
    r = correr(d)
    sa_el, _ = ame.espectro(r['p'], 4.0, 1.0, r['serie'], dl=True)
    assert r['Sa_modal'][:, 0] == pytest.approx(np.interp(r['T'], r['serie'], sa_el))
    assert r['a_min'] == 0.0


def un_nivel(**extra):
    d = copy.deepcopy(EJ2)
    d['niveles'] = [dict(h='3.5', ejes='todos', seccion='40x40', artic='', cargas='3', W='')]
    d.update(extra)
    return d


def test_distorsion_sv_amplifica_con_FC_Q_Rprima_del_plato_como_en_las_hojas():
    """Hojas 14 y 18: QR′ = 4(2) = 8 → R′ del plato (k2 = 0). La opción estricta usa el periodo fundamental."""
    r = correr(un_nivel())
    p, T1 = r['p'], r['T'][0]
    est = r['revision']['estados']
    assert est['SV']['amp'] == pytest.approx(1.1 * 4.0 * 2.0)                   # FC·Q·R′ con R′ = k1·R0 = 2
    estricto = correr(un_nivel(R_dist='Del periodo fundamental (estricto)'))
    R1 = float(ame._R(p, 4.0, T1))
    assert estricto['revision']['estados']['SV']['amp'] == pytest.approx(1.1 * 4.0 * R1) and R1 > 2.0
    # con un solo modo: γ_DL / γ_SV = Ks·Q′(T)·R′(T) / (FC·Q·R′ del plato)
    q, rr = (float(x[0]) for x in ame._q_r(p, 4.0, T1))
    assert est['DL']['dist'][0] / est['SV']['dist'][0] == pytest.approx(0.25 * q * rr / (1.1 * 4.0 * 2.0), rel=1e-6)


def test_ocupacion_inmediata_usa_Q_uno_y_R_prima_075R():
    r = correr(un_nivel(subgrupo='A2'))
    p = r['p']
    assert p['estado'] == 'Ocupación inmediata' and r['Qs'] == pytest.approx([1.0])
    assert r['revision']['estados']['OI']['amp'][0] == pytest.approx(1.1 * 1.0 * 0.75 * 1.75)   # R′ = 0.75·k1·R0 del plato
    estricto = correr(un_nivel(subgrupo='A2', R_dist='Del periodo fundamental (estricto)'))
    R = float(ame._R(dict(p, Q=np.array([1.0])), 1.0, r['T'][0]))
    assert estricto['revision']['estados']['OI']['amp'][0] == pytest.approx(1.1 * 1.0 * 0.75 * R)
    fila = {f[0]: f[1] for f in r['normativa']['factores']}
    assert float(fila["R′ = 0.75·R"]) == pytest.approx(0.75 * 1.75, abs=1e-4)


def test_theta_pdelta_a_mano():
    th, P = theta_pdelta([10.0, 8.0, 6.0], np.array([0.5, 0.4, 0.3]), np.array([20.0, 14.0, 7.0]),
                         np.array([400.0, 350.0, 350.0]))
    assert P == pytest.approx([24, 14, 6])
    assert th == pytest.approx([24 * 0.5 / (20 * 400), 14 * 0.4 / (14 * 350), 6 * 0.3 / (7 * 350)])


def test_irregularidades():
    masa, rigidez, k = irregularidades([10.0, 20.0, 9.0, 5.0], np.array([30.0, 20.0, 12.0, 5.0]),
                                       np.array([1.0, 1.0, 1.0, 1.0]))
    assert masa == [2]  # 20 > 1.5 * 10
    assert rigidez == []
    _, rigidez, k = irregularidades([10.0] * 4, np.array([30.0, 20.0, 12.0, 5.0]), np.array([3.0, 1.0, 1.0, 1.0]))
    assert k[:2] == pytest.approx([10.0, 20.0])
    assert rigidez == [1]  # k1 = 10 < 0.7 * k2 = 14
    _, rigidez, _ = irregularidades([10.0] * 4, np.array([30.0, 20.0, 12.0, 5.0]), np.array([2.0, 1.0, 1.0, 1.0]))
    assert rigidez == []  # k1 = 15 > 14


def test_los_ejemplos_marcan_los_avisos_esperados():
    r = correr(EJ4)
    textos = ' '.join(t for _, t in r['avisos'])
    assert 'piso blando' in textos.lower()
    # el estado de diseño (seguridad de vida) cumple su límite 0.03; la limitación de daños solo es referencia
    assert r['revision']['estados']['SV']['dist'].max() < 0.03
    assert 'Referencia (no es tu estado de diseño) · Limitación de daños' in textos
    assert not [t for s_, t in r['avisos'] if s_ == 'alerta']


def test_avisos_de_parametros_incoherentes():
    d = copy.deepcopy(EJ2)
    d.update(c='0.99', Ta='2.0', Tb='1.0', k='1.5', Q='7')
    p = ame.preparar_modelo(d)
    sev = {t[:30]: s for s, t in ame.calcular(p)['avisos']}
    textos = ' '.join(sev)
    assert 'parece estar en g' in ' '.join(t for _, t in ame.calcular(p)['avisos'])
    assert 'Ta (' in textos
    assert any('k = 1.5' in t for _, t in ame.calcular(p)['avisos'])
    assert any('Q = 7' in t for _, t in ame.calcular(p)['avisos'])


# ---------------------------------------------------------------- consistencia interna
@pytest.mark.parametrize('ej', [EJ1, EJ2, EJ3, EJ4, EJ5])
def test_propiedades_modales(ej):
    r = correr(ej)
    M, K, fi = r['p']['M'], r['p']['K'], r['fi']
    assert r['gam'].sum() > 0 and np.all(r['T'][:-1] > r['T'][1:])  # periodos decrecientes
    # ortogonalidad de los modos respecto a M y K
    Mm, Km = fi.T @ M @ fi, fi.T @ K @ fi
    fuera = lambda A: np.abs(A - np.diag(np.diag(A))).max() / np.abs(np.diag(A)).max()
    assert fuera(Mm) < 1e-8 and fuera(Km) < 1e-8
    # la masa efectiva modal acumulada suma la masa total
    m_ef = [(fi[:, j] @ M @ np.ones(r['p']['n'])) ** 2 / (fi[:, j] @ M @ fi[:, j]) for j in range(r['p']['n'])]
    assert sum(m_ef) == pytest.approx(np.trace(M))
    # el desplazamiento y las fuerzas combinados son coherentes con el cortante basal
    assert r['tabla_din']['Vu'][-1] == pytest.approx(r['p']['factor_Fu'] * r['F_final'].sum() / 1000)


def test_modo_principal_de_estado_coincide_con_la_revision():
    r = correr(EJ2)  # estado SV: d del análisis principal = d del estado SV sin amplificar
    assert r['d'] == pytest.approx(r['revision']['estados']['SV']['d'])


# ---------------------------------------------------------------- marco normativo NTC-S 2023 (apuntes)
@pytest.mark.parametrize('grupo, intensidad, clave', [
    ('B', 'Frecuente', 'DL'), ('A', 'Frecuente', 'DL'), ('B', 'Base de diseño', 'SV'), ('A', 'Base de diseño', 'OI'),
    ('A', 'Infrecuente', 'SV'), ('B', 'Infrecuente', 'SV'), ('A', 'Ocasional', 'DL'), ('B', 'Ocasional', 'OI')])
def test_objetivo_de_diseno_segun_la_matriz_de_los_apuntes(grupo, intensidad, clave):
    assert ntc.objetivo_diseno(grupo, intensidad)['clave'] == clave


def test_prevencion_de_colapso_del_grupo_B_avisa_que_requiere_acelerogramas():
    assert 'acelerogramas' in ntc.objetivo_diseno('B', 'Infrecuente')['nota']
    assert ntc.objetivo_diseno('B', 'Infrecuente')['desempeno'] == 'Prevención de colapso'


def test_tipo_del_formato_anterior_se_traduce_a_grupo_e_intensidad():
    d = copy.deepcopy(EJ2)
    d['tipo'] = 'Grupo A - sismo infrecuente'
    p = ame.preparar_modelo(d)
    assert (p['grupo'], p['intensidad'], p['estado']) == ('A', 'Infrecuente', 'Seguridad de vida')
    d['tipo'] = 'Grupo A - sismo base'
    assert ame.preparar_modelo(d)['estado'] == 'Ocupación inmediata'


def test_subgrupos_y_categoria_de_riesgo_asce():
    assert [ntc.SUBGRUPOS[s] for s in ('B', 'A2', 'A1')] == ['B', 'A', 'A']
    assert [ntc.RIESGO_ASCE[s] for s in ('B', 'A2', 'A1')] == ['I–II', 'III', 'IV']
    assert [ntc.normalizar_subgrupo(s) for s in ('B1', 'b2', 'B', 'A1')] == ['B', 'B', 'B', 'A1']
    r1, r2, r3 = (correr(dict(EJ2, subgrupo=s)) for s in ('B1', 'B2', 'B'))      # proyectos anteriores con B1 o B2
    assert r1['p']['subgrupo'] == r2['p']['subgrupo'] == r3['p']['subgrupo'] == 'B'
    assert r1['V_din'] == pytest.approx(r3['V_din']) and r2['V_din'] == pytest.approx(r3['V_din'])
    d = copy.deepcopy(EJ2)
    d['subgrupo'] = 'A2'
    p = ame.preparar_modelo(d)
    assert p['grupo'] == 'A' and p['subgrupo'] == 'A2'
    d['subgrupo'] = 'Z9'
    with pytest.raises(ValueError, match='Subgrupo no válido'):
        ame.preparar_modelo(d)


def test_espectro_por_intensidad():
    d = copy.deepcopy(EJ2)
    bd = dict(a0='224', c='975', Ta='0.8', Tb='1.7', k='0.445', Ts='1.0')
    inf = dict(a0='300', c='1300', Ta='0.8', Tb='1.7', k='0.445', Ts='1.0')
    d['espectros'] = {'Base de diseño': bd, 'Infrecuente': inf}
    v_bd = correr(dict(d, intensidad='Base de diseño'))['V_din']
    v_inf = correr(dict(d, intensidad='Infrecuente'))['V_din']
    assert v_inf / v_bd == pytest.approx(1300 / 975, rel=0.05)  # sube el plato del espectro
    # frecuente y ocasional salen del espectro de base de diseño (× Ks): no piden uno propio
    assert correr(dict(d, intensidad='Frecuente'))['p']['c'] == 975
    sin_bd = dict(d, espectros={'Infrecuente': inf})
    with pytest.raises(ValueError, match='intensidad "Base de diseño" \\(de él salen'):
        ame.preparar_modelo(dict(sin_bd, intensidad='Ocasional'))
    with pytest.raises(ValueError, match='intensidad "Infrecuente"'):
        ame.preparar_modelo(dict(d, espectros={'Base de diseño': bd}, intensidad='Infrecuente'))


def test_limitacion_de_danos_sin_factor_de_carga():
    r = correr(dict(copy.deepcopy(EJ2), intensidad='Frecuente'))
    t = r['tabla_din']
    assert r['p']['estado'] == 'Limitación de daños'
    assert t['Fu'] == pytest.approx(t['F'])  # "sin FC" (apuntes p. 12)
    assert r['revision']['estados']['DL']['gamma_c'] == 1.0


@pytest.mark.parametrize('irreg, fuerte_t, fuerte_e, esperado', [
    ([], False, False, 1.0), (['5.3.2'], False, False, 0.8), (['5.2.1', '5.2.3'], False, False, 0.7),
    (['5.2.1', '5.2.3', '5.3.1'], False, False, 0.6), ([], True, False, 0.6), (['5.3.1'], True, False, 0.5),
    (['5.2.1'], True, False, 0.6),  # 5.2.1 no cuenta como condición adicional de la fuerte irregularidad por torsión
    ([], False, True, 0.33)])
def test_gamma_c_segun_las_tablas_c542_c553_c563(irreg, fuerte_t, fuerte_e, esperado):
    assert ntc.gamma_c(irreg, fuerte_t, fuerte_e)[0] == pytest.approx(esperado)


def test_irregularidad_reduce_el_limite_de_ocupacion_inmediata_y_vida_pero_no_el_de_danos():
    r = correr(dict(copy.deepcopy(EJ2), irreg=['5.3.2', '5.3.1']))
    est = r['revision']['estados']
    assert est['SV']['limite'] == pytest.approx(0.03 * 0.7)
    assert est['OI']['limite'] == pytest.approx(0.005 * 0.7)
    assert est['DL']['limite'] == pytest.approx(0.004)
    with pytest.raises(ValueError, match='no reconocida'):
        ntc.gamma_c(['9.9.9'])


def test_excentricidad_accidental_de_los_apuntes():
    assert ntc.excentricidad_accidental(3, 12) == pytest.approx([0.6, 0.9, 1.2])   # 0.05b, 0.075b, 0.10b
    assert ntc.excentricidad_accidental(1, 10) == pytest.approx([0.5])
    r = correr(dict(copy.deepcopy(EJ2), b_planta='12'))
    tor = r['normativa']['torsion']
    assert tor['Mt'] == pytest.approx(tor['e_a'] * tor['Fu'])
    assert 'torsion' not in correr(EJ2)['normativa']


def test_aplicabilidad_del_metodo_estatico():
    assert ntc.altura_maxima_estatico('I', True) == 40 and ntc.altura_maxima_estatico('I', False) == 30
    assert ntc.altura_maxima_estatico('II', True) == 30 and ntc.altura_maxima_estatico('III', False) == 20
    assert ntc.altura_maxima_estatico('', True) is None
    assert correr(dict(copy.deepcopy(EJ2), zona='III'))['normativa']['estatico']['aplica']       # H = 12 m ≤ 30 m
    assert not correr(dict(copy.deepcopy(EJ2), zona='III', subgrupo='A1'))['normativa']['estatico']['aplica']
    alto = copy.deepcopy(EJ2)
    alto['niveles'] = [dict(alto['niveles'][1]) for _ in range(9)]          # 31.5 m en zona III regular
    assert not correr(dict(alto, zona='III'))['normativa']['estatico']['aplica']
    with pytest.raises(ValueError, match='zona geotécnica'):
        ame.preparar_modelo(dict(copy.deepcopy(EJ2), zona='IV'))


def test_cimentacion_con_065_R_prima():
    r = correr(EJ2)
    c = r['normativa']['cimentacion']
    assert c['factor'] == pytest.approx(0.65 * 2.0)                              # R′ del plato (hojas)
    assert c['V'] == pytest.approx(c['factor'] * r['V_din'])
    estricto = correr(dict(EJ2, R_dist='Del periodo fundamental (estricto)'))
    assert estricto['normativa']['cimentacion']['factor'] == pytest.approx(
        0.65 * float(ame._R(estricto['p'], 4.0, estricto['T'][0])))
    assert 'cimentacion' not in correr(dict(copy.deepcopy(EJ2), intensidad='Frecuente'))['normativa']


def test_cortante_basal_minimo():
    r = correr(EJ2)
    v = r['normativa']['vmin']
    assert v['V_min'] == pytest.approx(1.1 * r['a_min_g'] * sum(i['W'] for i in r['p']['info']))
    assert v['cumple'] == (r['V_din'] >= v['V_min'])
    assert v['factor'] >= 1.0
    # un edificio muy rígido y pesado con espectro bajo debe poder incumplirlo y avisarlo
    d = copy.deepcopy(EJ2)
    d.update(a0='1', c='5.1', Ts='2')
    r2 = correr(d)
    assert (not r2['normativa']['vmin']['cumple']) == any('Cortante basal mínimo' in t for _, t in r2['avisos'])


def test_R_igual_a_uno_en_materiales_distintos_del_concreto():
    r = correr(EJ5)
    assert r['p']['R_unitaria'] and all(f[4] == pytest.approx(1.0 * 1.0) for f in r['QR'])
    d = copy.deepcopy(EJ5)
    d['R1_otros'] = False
    r2 = correr(d)
    assert not r2['p']['R_unitaria'] and r2['QR'][0][4] > 1.7
    assert r2['V_din'] < r['V_din']                  # con R = 1 las fuerzas de diseño son mayores
    assert not correr(EJ2)['p']['R_unitaria']        # concreto: R = k1·R0 + k2


def test_amortiguamiento_por_material():
    assert correr(EJ2)['p']['zeta'] == 0.05
    d = copy.deepcopy(EJ5)
    d.pop('zeta')
    assert correr(d)['p']['zeta'] == 0.03
    d['zeta'] = '0.02'
    assert correr(d)['p']['zeta'] == 0.02


def test_participacion_modal_acumulada():
    r = correr(EJ2)
    m = r['normativa']['modos']
    assert m['acum'][-1] == pytest.approx(1.0) and m['n95'] <= 3 and m['n_T04'] == 1


def test_avisos_normativos():
    textos = lambda r: ' | '.join(t for _, t in r['avisos'])
    # k1 distinto del sugerido para el número de crujías (hoja 14: 1 crujía con k1 = 1.0)
    assert 'k1 = 0.8 con menos de 3 crujías' in textos(correr(EJ1))
    d3 = copy.deepcopy(EJ2)
    d3['crujias'] = '4 5 4'
    d3['niveles'] = [dict(n, cargas='4') for n in d3['niveles']]
    assert 'k1 = 0.8' not in textos(correr(d3))
    # grupo B, infrecuente: prevención de colapso
    assert 'acelerogramas' in textos(correr(dict(copy.deepcopy(EJ2), intensidad='Infrecuente')))
    # grupo A: el estático no aplica
    assert 'método estático no aplica' in textos(correr(dict(copy.deepcopy(EJ2), subgrupo='A1')))
    # R = 1 y ζ del acero
    assert 'R = 1' in textos(correr(EJ5))
    # intensidad distinta de la base: recordar el espectro
    assert 'espectro SASID de esa intensidad' in textos(correr(dict(copy.deepcopy(EJ2), intensidad='Infrecuente')))
    # solo el estado de diseño marca alerta; los demás son referencia
    r = correr(EJ4)
    alertas = [t for s_, t in r['avisos'] if s_ == 'alerta']
    assert all('Seguridad de vida' in t for t in alertas)
    assert any('Referencia' in t for s_, t in r['avisos'] if s_ == 'info')


def test_deteccion_de_fuerte_irregularidad_en_elevacion():
    d = copy.deepcopy(EJ2)
    d['niveles'][1] = dict(d['niveles'][1], seccion='15x15')     # el nivel 2 queda muchísimo más flexible
    r = correr(d)
    textos = ' | '.join(t for _, t in r['avisos'])
    assert r['revision']['irreg_rigidez'] == [2]
    assert 'piso blando' in textos.lower() and '(5.3.3)' in textos
    # al marcarla, el aviso desaparece y se aplica γc = 0.33 a los límites de OI y SV
    r2 = correr(dict(d, fuerte_elev=True))
    assert 'posible irregularidad fuerte' in textos
    assert 'posible irregularidad fuerte' not in ' | '.join(t for _, t in r2['avisos'])
    assert r2['p']['gamma_c'] == pytest.approx(0.33)
    assert r2['revision']['estados']['SV']['limite'] == pytest.approx(0.03 * 0.33)


# ---------------------------------------------------------------- modo "K y M dadas" (hoja 19)
def test_modo_matrices_reproduce_la_hoja_19():
    r = correr(EJ6)
    p = r['p']
    assert p['manual'] and p['estado'] == 'Ocupación inmediata' and r['Qs'] == pytest.approx([1.0])
    assert r['T'] == pytest.approx([2.1790, 0.9939, 0.5735, 0.5020], abs=5e-5)
    assert r['Sa_modal'][:, 0] == pytest.approx([210.02, 476.19, 476.19, 476.19], abs=0.01)
    assert r['F_final'] / 1000 == pytest.approx([608.39, 763.3, 426.32, 601.01], abs=0.02)
    assert r['tabla_din']['Vu'] == pytest.approx([661.11, 1130.1, 1969.7, 2638.9], abs=0.05)
    assert r['tabla_din']['Mvu'].sum() == pytest.approx(31999, abs=1)


def test_modo_matrices_pesos_y_unidades():
    r = correr(EJ6)
    assert r['p']['cargas'] == pytest.approx(np.array([3, 3, 1.5, 1.5]) * 981)     # W = m·g (t)
    kg = dict(EJ6, unidades='kg, cm',
              K='\n'.join(' '.join(str(x * 1000) for x in f) for f in
                          [[300, -150, 0, 0], [-150, 225, -75, 0], [0, -75, 150, -75], [0, 0, -75, 75]]),
              M='3000 3000 1500 1500')  # M como vector de masas
    assert correr(kg)['F_final'] == pytest.approx(r['F_final'])
    r2 = correr(dict(EJ6, pesos='100 100 100 100'))
    assert r2['p']['cargas'] == pytest.approx([100] * 4) and r2['V_est'] != pytest.approx(r['V_est'])


@pytest.mark.parametrize('cambio, texto', [
    (dict(alturas='5 5 5'), 'Hay 3 alturas'), (dict(alturas='5 5 5 -1'), 'mayores que cero'),
    (dict(pesos='1 2'), 'Hay 2 pesos'), (dict(K='1 2\n3 4'), 'M debe ser'), (dict(K='1 2 3\n4 5 6'), 'cuadrada')])
def test_modo_matrices_errores_de_captura(cambio, texto):
    with pytest.raises(ValueError, match=texto):
        ame.preparar(dict(EJ6, **cambio))


def test_modo_matrices_avisa_si_K_es_incoherente():
    sim = correr(dict(EJ6, K='300 -150 0 0\n-100 225 -75 0\n0 -75 150 -75\n0 0 -75 75'))
    assert any('K no es simétrica' in t for s_, t in sim['avisos'] if s_ == 'alerta')
    neg = correr(dict(EJ6, K='-300 150 0 0\n150 -225 75 0\n0 75 -150 75\n0 0 75 -75'))
    assert any('definida positiva' in t for s_, t in neg['avisos'])


def test_modo_matrices_detecta_K_y_M_al_reves():
    al_reves = correr(dict(EJ6, K=EJ6['M'], M=EJ6['K']))
    assert any('al revés' in t for s_, t in al_reves['avisos'] if s_ == 'alerta')
    assert al_reves['T'][0] > 60                      # los periodos salen absurdos (hoja 19: 2.18 s)
    assert not any('al revés' in t for _, t in correr(EJ6)['avisos'])


def test_modo_matrices_con_formato_anterior_sigue_funcionando():
    r = ame.calcular(ame.preparar_datos(datos_hoja19()))
    assert r['T'][0] == pytest.approx(2.1790, abs=5e-5)


def test_modo_matrices_exporta_excel_y_pdf():
    r = correr(EJ6)
    buf = io.BytesIO()
    ame.exportar_excel(r, buf)
    pdf = io.BytesIO()
    ame_pdf.memoria_pdf(r, pdf, proyecto='K y M', autor='Pruebas')
    assert buf.getvalue()[:2] == b'PK' and pdf.getvalue().startswith(b'%PDF')


# ---------------------------------------------------------------- método estático (hojas 20 y 21 de los apuntes)
def test_estatico_hoja21_marco_de_tres_niveles():
    """Grupo B, Q = 4, k1 = 1, hoja 14: C = máx(0.06208 plato; 0.0624 en T = 0) = 0.0624 g."""
    r = correr(EJ1)
    se, e = r['estatico'], r['tabla_est']
    assert se['q_pl'] == pytest.approx(5.2216, abs=1e-4) and se['Rp_pl'] == pytest.approx(2.0)   # Q′ = 5.2216, R′ = 2
    assert se['q_pl'] * se['Rp_pl'] == pytest.approx(10.4432, abs=1e-3)
    assert se['C_pl'] == pytest.approx(0.06208, abs=2e-5) and se['C_0'] == pytest.approx(0.0624, abs=2e-4)
    assert se['C'] == se['C_0'] and 'T = 0' in se['gobierna']                                   # "0.0624 → aplica"
    assert e['Fi'] == pytest.approx([0.91, 0.74, 0.41], abs=0.006)
    assert e['Fu'] == pytest.approx([1.0, 0.82, 0.45], abs=0.006)
    assert e['Vu'] == pytest.approx([1.0, 1.82, 2.26], abs=0.006)
    assert e['Mvu'].sum() == pytest.approx(22.11, abs=0.02)
    assert e['Fi'].sum() == pytest.approx(se['C'] * se['W_T'])                                  # ΣFi = V_B = C·ΣW


def test_estatico_hoja21_estructura_irregular_de_cinco_niveles():
    """Grupo B, Q = 2, k1 = 0.8, irregular: c/(Q′R′) = 0.2841 y a0/(Q′R′) = 0.1202."""
    r = correr(EJ7)
    se, e = r['estatico'], r['tabla_est']
    assert se['q_pl'] == pytest.approx(2.4991, abs=1e-4) and se['Rp_pl'] == pytest.approx(1.40)
    assert se['Rp_0'] == pytest.approx(1.90)                                                    # k2 = 0.5 en T = 0
    assert se['C_pl'] == pytest.approx(0.2841, abs=1e-4) and se['C_0'] == pytest.approx(0.1202, abs=1e-4)
    assert 'plato' in se['gobierna']
    assert e['Wi'] == pytest.approx([300, 400, 400, 400, 400]) and e['hWi'].sum() == pytest.approx(16500)
    assert e['Fi'] == pytest.approx([147.20, 157.02, 117.76, 78.51, 39.25], abs=0.01)
    assert e['Fu'] == pytest.approx([161.92, 172.72, 129.54, 86.36, 43.18], abs=0.01)
    assert e['Vu'] == pytest.approx([161.92, 334.64, 464.18, 550.53, 593.71], abs=0.01)
    assert e['Mvu'] == pytest.approx([2428.83, 2072.60, 1165.84, 518.15, 129.54], abs=0.01)
    assert e['Mvu'].sum() == pytest.approx(6314.95, abs=0.01)
    assert 2 * se['Rp_pl'] == pytest.approx(2.8)                                                # Q·R′ = 2(1.4)
    assert r['normativa']['estatico']['aplica']                                                 # H = 15 m ≤ 20 m (zona III, irregular)


def test_estatico_criterio_maximo_del_espectro_es_el_del_codigo_de_matlab():
    r = correr(dict(EJ1, criterio_estatico='Máximo del espectro de diseño (como el código de MATLAB)'))
    se = r['estatico']
    assert se['criterio'] == 'maximo' and se['C'] == pytest.approx(float(r['curvas_g'][0].max()))
    assert se['C'] > correr(EJ1)['estatico']['C']       # 0.063008 g contra 0.0624 g de los apuntes


def test_estatico_distorsiones_a_mano():
    """Un solo nivel: u = F/k y la distorsión = u·FC·Q·R′/h."""
    d = dict(EJ6, K='1000', M='0.4', alturas='3', pesos='400', subgrupo='B', intensidad='Base de diseño', Q='4', k1='1')
    r = correr(d)
    se = r['estatico']
    F_kg = se['C'] * 400 * 1000
    assert se['u'][0] == pytest.approx(F_kg / 1_000_000)                    # K = 1000 t/cm = 1e6 kg/cm
    assert se['dist'][0] == pytest.approx(se['u'][0] * 1.1 * 4 * se['Rp_pl'] / 300)


def test_estatico_limitacion_de_danos_sin_FC_ni_QR():
    r = correr(dict(EJ1, intensidad='Frecuente', espectros={'Frecuente': dict(a0='153', c='636', Ta='0.6', Tb='1.8',
                                                                                  k='0.505', Ts='0.9')}))
    se = r['estatico']
    assert r['clave'] == 'DL' and se['FC'] == 1.0 and se['amp'] == 1.0
    assert se['C'] == pytest.approx(636 / 981 / 4.4)                      # Sa(BD)·Ks con Ks = 1/[6-4(0.9-0.5)] (Ts = 0.9)
    assert r['tabla_est']['Fu'] == pytest.approx(r['tabla_est']['Fi'])


def test_estatico_ocupacion_inmediata_Q_uno_R_prima_075R():
    r = correr(dict(EJ7, subgrupo='A2'))
    se = r['estatico']
    assert r['clave'] == 'OI' and se['q_pl'] == pytest.approx(1.0)
    assert se['Rp_pl'] == pytest.approx(0.75 * 0.8 * 1.75) and se['Rp_0'] == pytest.approx(0.75 * (0.8 * 1.75 + 0.5))
    assert not r['normativa']['estatico']['aplica']                         # grupo A: el estático no aplica


def test_estatico_acero_R_igual_a_uno():
    r = correr(EJ5)
    assert r['estatico']['R_pl'] == 1.0 and r['estatico']['R_0'] == 1.0


def test_tablas_de_derivacion_y_comparacion_estatica():
    r = correr(EJ7)
    enc, filas = ame.tabla_derivacion_estatica(r)
    assert enc[0] == 'Concepto' and filas[-1][1] == pytest.approx(r['estatico']['C_pl'])
    enc, filas = ame.tabla_distorsiones_estatico(r)
    assert [f[0] for f in filas] == [5, 4, 3, 2, 1] and len(filas[0]) == len(enc)
    assert r['estatico']['dist_din'] == pytest.approx(r['revision']['estados'][r['clave']]['dist'])


def test_factor_de_escala_y_fuerzas_dinamicas_escaladas():
    r = correr(EJ7)
    f = r['factor_escala']
    assert f == pytest.approx(r['V_est'] / r['V_din']) and f > 1
    te = r['tabla_din_esc']
    assert r['factor_aplicado'] == pytest.approx(f)
    assert te['Vu'][-1] == pytest.approx(r['V_est'])                        # el dinámico escalado iguala al estático
    assert te['Fu'] == pytest.approx(r['tabla_din']['Fu'] * f)
    bajo = correr(dict(EJ7, mult_V='0.1'))                                  # factor < 1: no se escala
    assert bajo['factor_escala'] < 1 and bajo['factor_aplicado'] == 1.0
    assert bajo['tabla_din_esc']['Fu'] == pytest.approx(bajo['tabla_din']['Fu'])


def test_estatico_se_exporta_y_se_grafica():
    from openpyxl import load_workbook
    r = correr(EJ7)
    buf = io.BytesIO()
    ame.exportar_excel(r, buf)
    ws = load_workbook(io.BytesIO(buf.getvalue()))['Estático']
    texto = ' '.join(str(c.value) for fila in ws.iter_rows() for c in fila if c.value is not None)
    assert 'Coeficiente sísmico C' in texto and 'Fuerzas dinámicas escaladas' in texto and 'PASO 7' in texto
    fig = ame.fig_diagramas(r, 'estatico')
    buf2 = io.BytesIO()
    fig.savefig(buf2, format='png')
    assert len(buf2.getvalue()) > 5_000
    assert 'ANÁLISIS ESTÁTICO' in ame.reporte_texto(r) and 'Σ F_i' in ame.reporte_texto(r)


def test_aviso_de_distorsion_estatica_y_criterio_matlab():
    textos = lambda r: ' | '.join(t for _, t in r['avisos'])
    assert 'máximo del espectro de diseño (como el código de MATLAB)' in textos(
        correr(dict(EJ1, criterio_estatico='Máximo del espectro (MATLAB)')))
    flexible = correr(dict(EJ7, K='300 -150 0 0 0\n-150 300 -150 0 0\n0 -150 300 -150 0\n0 0 -150 300 -150\n0 0 0 -150 150'))
    assert 'Método estático: la distorsión supera' in textos(flexible)


# ---------------------------------------------------------------- análisis elástico (sin reducir) y comparación
def test_elastico_igual_al_de_diseno_cuando_no_hay_reduccion():
    """Acero con R = 1 y Q = 1: Q′ = 1 y R′ = 1, así que el espectro de diseño ES el elástico."""
    r = correr(dict(EJ5, Q='1'))
    assert r['Sa_modal'][:, 0] == pytest.approx(r['elastico']['Sa'])
    assert r['elastico']['V_el'] == pytest.approx(r['V_din'])
    assert r['elastico']['tabla']['Fu'] == pytest.approx(r['tabla_din']['Fu'])


def test_elastico_un_modo_la_razon_es_Qprima_por_Rprima():
    r = correr(un_nivel())
    q, rr = (float(x[0]) for x in ame._q_r(r['p'], 4.0, r['T'][0]))
    assert r['elastico']['V_el'] / r['V_din'] == pytest.approx(q * rr, rel=1e-6)
    assert r['elastico']['Sa'][0] / r['Sa_modal'][0, 0] == pytest.approx(q * rr, rel=1e-6)


def test_elastico_espectro_sin_reducir_en_el_plato():
    r = correr(EJ2)
    el = r['elastico']
    i = int(np.argmin(np.abs(r['serie'] - 1.2)))                 # T = 1.2 s está en el plato (0.8 a 1.7 s)
    assert el['curva'][i] == pytest.approx(975.0)                # c, sin dividir entre Q′R′
    assert np.all(el['Sa'] >= r['Sa_modal'][:, 0])               # elástico ≥ diseño
    assert el['V_el'] > r['V_din'] and el['tabla']['Mvu'].sum() > r['tabla_din']['Mvu'].sum()


def test_elastico_estatico_usa_el_maximo_de_c_y_a0():
    r = correr(EJ2)
    est = r['elastico']['est']
    assert est['C'] == pytest.approx(975 / 981) and est['tabla']['Fi'].sum() == pytest.approx(est['C'] * est['tabla']['Wi'].sum())


def test_comparacion_de_metodos():
    r = correr(EJ7)
    enc, filas = ame.tabla_comparacion_metodos(r)
    nombres = [f[0] for f in filas]
    assert enc[0] == 'Análisis' and any('Estático' in x for x in nombres) and any('Elástico dinámico' in x for x in nombres)
    assert any('escalado' in x for x in nombres) and any('Elástico estático' in x for x in nombres)
    enc, filas = ame.tabla_niveles_metodos(r)
    assert len(filas) == 5 and len(filas[0]) == len(enc) and 'Fu elástico' in enc
    enc, filas = ame.tabla_elastico_comparacion(r)
    assert all(f[4] > 1 for f in filas)                          # la fuerza elástica supera a la de diseño en todos los niveles
    for fig in (ame.fig_elastico(r), ame.fig_comparacion(r), ame.fig_diagramas(r, 'elastico'), ame.fig_diagramas(r, 'escalado')):
        buf = io.BytesIO()
        fig.savefig(buf, format='png')
        assert len(buf.getvalue()) > 5_000
    import ame_detalle
    enc, filas = ame_detalle.resumen_metodos(r)
    assert [f[0] for f in filas] == ['Modal espectral', 'Estático', 'Modal espectral escalado']
    assert 'ELÁSTICO' not in ame.reporte_texto(r).upper()


# ---------------------------------------------------------------- unidades y opciones contra "trampas"
def test_espectro_en_g_equivale_a_cm_s2():
    base = correr(EJ2)
    g = dict(copy.deepcopy(EJ2), a0=str(224 / 981), c=str(975 / 981), unidad_espectro='g')
    assert correr(g)['V_din'] == pytest.approx(base['V_din'], rel=1e-9)
    assert correr(g)['p']['c'] == pytest.approx(975.0)
    with pytest.raises(ValueError, match='cm/s² o g'):
        ame.preparar_modelo(dict(EJ2, unidad_espectro='ft/s'))


def test_unidades_de_E_y_fc():
    d = dict(copy.deepcopy(EJ2), E='2213594.4', unidad_E='t/m²')                    # hoja 7: E = 2,213,594.4 t/m²
    assert ame.preparar_modelo(d)['E'] == pytest.approx(221359.44, rel=1e-6)
    d = dict(copy.deepcopy(EJ2), E='', fc='24.516', unidad_fc='MPa')                # 24.516 MPa ≈ 250 kg/cm²
    assert ame.preparar_modelo(d)['E'] == pytest.approx(14000 * np.sqrt(250), rel=1e-3)
    assert ame.preparar_modelo(dict(copy.deepcopy(EJ2), E='2.04', unidad_E='GPa'))['E'] == pytest.approx(2.04 * 10197.16)
    with pytest.raises(ValueError, match='Unidad no válida'):
        ame.preparar_modelo(dict(EJ2, unidad_E='psi'))


def test_secciones_en_metros():
    a = ame.preparar_modelo(EJ2)
    d = copy.deepcopy(EJ2)
    d['niveles'] = [dict(n, seccion='0.25x0.5') for n in d['niveles']]
    b = ame.preparar_modelo(dict(d, unidad_seccion='m'))
    assert b['K'] == pytest.approx(a['K'])
    d['niveles'][0]['seccion'] = 'I:0.0026/A:0.125'
    with pytest.raises(ValueError, match='cm⁴ y cm²'):
        ame.preparar_modelo(dict(d, unidad_seccion='m'))
    with pytest.raises(ValueError, match='cm o m'):
        ame.preparar_modelo(dict(EJ2, unidad_seccion='pulg'))


def test_modos_que_se_combinan():
    todos = correr(EJ4)
    assert todos['n_modos'] == 3
    dos = correr(dict(EJ4, modos='Los primeros N', modos_n='2'))
    assert dos['n_modos'] == 2
    esperado = np.sqrt((dos['F'][:, :2] ** 2).sum(axis=1))                    # SRSS solo con 2 modos
    assert dos['F_final'] == pytest.approx(esperado)
    assert dos['F_final'][0] < todos['F_final'][0]
    uno = correr(dict(EJ4, modos='Los primeros N', modos_n='1'))
    assert uno['n_modos'] == 1 and uno['F_final'] == pytest.approx(np.abs(uno['F'][:, 0]))
    assert correr(dict(EJ4, modos='Criterio de los apuntes (95 % de masa o T ≥ 0.4 s)'))['n_modos'] >= 3
    assert correr(dict(EJ4, modos='Los primeros N', modos_n='9'))['n_modos'] == 3     # no pasa del total
    # las revisiones y el elástico usan los mismos modos
    assert dos['elastico']['F_final'] == pytest.approx(np.sqrt((dos['elastico']['F'][:, :2] ** 2).sum(axis=1)))
    assert 'Se combinan 2 de 3 modos' in ' '.join(t for _, t in dos['avisos'])


def test_ordenadas_espectrales_dadas_por_modo():
    r = correr(dict(EJ1, Sa_modos='60.475 59.00 59.452'))
    assert r['Sa_modal'][:, 0] == pytest.approx([60.475, 59.00, 59.452])
    assert r['F_final'] == pytest.approx(correr(EJ1)['F_final'], rel=2e-3)             # las de la hoja 14
    g = correr(dict(EJ1, Sa_modos='0.06165 0.06014 0.06060', unidad_espectro='g', a0='0.156', c='0.6483'))
    assert g['Sa_modal'][:, 0] == pytest.approx(np.array([0.06165, 0.06014, 0.06060]) * 981)
    with pytest.raises(ValueError, match='ordenadas Sa por modo'):
        correr(dict(EJ1, Sa_modos='60 59'))


def test_excentricidad_estatica_en_la_torsion():
    r = correr(dict(EJ2, b_planta='12', e_s='0.5'))
    t = r['normativa']['torsion']
    assert t['e1_din'] == pytest.approx(0.5 + t['e_a']) and t['e2_din'] == pytest.approx(0.5 - t['e_a'])
    assert t['e1_est'] == pytest.approx(1.5 * 0.5 + t['e_a'])                          # estático: 1.5·e_s ± e_a
    assert t['Mt1_din'] == pytest.approx(t['e1_din'] * t['Fu'])
    assert 'e_s' not in correr(dict(EJ2, b_planta='12'))['normativa']['torsion']


@pytest.mark.parametrize('cambio, texto', [
    (lambda d: d['niveles'][0].update(h='350'), 'es absurda'),
    (lambda d: d.update(unidad_seccion='cm', niveles=[dict(n, seccion='0.25x0.5') for n in d['niveles']]), 'mide menos de 3 cm'),
    (lambda d: d.update(Q='0.5'), 'menor que 1'),
    (lambda d: d.update(lim_sv='3'), '¿lo capturaste en %?'),
    (lambda d: d['niveles'][0].update(cargas='4700'), '¿Está en kg/m?'),
    (lambda d: d.update(k1='5'), 'fuera de los valores usuales'),
])
def test_avisos_de_capturas_absurdas(cambio, texto):
    d = copy.deepcopy(EJ2)
    cambio(d)
    assert texto in ' | '.join(t for _, t in correr(d)['avisos'])


def test_aviso_de_periodo_absurdo_si_K_y_M_estan_al_reves_o_en_otras_unidades():
    r = correr(dict(EJ6, K=EJ6['M'], M=EJ6['K']))
    assert any('es absurdo para un edificio' in t for s_, t in r['avisos'] if s_ == 'alerta')


# ---------------------------------------------------------------- lo que pide el examen del 29-oct-2025
def test_matriz_de_objetivos_del_examen_ocasional_y_MC():
    """Tabla del examen: Frec B→OC · Ocas A→OC, B→OI · BD A→OI, B→SV · MC A→SV, B→PC."""
    esperado = {('B', 'Frecuente'): 'DL', ('A', 'Ocasional'): 'DL', ('B', 'Ocasional'): 'OI', ('A', 'Base de diseño'): 'OI',
                ('B', 'Base de diseño'): 'SV', ('A', 'Infrecuente'): 'SV'}
    for (g, i), clave in esperado.items():
        assert ntc.objetivo_diseno(g, i)['clave'] == clave
    assert ntc.objetivo_diseno('B', 'Infrecuente')['desempeno'] == 'Prevención de colapso'
    for alias in ('MC', 'mc', 'Infrecuente (MC)', 'Máximo considerado', 'Infrecuente'):
        assert ntc.normalizar_intensidad(alias) == 'Infrecuente'
    assert ntc.normalizar_intensidad('Base de diseño (BD)') == 'Base de diseño'
    assert correr(dict(EJ8, intensidad='MC'))['p']['intensidad'] == 'Infrecuente'


def test_sismo_ocasional_multiplica_Ks_por_2_5():
    base = correr(dict(EJ2, intensidad='Frecuente'))
    ocas = correr(dict(EJ2, subgrupo='A1', intensidad='Ocasional'))
    assert base['p']['Ks'] == 0.25 and ocas['p']['Ks'] == pytest.approx(2.5 * 0.25) and ocas['clave'] == 'DL'
    assert ocas['V_din'] == pytest.approx(2.5 * base['V_din'])                 # Sa(BD)·2.5·Ks contra Sa(BD)·Ks
    assert correr(dict(EJ2, subgrupo='B', intensidad='Ocasional'))['clave'] == 'OI'   # grupo B ocasional: ocupación inmediata
    assert '2.5' in ' '.join(f[1] for f in ocas['normativa']['factores'] if f[0] == 'Ks')


def test_E_en_toneladas_por_cm2():
    d = dict(EJ8, E='158', unidad_E='t/cm²')
    assert ame.preparar_modelo(d)['E'] == pytest.approx(158_000)
    assert ame.preparar_modelo(dict(EJ8, E='158000', unidad_E='kg/cm²'))['E'] == pytest.approx(158_000)


def test_apoyo_de_rodillos_no_aporta_rigidez_lateral_ejemplo_del_examen():
    """Eje 4 con rodillos en la base: su columna no resiste cortante en el nivel 1 (k = 0)."""
    p = ame.preparar_modelo(EJ8)
    E, h = 158_000.0, 300.0
    I = {'20x40': 20 * 40 ** 3 / 12, '30x30': 30 * 30 ** 3 / 12, '30x60': 30 * 60 ** 3 / 12}
    k12 = lambda i: 12 * E * i / h ** 3
    k3 = lambda i: 3 * E * i / h ** 3
    assert p['info'][0]['k'] == pytest.approx(k12(I['20x40']) + k3(I['30x30']) + k12(I['20x40']))   # eje 4 → 0
    assert p['info'][1]['k'] == pytest.approx(2 * k12(I['20x40']) + k12(I['30x30']) + k12(I['30x60']))
    assert p['cargas'] == pytest.approx([20.0, 10.0])                        # 2 T/m y 1 T/m sobre 10 m
    cols = {c['eje']: c for c in p['geom']['niveles'][0]['cols']}
    assert cols[4]['rodillo'] and cols[4]['k'] == 0 and cols[2]['artic'] and not cols[1]['artic']
    assert 'Rodillo' in ame.apoyo_txt(cols[4]) and ame.apoyo_txt(cols[2]) == 'Articulado'
    assert any('apoyo de rodillos' in t for _, t in correr(EJ8)['avisos'])


def test_rodillo_y_articulacion_superior_con_vigas_flexibles():
    """El solver de marco plano da lo mismo que el modelo de cortante en el límite de vigas rígidas."""
    for extra in (dict(rod='2'), dict(art_sup='2'), dict(artic='2', art_sup='1')):
        d = copy.deepcopy(EJ4)
        d['niveles'] = [dict(n, viga='1000x1000') for n in d['niveles']]
        d['niveles'][0].update(extra)
        d['niveles'][0]['seccion'] = 'I:213333/A:100000000/I2:213333'                    # axialmente rígida
        d['niveles'][1]['seccion'] = d['niveles'][2]['seccion'] = 'I:213333/A:100000000/I2:213333'
        flex = ame.preparar_modelo(d)
        rig = ame.preparar_modelo(dict(d, vigas='rigidas'))
        assert flex['K'] == pytest.approx(rig['K'], rel=1e-4, abs=1.0), extra   # K ~ 1e4 kg/cm


def test_articulada_arriba_y_abajo_no_resiste_y_errores_de_apoyos():
    d = copy.deepcopy(EJ2)
    d['niveles'][1].update(artic='2', art_sup='2')
    p = ame.preparar_modelo(d)
    assert p['geom']['niveles'][1]['cols'][1]['k'] == 0
    assert any('articulada arriba y abajo' in t for _, t in ame.calcular(p)['avisos'])
    d = copy.deepcopy(EJ2)
    d['niveles'][1].update(rod='1')
    with pytest.raises(ValueError, match='solo aplican en la base'):
        ame.preparar_modelo(d)
    d = copy.deepcopy(EJ2)
    d['niveles'][0].update(rod='2', artic='2')
    with pytest.raises(ValueError, match='articulado y como rodillo'):
        ame.preparar_modelo(d)
    d = copy.deepcopy(EJ2)
    d['niveles'][0]['rod'] = '1 2 3'                                           # ninguna columna resiste: mecanismo
    with pytest.raises(ValueError, match='inestable'):
        correr(d)


def test_giro_de_secciones_con_arroba_90():
    a = parse_secciones('20x40', 'x')[0]
    b = parse_secciones('20x40@90', 'x')[0]
    assert (b['Ix'], b['Iy'], b['texto']) == (a['Iy'], a['Ix'], '20x40@90')
    perfil = parse_secciones('I:3688/A:26.8/I2:98@90', 'x')[0]
    assert (perfil['Ix'], perfil['Iy']) == (98, 3688)
    d = copy.deepcopy(EJ2)
    d['niveles'] = [dict(n, seccion='25x50 25x50@90 25x50 25x50@90') if False else dict(n, seccion='25x50@90') for n in d['niveles']]
    girado = ame.preparar_modelo(d)['info'][0]['k']
    assert girado == pytest.approx(ame.preparar_modelo(dict(EJ2, direccion='Y'))['info'][0]['k'])
    en_m = copy.deepcopy(EJ2)
    en_m['niveles'] = [dict(n, seccion='0.25x0.5@90') for n in en_m['niveles']]       # el 90 del giro no se escala
    assert ame.preparar_modelo(dict(en_m, unidad_seccion='m'))['info'][0]['k'] == pytest.approx(girado)


def test_problema_2a_peso_total_a_partir_del_periodo_medido():
    r = correr(EJ9)
    p = r['p']
    K = 4 * 12 * 2.1e6 * 3688 / 530 ** 3                                       # 4 columnas IR 305x21.1, base y techo rígidos
    m = K * 0.17 ** 2 / (4 * np.pi ** 2)
    assert p['K'][0, 0] == pytest.approx(K)
    assert r['T'][0] == pytest.approx(0.17)
    assert p['cargas'][0] == pytest.approx(m * 981 / 1000)                      # ≈ 1.793 t
    aj = p['ajuste_periodo']
    assert aj['W_total'] == pytest.approx(1.7932, abs=1e-3) and aj['T_medido'] == 0.17
    assert any('Periodo medido' in t for _, t in r['avisos'])
    # con cargas dadas el periodo medido solo fija la escala: el peso total sigue saliendo del periodo
    otro = correr(dict(EJ9, niveles=[dict(EJ9['niveles'][0], W='500')]))
    assert otro['p']['cargas'][0] == pytest.approx(m * 981 / 1000) and otro['T'][0] == pytest.approx(0.17)
    assert 'ajuste_periodo' not in correr(EJ2)['p']
    with pytest.raises(ValueError, match='mayor que cero'):
        ame.preparar_modelo(dict(EJ9, T_medido='0'))


def test_problema_2a_perfil_ir_con_el_eje_debil_en_la_direccion_del_sismo():
    fuerte = correr(EJ9)['p']['cargas'][0]
    debil = correr(dict(EJ9, niveles=[dict(EJ9['niveles'][0], seccion='I:3688/A:26.8/I2:98@90')]))['p']['cargas'][0]
    assert debil / fuerte == pytest.approx(98 / 3688)                           # k ∝ I → el mismo T implica menos masa
    mixto = ame.preparar_modelo(dict(EJ9, niveles=[dict(
        EJ9['niveles'][0], seccion='I:3688/A:26.8/I2:98 I:3688/A:26.8/I2:98@90 I:3688/A:26.8/I2:98 I:3688/A:26.8/I2:98@90')]))
    assert mixto['info'][0]['I'] == pytest.approx(2 * 3688 + 2 * 98)


def test_problema_2b_operacion_continua_con_el_espectro_de_base_de_diseno():
    r = correr(EJ9)
    se = r['estatico']
    assert r['p']['estado'] == 'Limitación de daños' and r['p']['objetivo']['desempeno'].startswith('Operación continua')
    assert se['FC'] == 1.0 and se['C'] == pytest.approx(0.25 * 975 / 981)       # Sa(BD)·Ks en el plato
    assert se['V_est'] == pytest.approx(se['C'] * r['p']['cargas'][0])          # V = C·W, sin factor de carga
    assert r['normativa']['estatico']['aplica']


def test_normalizacion_de_las_formas_modales_no_cambia_las_fuerzas():
    base = correr(EJ4)
    for opcion, verifica in (
            ('Último nivel (azotea) = 1', lambda fi: np.allclose(fi[-1], 1)),
            ('Valor máximo = 1', lambda fi: np.allclose(np.abs(fi).max(axis=0), 1)),
            ('Masa unitaria (φᵀMφ = 1)', None)):
        r = correr(dict(EJ4, norm_modal=opcion))
        assert r['F_final'] == pytest.approx(base['F_final'], rel=1e-9) and r['V_din'] == pytest.approx(base['V_din'])
        assert r['T'] == pytest.approx(base['T']) and r['elastico']['V_el'] == pytest.approx(base['elastico']['V_el'])
        if verifica:
            assert verifica(r['fi'])
        else:
            assert np.allclose(np.diag(r['fi'].T @ r['p']['M'] @ r['fi']), 1)
    assert np.allclose(base['fi'][0], 1)                                       # por defecto: primer nivel = 1


def test_momento_de_volteo_por_entrepiso():
    r = correr(EJ8)
    t = r['tabla_din']
    Fu, h = t['Fu'], t['h']                                                    # de arriba hacia abajo
    assert t['Mv'][0] == pytest.approx(Fu[0] * 3.0)                            # entrepiso superior: Fu·(h entrepiso)
    assert t['Mv'][1] == pytest.approx(Fu[0] * h[0] + Fu[1] * h[1])            # base: ΣFu·z
    assert t['Mv'][-1] == pytest.approx(t['Mvu'].sum())
    for tabla in (r['tabla_est'], r['tabla_din_esc'], r['elastico']['tabla'], r['elastico']['est']['tabla']):
        assert tabla['Mv'][-1] == pytest.approx(tabla['Mvu'].sum())


def test_ejemplo_del_problema_1_dinamico_a_mano():
    """Dos niveles, k = 16165.7 y 57640.7 kg/cm, W = 20 y 10 t: T, γ y las fuerzas con el espectro del ejemplo."""
    r = correr(EJ8)
    k1, k2 = 16165.7, 57640.7
    m = np.diag([20, 10]) * 1000 / 981
    K = np.array([[k1 + k2, -k2], [-k2, k2]])
    lam = np.sort(np.linalg.eigvals(np.linalg.solve(m, K)).real)
    assert r['T'] == pytest.approx(2 * np.pi / np.sqrt(lam), rel=1e-4)
    assert r['gam'].sum() == pytest.approx(1.0)
    assert r['p']['estado'] == 'Seguridad de vida' and r['Qs'][0] == 2 and r['p']['gamma_c'] == 1
    assert r['p']['subgrupo'] == 'A2' and r['p']['intensidad'] == 'Infrecuente'


def test_marco_mixto_concreto_y_acero_por_nivel():
    """Examen: C1 50x50 de concreto (221.36 t/cm²) abajo y C2 = IR 305x21.1 de acero (2.1e6 kg/cm²) arriba."""
    p = ame.preparar_modelo(dict(EJ10, vigas='rigidas'))
    Ec, Es, I50 = 221_359.44, 2.1e6, 50 ** 4 / 12
    esperado = [2 * 12 * Ec * I50 / 450 ** 3 + 2 * 3 * Ec * I50 / 450 ** 3, 4 * 12 * Ec * I50 / 350 ** 3,
                4 * 12 * Ec * I50 / 350 ** 3, 4 * 12 * Es * 3688 / 350 ** 3]
    assert [r['k'] for r in p['info']] == pytest.approx(esperado, rel=1e-6)
    assert [l['cols'][0]['E'] for l in p['geom']['niveles']] == pytest.approx([Ec, Ec, Ec, Es], rel=1e-6)
    flex = correr(EJ10)
    assert flex['p']['vigas'] == 'flexibles' and flex['T'][0] > correr(dict(EJ10, vigas='rigidas'))['T'][0]
    assert flex['p']['K'] == pytest.approx(flex['p']['K'].T) and np.all(np.linalg.eigvalsh(flex['p']['K']) > 0)
    sin = copy.deepcopy(EJ10)
    sin['niveles'][3]['E_col'] = ''                                          # sin E propio: usa el de concreto
    assert ame.preparar_modelo(dict(sin, vigas='rigidas'))['info'][3]['k'] == pytest.approx(esperado[3] * Ec / Es)
    acero = copy.deepcopy(EJ10)
    acero['niveles'][3]['E_col'] = 'acero'
    assert ame.preparar_modelo(dict(acero, vigas='rigidas'))['info'][3]['k'] == pytest.approx(esperado[3] * 2.04 / 2.1)


# ---------------------------------------------------------------- matriz del pizarrón (A1 / A2) y volados
@pytest.mark.parametrize('sub, intensidad, clave', [
    ('B', 'Frecuente', 'DL'), ('B', 'Ocasional', 'OI'), ('B', 'Base de diseño', 'SV'), ('B', 'Infrecuente', 'SV'),
    ('A2', 'Frecuente', 'DL'), ('A2', 'Ocasional', 'DL'), ('A2', 'Base de diseño', 'OI'), ('A2', 'Infrecuente', 'SV'),
    ('A1', 'Frecuente', 'DL'), ('A1', 'Ocasional', 'DL'), ('A1', 'Base de diseño', 'DL'), ('A1', 'Infrecuente', 'OI')])
def test_matriz_del_pizarron_distingue_A1_y_A2(sub, intensidad, clave):
    assert ntc.objetivo_diseno(ntc.SUBGRUPOS[sub], intensidad, sub)['clave'] == clave
    assert correr(dict(EJ2, subgrupo=sub, intensidad=intensidad, espectros={
        'Base de diseño': dict(a0='224', c='975', Ta='0.8', Tb='1.7', k='0.445', Ts='1.0'),
        'Infrecuente': dict(a0='300', c='1300', Ta='0.8', Tb='1.7', k='0.445', Ts='1.0')}))['clave'] == clave


def test_B_en_MC_es_prevencion_de_colapso_y_A1_en_SBD_avisa_su_supuesto():
    assert ntc.objetivo_diseno('B', 'Infrecuente', 'B')['desempeno'] == 'Prevención de colapso'
    assert 'confirma con la norma' in ntc.objetivo_diseno('A', 'Base de diseño', 'A1')['nota']
    assert ntc.objetivo_diseno('A', 'Base de diseño')['clave'] == 'OI'                # "A" sin subgrupo = A2


def test_volados_masas_y_rigidez_como_en_el_pizarron():
    """W1 = 4(3)+2(6)+2(3) = 30, W2 = 4(6)+2(3) = 30, W3 = 3(9) = 27 t; K con rodillo (0), empotrada (12EI/h³), articulada (3EI/h³)."""
    p = ame.preparar_modelo(EJ11)
    assert p['cargas'] == pytest.approx([30, 30, 27])
    assert np.diag(p['M']) == pytest.approx(np.array([30000, 30000, 27000]) / 981)
    E, I = 158_000.0, 40 ** 4 / 12
    k12 = lambda h: 12 * E * I / (h * 100) ** 3
    k3 = lambda h: 3 * E * I / (h * 100) ** 3
    ks = [0 + k12(3) + k3(3), 3 * k12(4), 2 * k12(3.5)]
    K = np.array([[ks[0] + ks[1], -ks[1], 0], [-ks[1], ks[1] + ks[2], -ks[2]], [0, -ks[2], ks[2]]])
    assert p['K'] == pytest.approx(K)
    assert p['geom']['volados'] == (3.0, 2.0) and len(p['geom']['spans']) == 4


def test_volados_no_cambian_la_rigidez_lateral_con_vigas_flexibles():
    d = copy.deepcopy(EJ11)
    d['vigas'] = 'flexibles'
    d['niveles'] = [dict(n, viga='30x60') for n in d['niveles']]
    con = ame.preparar_modelo(d)
    sin = copy.deepcopy(d)
    sin.update(volado_izq='', volado_der='')
    sin['niveles'] = [dict(n, cargas='2 2') for n in sin['niveles']]
    assert con['K'] == pytest.approx(ame.preparar_modelo(sin)['K'])                 # un volado isostático no aporta rigidez


def test_cargas_con_volados_errores_y_una_sola_carga():
    d = copy.deepcopy(EJ11)
    d['niveles'][0]['cargas'] = '2 2 4'
    with pytest.raises(ValueError, match='4 tramo'):
        ame.preparar_modelo(d)
    d['niveles'][0]['cargas'] = '1'
    assert ame.preparar_modelo(d)['cargas'][0] == pytest.approx(1 * (3 + 6 + 3 + 2))   # una carga para todos los tramos
    d['volado_izq'] = '-1'
    with pytest.raises(ValueError, match='no pueden ser negativos'):
        ame.preparar_modelo(d)


def test_volados_se_dibujan_y_se_exportan():
    r = correr(EJ11)
    fig = ame.fig_estructura(r['p'])
    buf = io.BytesIO()
    fig.savefig(buf, format='png')
    assert len(buf.getvalue()) > 5_000
    out = io.BytesIO()
    ame.exportar_excel(r, out)
    pdf = io.BytesIO()
    ame_pdf.memoria_pdf(r, pdf, proyecto='Volados', autor='Pruebas')
    assert out.getvalue()[:2] == b'PK' and pdf.getvalue().startswith(b'%PDF')


# ---------------------------------------------------------------- columnas con distinto EI, piso sin losa en una crujía
K_HOJA = np.array([[167028.5680, -129072.5598, 0, 0], [-129072.5598, 225944.7184, -96804.4198, -67.7388],
                   [0, -96804.4198, 98430.1504, -1625.7306], [0, -67.7388, -1625.7306, 1693.4694]])


def con_acero_en_3EI(d):
    d = copy.deepcopy(d)
    d['niveles'][2]['artic'] = '4'          # columna 12, que cruza del piso 2 al 4
    d['niveles'][3]['artic'] = '1 2 3'      # columnas 13, 14 y 15
    return d


def test_tarea_masas_y_K_como_la_hoja():
    p = ame.preparar_modelo(EJ12)
    assert p['cargas'] == pytest.approx([18.9, 27.0, 18.0, 24.3])      # el piso 3 solo lleva losa en las crujías 1 y 2
    assert np.diag(p['M']) == pytest.approx([19.2661, 27.5229, 18.34862, 24.7706], abs=2e-4)
    p2 = ame.preparar_modelo(con_acero_en_3EI(EJ12))
    assert p2['K'] == pytest.approx(K_HOJA, rel=2e-5, abs=0.05)         # K de la hoja (acero con 3EI/h³)


def test_columna_que_cruza_un_piso_sin_viga():
    p = ame.preparar_modelo(EJ12)
    cols3 = {c['eje']: c for c in p['geom']['niveles'][2]['cols']}
    cols4 = {c['eje']: c for c in p['geom']['niveles'][3]['cols']}
    L = 700.0
    assert cols3[4]['hasta'] == 4 and cols4[4]['k_cont'] and cols4[4]['k'] == 0
    assert cols3[4]['k'] == pytest.approx(12 * 2.1e6 * 3688 / L ** 3)    # una sola columna de 7 m, 12EI/L³
    assert ame.apoyo_txt(cols4[4]) == 'Continuación de la columna del nivel inferior'
    assert 'cruza hasta el piso 4' in ame.apoyo_txt(cols3[4])
    K = p['K']
    assert K[1, 3] == pytest.approx(-cols3[4]['k'])
    assert K[2, 3] == pytest.approx(-sum(c['k'] for c in cols4.values() if c['eje'] != 4))
    assert any('cruza esos pisos' in n for n in p['notas_modelo'])
    assert any('se ignoró su carga' in n for n in p['notas_modelo'])


def test_E_por_columna_y_por_viga():
    p = ame.preparar_modelo(EJ12)
    Ec, Es = 221_359.44, 2.1e6
    assert [c['E'] for c in p['geom']['niveles'][2]['cols']] == pytest.approx([Ec, Ec, Ec, Es])
    assert [c['E'] for c in p['geom']['niveles'][3]['cols']] == pytest.approx([Es] * 4)
    assert [l['viga']['E'] for l in p['geom']['niveles']] == pytest.approx([Ec, Ec, Es, Es])
    d = copy.deepcopy(EJ12)
    d['niveles'][2]['E_col'] = 'general general general acero'
    assert ame.preparar_modelo(d)['geom']['niveles'][2]['cols'][3]['E'] == pytest.approx(2_040_000)
    d['niveles'][2]['E_col'] = '221 221'
    with pytest.raises(ValueError, match='escribiste 2 valores de E'):
        ame.preparar_modelo(d)


def test_tramo_sin_losa_pone_la_carga_en_cero_y_valida():
    d = copy.deepcopy(EJ2)
    d['niveles'][1]['sin_losa'] = '2'
    p = ame.preparar_modelo(d)
    assert p['cargas'][1] == pytest.approx(4.6667 * 4)                 # solo la crujía 1 (4 m) tiene losa
    d['niveles'][1]['sin_losa'] = '9'
    with pytest.raises(ValueError, match='el tramo 9 no existe'):
        ame.preparar_modelo(d)


def test_columna_sin_arriostrar_da_error_claro():
    d = copy.deepcopy(EJ12)
    d['niveles'][3]['ejes'] = '1 2 3'                                  # el eje 4 termina en un piso sin viga
    with pytest.raises(ValueError, match='Eje 4: en el piso del nivel 3 no hay viga ni losa'):
        ame.preparar_modelo(d)


def test_columna_que_cruza_con_vigas_flexibles_coincide_en_el_limite_de_vigas_rigidas():
    d = copy.deepcopy(EJ12)
    d['vigas'] = 'flexibles'
    for n in d['niveles']:
        n['viga'], n['E_viga'] = 'I:100000000000/A:100000000000', '221.35944'
        n['seccion'] = n['seccion'].replace('50x50', 'I:520833.3333/A:100000000').replace('I:3688/A:26.8', 'I:3688/A:100000000')
    flex = ame.preparar_modelo(d)['K']
    rig = ame.preparar_modelo(dict(d, vigas='rigidas'))['K']
    assert flex == pytest.approx(rig, rel=1e-4, abs=1.0)
    real = ame.preparar_modelo(dict(EJ12, vigas='flexibles'))
    assert real['K'] == pytest.approx(real['K'].T) and np.all(np.linalg.eigvalsh(real['K']) > 0)
    assert correr(dict(EJ12, vigas='flexibles'))['T'][0] > correr(EJ12)['T'][0]


def test_calculadora_de_EI():
    ir = ame.propiedades_miembros('I:3688/A:26.8/I2:98', '2.1e6', 'kg/cm²', 'X', 'cm', 3.5)[0]
    assert ir['EI'] == pytest.approx(2.1e6 * 3688) and ir['EI_tm2'] == pytest.approx(774.48)
    assert ir['I_otra'] == 98 and ir['r'] == pytest.approx(np.sqrt(3688 / 26.8))          # rx = 11.7 cm
    assert ir['k12'] == pytest.approx(12 * ir['EI'] / 350 ** 3) and ir['k3'] == pytest.approx(ir['k12'] / 4)
    giro = ame.propiedades_miembros('I:3688/A:26.8/I2:98', '2100', 't/cm²', 'Y')[0]
    assert giro['I'] == 98 and giro['EI'] == pytest.approx(2.1e6 * 98) and 'k12' not in giro
    c = ame.propiedades_miembros('50x50 Ø60', '221.35944', 't/cm²', 'X')
    assert c[0]['EI_tm2'] == pytest.approx(221359.44 * 50 ** 4 / 12 / 1e7) and c[1]['I'] == pytest.approx(np.pi * 60 ** 4 / 64)
    assert ame.propiedades_miembros('0.5x0.5', '221.35944', 't/cm²', 'X', 'm')[0]['I'] == pytest.approx(50 ** 4 / 12)
    with pytest.raises(ValueError, match='E debe ser mayor que cero'):
        ame.propiedades_miembros('50x50', '0')


def test_tarea_dibuja_y_exporta():
    r = correr(EJ12)
    buf = io.BytesIO()
    ame.fig_estructura(r['p']).savefig(buf, format='png')
    assert len(buf.getvalue()) > 5_000
    out, pdf = io.BytesIO(), io.BytesIO()
    ame.exportar_excel(r, out)
    ame_pdf.memoria_pdf(r, pdf, proyecto='Tarea', autor='Pruebas')
    assert out.getvalue()[:2] == b'PK' and pdf.getvalue().startswith(b'%PDF')


def test_apoyos_de_la_base_por_eje():
    assert ame.estados_apoyo('3', '1', 3) == ['Rodillo', 'Empotrado', 'Articulado']
    assert ame.estados_apoyo('', '', 4) == ['Empotrado'] * 4
    assert ame.estados_apoyo('todos', '', 2) == ['Articulado', 'Articulado']
    assert ame.apoyos_a_texto(['Rodillo', 'Empotrado', 'Articulado']) == ([3], [1])
    with pytest.raises(ValueError, match='Apoyo no válido'):
        ame.apoyos_a_texto(['Fijo'])
    # el ida y vuelta reproduce el modelo del pizarrón
    n1 = EJ11['niveles'][0]
    art, rod = ame.apoyos_a_texto(ame.estados_apoyo(n1['artic'], n1['rod'], 3))
    d = copy.deepcopy(EJ11)
    d['niveles'][0].update(artic=' '.join(map(str, art)), rod=' '.join(map(str, rod)))
    assert ame.preparar_modelo(d)['K'] == pytest.approx(ame.preparar_modelo(EJ11)['K'])


# ---------------------------------------------------------------- captura por columna, volados por nivel
def test_captura_por_columna_equivale_a_la_captura_por_nivel():
    a, b = ame.preparar_modelo(EJ12), ame.preparar_modelo(EJ13)
    assert b['K'] == pytest.approx(a['K']) and b['M'] == pytest.approx(a['M']) and b['cargas'] == pytest.approx(a['cargas'])
    assert correr(EJ13)['T'] == pytest.approx(correr(EJ12)['T'])
    cols3 = {c['eje']: c for c in b['geom']['niveles'][2]['cols']}
    assert cols3[4]['hasta'] == 4 and ame.apoyo_txt(cols3[2]) == 'Empotrado'
    assert {c['eje']: ame.apoyo_txt(c) for c in b['geom']['niveles'][0]['cols']} == {
        1: 'Empotrado', 2: 'Articulado', 3: 'Articulado', 4: 'Empotrado'}


def test_ida_y_vuelta_entre_tabla_por_columna_y_por_nivel():
    filas = ame.columnas_desde_niveles(EJ12)
    assert {f['eje'] for f in filas} == {1, 2, 3, 4}
    assert any(f['eje'] == 4 and f['desde'] == 3 and f['hasta'] == 4 for f in filas)       # la columna que cruza es una fila
    por_nivel = {k: v for k, v in EJ12.items() if k != 'niveles'}
    por_nivel['niveles'] = [{k: v for k, v in n.items() if k in ('h', 'viga', 'E_viga', 'sin_losa', 'cargas')}
                            for n in EJ12['niveles']]
    vuelta = ame.preparar_modelo(dict(por_nivel, columnas=filas))
    assert vuelta['K'] == pytest.approx(ame.preparar_modelo(EJ12)['K'])


@pytest.mark.parametrize('cambio, texto', [
    (lambda c: c[0].update(eje='9'), 'el eje 9 no existe'),
    (lambda c: c[0].update(desde='3', hasta='1'), 'deben cumplir 1 ≤ desde ≤ hasta'),
    (lambda c: c[0].update(seccion=''), 'exactamente una sección'),
    (lambda c: c[0].update(seccion='50x50 40x40'), 'exactamente una sección'),
    (lambda c: c[0].update(apoyo='Fijo'), 'apoyo no válido'),
    (lambda c: c[1].update(apoyo='Rodillo'), 'rodillos solo aplican en la base'),
    (lambda c: c.append(dict(eje='1', desde='1', hasta='4', seccion='40x40')), 'tiene dos columnas'),
    (lambda c: None, 'no tiene ninguna columna'),
])
def test_errores_de_la_tabla_por_columna(cambio, texto):
    d = copy.deepcopy(EJ13)
    cambio(d['columnas'])
    if texto == 'no tiene ninguna columna':                       # deja el nivel 4 sin columnas
        d['columnas'] = [f for f in d['columnas'] if f['hasta'] != '4' and f['desde'] != '4']
        d['columnas'] = [dict(f, hasta='3') if f['eje'] == '4' else f for f in d['columnas']]
    with pytest.raises(ValueError, match=texto):
        ame.preparar_modelo(d)


def test_filas_vacias_y_columnas_con_valores_en_blanco():
    d = copy.deepcopy(EJ13)
    d['columnas'].append(dict(eje='', desde='', hasta='', seccion='', E='', apoyo='', art_sup=False))
    d['columnas'][0]['desde'] = ''                                   # vacío = nivel 1
    d['columnas'][0]['apoyo'] = None
    assert ame.preparar_modelo(d)['K'] == pytest.approx(ame.preparar_modelo(EJ13)['K'])


def test_volados_distintos_por_nivel():
    d = copy.deepcopy(EJ11)
    d['volado_izq'], d['volado_der'] = '3', '2'
    d['niveles'][1].update(volado_izq='1', cargas='2 4 2 0')              # nivel 2: volado izquierdo de 1 m
    d['niveles'][2].update(volado_izq='0', volado_der='1', cargas='3 3 1')  # nivel 3: sin volado izq., der. de 1 m
    p = ame.preparar_modelo(d)
    assert p['cargas'] == pytest.approx([30.0, 1 * 2 + 6 * 4 + 3 * 2, 6 * 3 + 3 * 3 + 1 * 1])
    assert p['geom']['volados'] == (3.0, 2.0)
    d['niveles'][1]['cargas'] = '2 2 4'
    with pytest.raises(ValueError, match='Nivel 2: hay 4 tramo'):
        ame.preparar_modelo(d)
    d['niveles'][1].update(cargas='2 2 4 0', volado_izq='-1')
    with pytest.raises(ValueError, match='Nivel 2: los volados no pueden ser negativos'):
        ame.preparar_modelo(d)
    buf = io.BytesIO()
    ame.fig_estructura(ame.preparar_modelo(dict(d, niveles=[dict(n, **({'volado_izq': '1'} if i == 1 else {}))
                                                          for i, n in enumerate(d['niveles'])]))).savefig(buf, format='png')
    assert len(buf.getvalue()) > 5_000


# ---------------------------------------------------------------- resolver al revés (dimensionado)
def con_columnas_25x50():
    d = copy.deepcopy(EJ2)
    d['niveles'] = [dict(n, seccion='25x50') for n in d['niveles']]
    return d


@pytest.mark.parametrize('limite', [0.02, 0.01, 0.005])
def test_dimensionado_por_distorsion_cumple_y_es_minimo(limite):
    d = con_columnas_25x50()
    r = dim.dimensionar_distorsion(d, [1, 2, 3], paso=5, lado_ini=20, limite=limite)
    assert r['cumple'] and max(r['dist']) <= limite and r['razones'].max() <= 1 + 1e-9
    base = dim._base(d)
    for nivel, lado in r['lados'].items():
        if lado > 20:                                                   # un paso menos en ese nivel ya incumple
            menos = dict(r['lados'])
            menos[nivel] -= 5
            res = ame.calcular(ame.preparar_modelo(dim._con_lados(base, menos, 'Cuadrada', 1.0)))
            assert dim._razones(res, 'diseno', limite)[0].max() > 1
    assert all(texto == f'{l:g}x{l:g}' for l, texto in zip(r['lados'].values(), r['textos'].values()))


def test_dimensionado_formas_y_estados():
    d = con_columnas_25x50()
    circ = dim.dimensionar_distorsion(d, [1, 2, 3], forma='Circular', limite=0.01)
    assert circ['cumple'] and all(t.startswith('Ø') for t in circ['textos'].values())
    rect = dim.dimensionar_distorsion(d, [1, 2, 3], forma='Rectangular', rel=1.5, limite=0.01)
    assert rect['cumple'] and all('x' in t for t in rect['textos'].values())
    assert rect['textos'][1] == f'{rect["lados"][1]:g}x{1.5 * rect["lados"][1]:g}'
    todos = dim.dimensionar_distorsion(d, [1, 2, 3], estados='todos')       # SV (0.03), DL (0.004) y OI (0.005) a la vez
    assert todos['cumple'] and len(todos['usados']) == 3 and max(todos['dist']) <= 0.03
    assert max(todos['lados'].values()) >= max(dim.dimensionar_distorsion(d, [1, 2, 3], estados='diseno')['lados'].values())


def test_dimensionado_por_periodo():
    d = con_columnas_25x50()
    r = dim.dimensionar_periodo(d, [1, 2, 3], 0.5)
    assert r['T_exacto'] == pytest.approx(0.5, abs=1e-6) and r['T'] <= 0.5 and r['lado'] >= r['lado_exacto']
    assert r['lado'] - r['lado_exacto'] < 5 and r['lado'] % 5 == 0
    # con la sección exacta, el periodo del modelo completo es el pedido
    d2 = dim._con_lados(dim._base(d), {1: r['lado_exacto'], 2: r['lado_exacto'], 3: r['lado_exacto']}, 'Cuadrada', 1.0)
    assert ame.calcular(ame.preparar_modelo(d2))['T'][0] == pytest.approx(0.5, abs=1e-6)


@pytest.mark.parametrize('llamada, texto', [
    (lambda d: dim.dimensionar_distorsion(d, [], limite=0.01), 'al menos un nivel'),
    (lambda d: dim.dimensionar_distorsion(d, [9], limite=0.01), 'El nivel 9 no existe'),
    (lambda d: dim.dimensionar_distorsion(d, [1, 2, 3], forma='Hexagonal', limite=0.01), 'Forma no válida'),
    (lambda d: dim.dimensionar_distorsion(d, [1, 2, 3], paso=0, limite=0.01), 'paso debe ser mayor'),
    (lambda d: dim.dimensionar_distorsion(d, [1, 2, 3], forma='Rectangular', rel=0, limite=0.01), 'relación h/b'),
    (lambda d: dim.dimensionar_periodo(d, [1, 2, 3], 0.001), 'no se puede llegar a'),
    (lambda d: dim.dimensionar_periodo(d, [1, 2, 3], -1), 'mayor que cero'),
    (lambda d: dim.dimensionar_distorsion(dict(d, unidad_seccion='m'), [1, 2, 3], limite=0.01), 'secciones en cm'),
    (lambda d: dim.dimensionar_distorsion(dict(d, modo='matrices'), [1, 2, 3], limite=0.01), 'matrices K y M'),
])
def test_dimensionado_errores(llamada, texto):
    with pytest.raises(ValueError, match=texto):
        llamada(con_columnas_25x50())


def test_dimensionado_critico_fuera_de_los_niveles_elegidos():
    d = con_columnas_25x50()
    d['niveles'][1]['seccion'] = '15x15'                                    # el nivel 2 es el flexible y no se dimensiona
    with pytest.raises(ValueError, match='entrepiso crítico es el 2'):
        dim.dimensionar_distorsion(d, [1, 3], limite=0.005)


def test_dimensionado_con_captura_por_columna_y_el_ejemplo_de_la_tarea():
    r = dim.dimensionar_distorsion(EJ13, [1, 2], limite=0.02)
    assert r['cumple'] and set(r['lados']) == {1, 2}
    with pytest.raises(ValueError, match='entrepiso crítico es el 4'):       # el nivel 4 es de acero y no se dimensiona
        dim.dimensionar_distorsion(EJ13, [1, 2], limite=0.005)


def test_el_pdf_ya_no_trae_la_seccion_final_de_alcances():
    pymupdf = pytest.importorskip('pymupdf')
    buf = io.BytesIO()
    ame_pdf.memoria_pdf(correr(EJ2), buf, proyecto='Prueba', autor='Pruebas')
    texto = ''.join(pg.get_text() for pg in pymupdf.open(stream=buf.getvalue(), filetype='pdf'))
    assert 'Alcances y limitaciones' not in texto and 'revisión de un ingeniero' not in texto
    assert 'Avisos de coherencia' in texto


def test_pagina_organizada_en_tabs_y_errores_no_la_rompen():
    pytest.importorskip('streamlit')
    from streamlit.testing.v1 import AppTest
    raiz = Path(__file__).resolve().parents[1]
    at = AppTest.from_file(str(raiz / 'app_web.py'), default_timeout=120).run()
    etiquetas = [t_.label for t_ in at.tabs]
    assert etiquetas[:3] == ['1 · Estructura', '2 · Sismo y norma', '3 · Resultados']
    assert {'Resumen', 'Calculadora de EI', 'Dimensionar · resolver al revés', '4 · Herramientas'} <= set(etiquetas)
    assert sorted({w.key for w in at.sidebar.text_input}) == ['autor', 'proyecto']          # la barra lateral solo trae el proyecto
    # datos inválidos: el error aparece, las demás pestañas siguen disponibles y la calculadora funciona
    at.text_input(key='crujias').set_value('x').run()
    assert not at.exception and any('Revisa los datos' in e.value for e in at.error)
    assert 'Calculadora de EI' in [t_.label for t_ in at.tabs] and len(at.session_state['calc_df']) >= 1
    assert not any(w_.value.startswith('Fila') for w_ in at.warning)               # la tabla de ejemplo no da errores
    at.text_input(key='crujias').set_value('4 5').run()
    assert not at.exception and not any('Revisa los datos' in e.value for e in at.error)


def test_calculadora_por_filas_cada_seccion_con_su_E():
    filas = [dict(seccion='I:3688/A:26.8/I2:98', E='2.1e6', unidad_E='kg/cm²', orientacion='Eje fuerte (Ix)', h=''),
             dict(seccion='50x50', E='221.35944', unidad_E='t/cm²', orientacion='Eje fuerte (Ix)', h='3.5'),
             dict(seccion='I:3688/A:26.8/I2:98', E='2100', unidad_E='t/cm²', orientacion='Eje débil (Iy)', h='3.5'),
             dict(seccion='', E='', unidad_E='', orientacion='', h='')]                  # fila vacía: se omite
    res, errores = ame.propiedades_tabla(filas)
    assert not errores and [r['fila'] for r in res] == [1, 2, 3]
    assert res[0]['EI'] == pytest.approx(2.1e6 * 3688) and 'k12' not in res[0]
    assert res[1]['E'] == pytest.approx(221359.44) and res[1]['EI'] == pytest.approx(221359.44 * 50 ** 4 / 12)
    assert res[1]['k12'] == pytest.approx(12 * res[1]['EI'] / 350 ** 3)
    assert res[2]['I'] == 98 and res[2]['orientacion'] == 'Iy' and res[2]['EI'] == pytest.approx(2.1e6 * 98)
    assert [r['unidad_E'] for r in res] == ['kg/cm²', 't/cm²', 't/cm²']
    # una fila con varias secciones usa el mismo E para todas
    multiple, _ = ame.propiedades_tabla([dict(seccion='30x30 40x40', E='200000')])
    assert [r['fila'] for r in multiple] == [1, 1] and multiple[1]['EI'] == pytest.approx(200000 * 40 ** 4 / 12)
    # n filas: se pueden agregar tantas como se quiera
    muchas, err = ame.propiedades_tabla([dict(seccion=f'{20 + 5 * i}x{20 + 5 * i}', E='221359') for i in range(30)])
    assert len(muchas) == 30 and not err and muchas[-1]['I'] == pytest.approx(165 ** 4 / 12)


def test_calculadora_por_filas_errores_por_fila_sin_detener_las_demas():
    res, errores = ame.propiedades_tabla([dict(seccion='50x50', E='221359'), dict(seccion='redonda', E='221359'),
                                          dict(seccion='40x40', E=''), dict(seccion='', E='2e6'),
                                          dict(seccion='30x30', E='2e6', h='-3'), dict(seccion='30x30', E='0'),
                                          dict(seccion='30x30', E='2e6', unidad_E='psi')])
    assert [r['fila'] for r in res] == [1]
    assert [i for i, _ in errores] == [2, 3, 4, 5, 6, 7]
    mensajes = ' | '.join(m for _, m in errores)
    for esperado in ('no es válida', 'falta E', 'falta la sección', 'altura debe ser mayor', 'E debe ser mayor que cero',
                     'Unidad no válida'):
        assert esperado in mensajes


# ---------------------------------------------------------------- E e I de cada columna (también las que cruzan niveles)
def test_I_y_E_directos_por_columna_equivalen_a_la_seccion():
    d = copy.deepcopy(EJ13)
    for f in d['columnas']:
        if f['eje'] == '4' and f['desde'] == '3':
            f.update(seccion='', I='3688', A='26.8', E='2100')               # la columna 12, que cruza los niveles 3 y 4
    assert ame.preparar_modelo(d)['K'] == pytest.approx(ame.preparar_modelo(EJ13)['K'])
    d['columnas'][-1].update(I='5000')                                       # editar solo su I: la rigidez cambia
    assert abs(ame.preparar_modelo(d)['K'][1, 3]) == pytest.approx(abs(ame.preparar_modelo(EJ13)['K'][1, 3]) * 5000 / 3688)


@pytest.mark.parametrize('cambio, texto', [
    (dict(I='-5'), 'la inercia debe ser mayor que cero'), (dict(I='0'), 'la inercia debe ser mayor que cero'),
    (dict(I='abc'), 'Valor numérico no válido'), (dict(seccion='', I='', A='10'), 'escribiste el área A pero no la inercia I'),
    (dict(I='3688', A='-2'), 'el área debe ser mayor que cero'), (dict(seccion='', I=''), 'exactamente una sección')])
def test_errores_de_I_y_A_por_columna(cambio, texto):
    d = copy.deepcopy(EJ13)
    d['columnas'][-1].update(cambio)
    with pytest.raises(ValueError, match=texto):
        ame.preparar_modelo(d)


def test_tabla_de_valores_efectivos_por_columna():
    filas = ame.tabla_columnas(ame.preparar_modelo(EJ12))
    assert len(filas) == 15                                                  # 16 tramos de columna − 1 continuación (la 12 cruza)
    c12 = [f for f in filas if f['eje'] == 4 and f['desde'] == 3][0]
    assert (c12['hasta'], c12['L']) == (4, 7.0) and c12['E'] == pytest.approx(2.1e6) and c12['I'] == 3688
    assert c12['EI_tm2'] == pytest.approx(774.48) and c12['k'] == pytest.approx(12 * 2.1e6 * 3688 / 700 ** 3)
    assert 'cruza hasta el piso 4' in c12['apoyo']
    c1 = [f for f in filas if f['eje'] == 1 and f['desde'] == 1][0]
    assert c1['L'] == 4.5 and c1['I'] == pytest.approx(50 ** 4 / 12) and c1['A'] == pytest.approx(2500)
    assert [f['apoyo'] for f in filas if f['desde'] == 1] == ['Empotrado', 'Articulado', 'Articulado', 'Empotrado']


def test_pagina_cambiar_de_modo_conserva_la_estructura_y_los_apoyos():
    pytest.importorskip('streamlit')
    from streamlit.testing.v1 import AppTest
    raiz = Path(__file__).resolve().parents[1]
    at = AppTest.from_file(str(raiz / 'app_web.py'), default_timeout=180).run()
    t1 = lambda: re.search(r'class="n">([^<]*)<', [m.value for m in at.markdown if 'PERIODO' in m.value][0]).group(1)
    for idx in (11, 10):                                                     # tarea (cruza niveles) y pizarrón (volados, rodillo)
        at.selectbox(key='ejemplo').select(list(EJ)[idx])
        at.sidebar.button[0].click()
        at.run()
        inicial = t1()
        at.radio(key='captura').set_value('Por columna (una fila por columna)').run()
        assert not at.exception and t1() == inicial
        assert {'I', 'A', 'E'} <= set(at.session_state['df_cols'].columns)
        at.radio(key='captura').set_value('Por nivel (una fila por nivel)').run()
        assert not at.exception and t1() == inicial                          # los apoyos de la base también se conservan
    at.selectbox(key='apoyo_1').select('Articulado').run()
    cambiado = t1()
    at.radio(key='captura').set_value('Por columna (una fila por columna)').run()
    assert t1() == cambiado                                                  # un apoyo cambiado a mano pasa a la otra tabla


# ---------------------------------------------------------------- exportación
def test_excel_y_pdf_se_generan():
    from openpyxl import load_workbook
    for ej in (EJ1, EJ4):
        r = correr(ej)
        buf = io.BytesIO()
        ame.exportar_excel(r, buf)
        hojas = load_workbook(io.BytesIO(buf.getvalue())).sheetnames
        assert {'Datos', 'Modal espectral', 'Estático', 'Normativa', 'Revisiones', 'Gráficas'} <= set(hojas)
        assert not {'Elástico', 'Comparación'} & set(hojas)
        pdf = io.BytesIO()
        ame_pdf.memoria_pdf(r, pdf, proyecto='Prueba', autor='Pruebas')
        assert pdf.getvalue().startswith(b'%PDF') and len(pdf.getvalue()) > 30_000


def test_pdf_con_formulas_y_letras_griegas():
    pymupdf = pytest.importorskip('pymupdf')
    buf = io.BytesIO()
    ame_pdf.memoria_pdf(correr(EJ2), buf, proyecto='Torre Ø', autor='Á')
    texto = ''.join(pg.get_text() for pg in pymupdf.open(stream=buf.getvalue(), filetype='pdf'))
    for esperado in ('Memoria de cálculo', 'λ', 'γ', 'Σ', 'T < Tₐ', 'Torre Ø', 'Marco normativo (NTC-S 2023)',
                     'Cortante basal mínimo', 'Cimentación', 'ANÁLISIS MODAL ESPECTRAL', 'ANÁLISIS ESTÁTICO',
                     'Matriz de rigidez K', 'Factores de participación modal'):
        assert esperado in texto
    assert 'elástico' not in texto.lower()


def test_figuras_se_dibujan():
    r = correr(EJ3)
    for f in (ame.fig_estructura(r['p']), ame.fig_modal(r), ame.fig_espectro(r), ame.fig_fuerzas(r), ame.fig_diagramas(r),
              ame.fig_distorsiones(r)):
        buf = io.BytesIO()
        f.savefig(buf, format='png')
        assert len(buf.getvalue()) > 5_000


# ---------------------------------------------------------------- lectura flexible de matrices (modo manual)
def test_lectura_flexible_de_matrices():
    for texto in ('1 2\n3 4', '1,2;3,4', '[1 2; 3 4]', 'K = 1\t2\n3\t4', '1.0, 2.0\n3.0, 4.0'):
        assert ame.parse_matriz(texto, 'K') == pytest.approx(np.array([[1, 2], [3, 4]]))
    assert ame.parse_matriz('20,985.1162 -14,212.1395', 'K')[0] == pytest.approx([20985.1162, -14212.1395])
    with pytest.raises(ValueError, match='rectangular'):
        ame.parse_matriz('1 2\n3', 'K')


# ---------------------------------------------------------------- página web
def test_pagina_web_corre_con_todos_los_ejemplos():
    pytest.importorskip('streamlit')
    from streamlit.testing.v1 import AppTest
    raiz = Path(__file__).resolve().parents[1]
    at = AppTest.from_file(str(raiz / 'app_web.py'), default_timeout=120).run()
    assert not at.exception
    for nombre in EJ:
        at.sidebar.selectbox(key='ejemplo').select(nombre)
        at.sidebar.button[0].click()
        at.run()
        assert not at.exception, nombre
        # la página debe dar el mismo periodo que el cálculo directo del ejemplo (detecta datos mal cargados en la página)
        t1_pagina = re.search(r'class="n">([^<]*)<', [m.value for m in at.markdown if 'PERIODO' in m.value][0]).group(1)
        assert t1_pagina == f'{correr(EJ[nombre])["T"][0]:.4f} s', nombre
    # controles nuevos (sobre un ejemplo por secciones: con matrices algunos controles se deshabilitan)
    at.sidebar.selectbox(key='ejemplo').select(list(EJ)[1])
    at.sidebar.button[0].click()
    at.run()
    at.selectbox(key='combinacion').select('CQC').run()
    at.radio(key='unidad_espectro').set_value('g').run()                 # solo cambia las unidades
    assert at.error and not at.exception                                           # a0 y c quedan enormes: alerta de captura
    at.radio(key='unidad_espectro').set_value('cm/s²').run()
    at.selectbox(key='modos').select('Los primeros N').run()
    at.text_input(key='modos_n').set_value('2').run()
    assert not at.exception
    at.selectbox(key='modos').select('Todos los modos').run()
    at.selectbox(key='R_dist').select('Del periodo fundamental (estricto)').run()
    assert not at.exception
    at.selectbox(key='estado').select('Limitación de daños').run()
    at.radio(key='direccion').set_value('Y').run()
    at.selectbox(key='material').select('Acero').run()
    assert not at.exception
    # objetivo de diseño: grupo A / base -> ocupación inmediata; infrecuente exige su propio espectro
    at.selectbox(key='subgrupo').select('A1').run()
    assert not at.exception and not any('Revisa los datos' in e.value for e in at.error)
    # captura por columna y dimensionado
    at.sidebar.selectbox(key='ejemplo').select(list(EJ)[12])
    at.sidebar.button[0].click()
    at.run()
    assert not at.exception and at.radio(key='captura').value.startswith('Por columna')
    assert 'Dimensionar · resolver al revés' in [t.label for t in at.tabs]
    at.sidebar.selectbox(key='ejemplo').select(list(EJ)[1])
    at.sidebar.button[0].click()
    at.run()
    [b for b in at.button if b.key == 'dim_ir'][0].click()
    at.run()
    assert not at.exception and any('Todos los entrepisos cumplen' in s_.value for s_ in at.success)
    at.text_input(key='crujias').set_value('4 5').run()
    # apoyos de la base: un selector por eje
    at.sidebar.selectbox(key='ejemplo').select(list(EJ)[1])
    at.sidebar.button[0].click()
    at.run()
    assert at.selectbox(key='apoyo_1').value == 'Empotrado'
    at.selectbox(key='apoyo_1').select('Rodillo').run()
    at.selectbox(key='apoyo_2').select('Articulado').run()
    assert not at.exception and not any('Revisa los datos' in e.value for e in at.error)
    at.sidebar.selectbox(key='ejemplo').select(list(EJ)[10])
    at.sidebar.button[0].click()
    at.run()
    assert [at.selectbox(key=f'apoyo_{i}').value for i in (1, 2, 3)] == ['Rodillo', 'Empotrado', 'Articulado']
    # modo matrices: se puede alternar sin perder los datos de secciones
    at.sidebar.selectbox(key='ejemplo').select(list(EJ)[5])
    at.sidebar.button[0].click()
    at.run()
    assert not at.exception and at.radio(key='modo').value.startswith('Matrices')
    assert {'Resumen', 'Detalle paso a paso', 'Modal espectral', 'Estático'} <= {t.label for t in at.tabs}
    assert not {'Análisis elástico', 'Análisis estático'} & {t.label for t in at.tabs}
    at.radio(key='modo').set_value('Por secciones y cargas').run()
    assert not at.exception and at.text_input(key='crujias').value
    at.selectbox(key='intensidad').select('Infrecuente').run()
    assert at.error and 'espectro SASID' in at.error[0].value and not at.exception
    for k, v in dict(a0='300', c='1300', Ta='0.8', Tb='1.7', k='0.445', Ts='1.0').items():
        at.text_input(key=k).set_value(v)
    at.run()
    assert not at.exception and not any('Revisa los datos' in e.value for e in at.error)
    at.selectbox(key='intensidad').select('Base de diseño').run()
    assert at.text_input(key='a0').value == '80'         # el espectro base (el del ejemplo 6 cargado antes) se conservó
    at.multiselect(key='irreg').set_value(['5.3.2']).run()
    at.text_input(key='b_planta').set_value('12').run()
    at.selectbox(key='zona').select('III').run()
    assert not at.exception
    at.text_input(key='crujias').set_value('').run()
    assert at.error and not at.exception  # mensaje claro, sin caerse


def test_dibujo_de_columnas_capturadas_marca_niveles_sin_columnas_y_se_completa():
    d = dict(crujias='5 5', niveles=[dict(h='4'), dict(h='3'), dict(h='3')],
             columnas=[dict(eje=str(e), desde=str(k), hasta=str(k), seccion='30x30') for e in (1, 2, 3) for k in (1, 2)])
    assert ame.cobertura_columnas(d)[3] == [3]
    assert ame.fig_captura_columnas(d) is not None
    nuevas, hechos = ame.completar_columnas(d['columnas'], 3)
    assert hechos == [3]
    d['columnas'] = nuevas
    assert ame.niveles_desde_columnas(d)['niveles'][2]['ejes'] == '1 2 3'


def test_solucion_del_examen_2_problema_1_coincide_con_el_pdf():
    """Periodos, Sa, γ, fuerzas modales, Fu/Vu/Mvu, cs y factor de escala de la solución del segundo examen parcial."""
    for ej in (EJ12_SOL, EJ13_SOL):
        r = ame.calcular(ame.preparar(copy.deepcopy(ej)))
        assert r['T'] == pytest.approx([0.79109, 0.27385, 0.076919, 0.051217], abs=2e-5)
        assert r['Sa_modal'][:, 0] == pytest.approx([1000.735, 433.5606, 245.1572, 221.4075], abs=0.01)
        assert r['gam'] == pytest.approx([0.055138, 0.77573, 0.11784, 0.051296], abs=2e-5)
        assert r['F'][0] == pytest.approx([1063.1, 6479.7, 556.57, 218.81], abs=0.05)
        assert r['F_final'] == pytest.approx([6593.5, 11426, 8347.3, 28586], abs=0.5)
        t = r['tabla_din']
        assert t['Vu'] * 1000 == pytest.approx([31444, 40626, 53194, 60447], abs=1)
        assert t['Mvu'] * 1e5 == pytest.approx([4.7166e7, 1.0559e7, 1.0055e7, 3.2638e6], rel=1e-4)
        assert r['estatico']['C'] == pytest.approx(1.03102, abs=1e-5)
        assert r['factor_escala'] == pytest.approx(1.6548, abs=1e-4)
        assert r['tabla_din_esc']['Vu'][-1] == pytest.approx(100.0294, abs=1e-3)


def test_detalle_completo_muestra_las_matrices_y_los_pasos_de_la_solucion():
    import ame_detalle
    r = ame.calcular(ame.preparar(copy.deepcopy(EJ12_SOL)))
    modal = ame_detalle.detalle_modal(r)
    pasos = [b['texto'] for b in modal if b['t'] == 'paso']
    assert pasos == ['Datos del problema', 'Matriz de masas M', 'Matriz de rigidez K', 'Periodos y formas modales',
                     'Factores de participación modal', 'Ordenadas espectrales de cada modo', 'Desplazamientos modales',
                     'Fuerzas modales', 'Combinación modal', 'Fuerzas sísmicas de diseño']
    assert all(len(b['expl']) == 2 for b in modal if b['t'] == 'paso')            # explicación en dos líneas
    nombres = [b['nombre'] for b in modal if b['t'] == 'matriz']
    assert {'M', 'K', 'K (simbólica)', 'A = M⁻¹K', 'M⁻¹'} <= set(nombres)
    sim = [b for b in modal if b['t'] == 'matriz' and b['simbolica'] is not None][0]['simbolica']
    assert sim[0][0] == 'k_{1}+k_{2}' and sim[1][3] == '-k_{c4}'                    # la columna del eje 4 cruza del nivel 3 al 4
    K = [b for b in modal if b['t'] == 'matriz' and b['nombre'] == 'K'][0]['datos']
    assert K == pytest.approx(r['p']['K'])
    txt = ame.reporte_texto(r)
    for frase in ('PASO 3 · MATRIZ DE RIGIDEZ K', 'k_{1}+k_{2}', 'Valores propios', 'Σγ = 1.000000', 'PASO 8 · FUERZAS MODALES',
                  'Σ Mvu = 710.4410', 'ANÁLISIS ESTÁTICO', 'Factor = 90.9358 / 54.9521 = 1.654818'):
        assert frase in txt, frase
    est = ame_detalle.detalle_estatico(r)
    assert [b['texto'] for b in est if b['t'] == 'paso'][-1] == 'Factor de escala del análisis modal espectral'
    # con K y M capturadas directamente se explican igual, sin la tabla de columnas
    r6 = ame.calcular(ame.preparar(copy.deepcopy(EJ6)))
    tx6 = ame_detalle.bloques_a_texto(ame_detalle.detalle_modal(r6))
    assert 'viene dada directamente' in tx6 and 'PASO 4' in tx6


def test_ejemplo_14_examen_con_rodillo_y_columnas_de_distinta_seccion():
    p = ame.preparar(copy.deepcopy(EJ14))
    assert p['cargas'] == pytest.approx([20.0, 8.0])                    # techo con 1 T/m solo en las dos primeras crujías
    k_emp = 12 * 158000 * (20 * 40 ** 3 / 12) / 300 ** 3                    # columnas de 20x40, empotradas
    k_30 = 12 * 158000 * (30 ** 4 / 12) / 300 ** 3
    assert p['info'][0]['k'] == pytest.approx(2 * k_emp + k_30 / 4)    # eje 2 articulado abajo (3EI/h³); eje 4 sobre rodillo: 0
    assert p['info'][1]['k'] == pytest.approx(2 * k_emp + k_30)
    r = ame.calcular(p)
    assert r['T'] == pytest.approx([0.27427, 0.10395], abs=2e-5)
    assert 'Eje 4' in ' '.join(t for _, t in r['avisos'])


def test_pagina_recupera_la_captura_con_la_misma_direccion(tmp_path):
    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file(str(Path(ame.__file__).parent / 'app_web.py'), default_timeout=180).run()
    pid = at.query_params['p'][0] if isinstance(at.query_params['p'], list) else at.query_params['p']
    at.selectbox(key='ejemplo').select(list(ame.EJEMPLOS_MODELO)[13])
    at.sidebar.button[0].click()
    at.run()
    at2 = AppTest.from_file(str(Path(ame.__file__).parent / 'app_web.py'), default_timeout=180)
    at2.query_params['p'] = pid
    at2.run()
    assert at2.radio(key='captura').value == at.radio(key='captura').value
    assert at2.text_input(key='crujias').value == '4 4 2'
    assert not at2.exception


def test_Ks_de_la_ecuacion_3_1_1_y_valor_dado_a_mano():
    assert [ntc.ks_norma(t) for t in (0.3, 0.5, 1.0, 1.8)] == pytest.approx([1 / 6, 1 / 6, 0.25, 0.25])
    assert ntc.ks_norma(0.75) == pytest.approx(1 / 5.0) and ntc.ks_norma(0.9) == pytest.approx(1 / 4.4)
    assert [ntc.zona_ks(t) for t in (0.4, 0.8, 1.2)] == ['A', 'B', 'C']
    d = dict(EJ2, intensidad='Frecuente')
    d['espectros'] = {'Base de diseño': dict(a0='114', c='320', Ta='0.9', Tb='2.2', k='0.39', Ts='0.5')}
    auto = correr(dict(d, Ks=''))
    assert auto['p']['Ks'] == pytest.approx(1 / 6) and auto['p']['Ks_auto'] and auto['p']['Ks_zona'] == 'A'
    a_mano = correr(dict(d, Ks='0.25'))
    assert a_mano['p']['Ks'] == 0.25 and not a_mano['p']['Ks_auto']
    assert a_mano['V_din'] == pytest.approx(auto['V_din'] * 0.25 * 6)
    assert 'ec. 3.1.1' in ' '.join(f[1] for f in auto['normativa']['factores'] if f[0] == 'Ks')


def test_en_limitacion_de_dano_los_rotulos_dicen_Sa_por_Ks_y_no_Q_R():
    import io
    import ame_pdf
    r = correr(dict(EJ11, intensidad='Frecuente', Ks=''))
    assert r['clave'] == 'DL' and ame.rotulo_sa(r) == 'Sa(BD)·Ks'
    txt = ame.reporte_texto(r)
    assert 'Sa·Ks (cm/s²)' in txt and "Q'·R'" not in txt
    assert 'sin factor de carga' in txt.lower() and 'FC = 1' in txt
    ame.exportar_excel(r, io.BytesIO())
    ame_pdf.memoria_pdf(r, io.BytesIO(), proyecto='x', autor='y')
    sv = correr(dict(EJ11, intensidad='Base de diseño'))
    assert ame.rotulo_sa(sv) == "Sa/(Q'R')" and "Q'·R'" in ame.reporte_texto(sv)


def test_exportar_modal_y_estatico_por_separado():
    import io
    import ame_pdf
    from openpyxl import load_workbook
    r = correr(EJ12_SOL)
    hojas = {}
    for m in ('modal', 'estatico', 'ambos'):
        buf = io.BytesIO()
        ame.exportar_excel(r, buf, m)
        hojas[m] = set(load_workbook(io.BytesIO(buf.getvalue())).sheetnames)
    assert 'Modal espectral' in hojas['modal'] and 'Estático' not in hojas['modal']
    assert 'Estático' in hojas['estatico'] and 'Modal espectral' not in hojas['estatico']
    assert {'Modal espectral', 'Estático'} <= hojas['ambos']
    pymupdf = pytest.importorskip('pymupdf')
    textos = {}
    for m in ('modal', 'estatico'):
        buf = io.BytesIO()
        ame_pdf.memoria_pdf(r, buf, proyecto='P', autor='A', metodo=m)
        textos[m] = ''.join(pg.get_text() for pg in pymupdf.open(stream=buf.getvalue(), filetype='pdf'))
    assert 'ANÁLISIS MODAL ESPECTRAL' in textos['modal'] and 'ANÁLISIS ESTÁTICO' not in textos['modal']
    assert 'ANÁLISIS ESTÁTICO' in textos['estatico'] and 'ANÁLISIS MODAL ESPECTRAL' not in textos['estatico']
    assert 'Coeficiente sísmico C' in textos['estatico'] and 'Valores propios' in textos['modal']



def test_ejemplo_15_pizarron_columnas_que_faltan_y_eje_de_13_m():
    p = ame.preparar(copy.deepcopy(EJ15))
    assert p['cargas'] == pytest.approx([24.0, 57.0, 58.0])             # incluye el voladizo del techo (2 T/m × 5 m)
    I50, I40 = 50 * 30 ** 3 / 12, 40 * 10 ** 3 / 12
    E = 150_000.0
    k = lambda I, h, f: f * E * I / (h * 100) ** 3
    assert p['info'][1]['k'] == pytest.approx(2 * k(I50, 5, 12))      # solo los ejes 1 y 4 (el 5 pasa de largo)
    assert p['info'][2]['k'] == pytest.approx(2 * k(I50, 4, 12) + 2 * k(I40, 4, 12))
    assert p['K'][0, 1] == pytest.approx(-p['info'][1]['k'])
    col5 = [t for t in ame.tabla_columnas(p) if t['eje'] == 5][0]
    assert col5['L'] == 13.0 and col5['hasta'] == 2 and col5['k'] == pytest.approx(k(I50, 13, 3))
    r = ame.calcular(p)
    assert r['clave'] == 'DL' and r['p']['Ks'] == pytest.approx(1 / 6)
    assert r['Sa_modal'][:, 0].max() <= 400 / 6 + 1e-9                  # nunca pasa de c·Ks


def test_pagina_abre_un_proyecto_json_desde_el_cargador_sin_error():
    import json
    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file(str(Path(ame.__file__).parent / 'app_web.py'), default_timeout=180).run()
    ejemplo = list(ame.EJEMPLOS_MODELO.values())[13]
    at.sidebar.file_uploader[0].upload('proyecto.json', json.dumps(ejemplo, ensure_ascii=False).encode('utf-8')).run()
    assert not at.exception
    assert at.text_input(key='crujias').value == '4 4 2'
    assert at.radio(key='captura').value == at.radio(key='captura').options[1]          # el ejemplo 14 es por columna
    at.sidebar.file_uploader[0].upload('malo.json', b'esto no es json').run()
    assert not at.exception and any('No se pudo leer' in e.value for e in at.error)


def test_estructura_A_de_la_hoja_de_excel_de_un_companero_coincide_en_las_dos_versiones():
    """Hoja 19 (K y M dadas): ordenadas de la hoja con SASID y, en la segunda versión, ordenadas dadas a mano."""
    r = correr(copy.deepcopy(EJ6))
    assert r['T'] == pytest.approx([2.17897, 0.99388, 0.57350, 0.50195], abs=2e-5)
    assert r['Sa_modal'][:, 0] == pytest.approx([210.02, 476.19, 476.19, 476.19], abs=0.01)
    assert r['gam'] == pytest.approx([0.413132716, 0.33423369, 0.17398604, 0.078647486], abs=2e-6)
    assert r['d'][:, 0] == pytest.approx([10.4350, 19.1347, 30.1699, 36.1878], abs=1e-3)
    assert r['F'][:, 0] == pytest.approx([260296.08, 477314.09, 376289.14, 451349.22], rel=1e-4)
    assert r['F_final'] == pytest.approx([608390.70, 763306.42, 426319.26, 601005.68], rel=1e-5)
    t = r['tabla_din']
    assert t['Fu'] == pytest.approx([661.11, 468.95, 839.64, 669.23], abs=0.01)                 # niveles 4, 3, 2, 1
    assert t['Vu'] == pytest.approx([661.11, 1130.06, 1969.69, 2638.93], abs=0.01)   # acumulado (la hoja de Excel suma mal el Vu)
    assert t['Mvu'].sum() == pytest.approx(31998.9123, rel=1e-6)
    d2 = copy.deepcopy(EJ6)
    d2.update(Sa_modos='41.44 87.50 87.50 87.50', unidad_espectro='cm/s²')
    r2 = correr(d2)
    assert r2['d'][:, 0] == pytest.approx([2.0590, 3.7756, 5.9530, 7.1404], abs=1e-3)
    assert r2['F'][:, 0] == pytest.approx([51360.34, 94181.26, 74247.52, 89058.00], rel=1e-4)
    assert r2['F_final'] == pytest.approx([113347.42, 144394.84, 82876.27, 115103.81], rel=1e-5)
    t2 = r2['tabla_din']
    assert t2['Fu'] == pytest.approx([126.61, 91.16, 158.83, 124.68], abs=0.01)
    assert t2['Vu'] == pytest.approx([126.61, 217.78, 376.61, 501.29], abs=0.01)
    assert t2['Mvu'].sum() == pytest.approx(6111.4964, rel=1e-5)


def test_detalle_explica_el_peso_obtenido_del_periodo_medido():
    """Problema al revés (nave con T medido): el detalle muestra k, T, m = kT²/4π² y W = m·g."""
    import ame_detalle
    r = correr(dict(EJ9, intensidad='Base de diseño', Q='2', k1='1.0'))
    modal = ame_detalle.detalle_modal(r)
    pasos = [b['texto'] for b in modal if b['t'] == 'paso']
    assert pasos[1:3] == ['Matriz de rigidez K', 'Peso del marco a partir del periodo medido']
    assert [b['n'] for b in modal if b['t'] == 'paso'] == list(range(1, len(pasos) + 1))
    k = r['p']['K'][0, 0]
    m = k * 0.17 ** 2 / (4 * np.pi ** 2)
    assert m == pytest.approx(r['p']['M'][0, 0]) and m * 981 / 1000 == pytest.approx(1.7932, abs=1e-4)
    txt = ame_detalle.bloques_a_texto(modal)
    assert 'm = k·T² / (4π²)' in txt and '1.79321 t' in txt
    # sin periodo medido, el paso de masas sigue siendo el de siempre
    r12 = correr(copy.deepcopy(EJ12_SOL))
    assert 'Matriz de masas M' in [b['texto'] for b in ame_detalle.detalle_modal(r12) if b['t'] == 'paso']


def test_ejemplo_16_torre_de_acero_coincide_con_la_hoja_de_un_companero():
    """Examen abr-2025: K, periodos, ordenadas y fuerzas de la hoja de Excel; los errores típicos de captura se detectan."""
    KC = np.array([[4061.0500, -3625.0667, 0, 0, 0], [-3625.0667, 10819.7531, -7067.2, 0, 0],
                   [0, -7067.2, 8867.4801, -1529.3250, -270.9551], [0, 0, -1529.3250, 5154.3917, -3625.0667],
                   [0, 0, -270.9551, -3625.0667, 3896.0218]])
    r = correr(EJ16)
    assert r['p']['K'] == pytest.approx(KC, abs=1e-3)
    assert np.diag(r['p']['M']) == pytest.approx([18.35, 40.77, 26.50, 18.35, 40.77], abs=0.005)
    assert r['T'] == pytest.approx([3.431511, 0.956727, 0.400517, 0.3343, 0.267441], abs=1e-4)
    assert r['Sa_modal'][:, 0] == pytest.approx([38.04, 187.56, 174.50, 164.94, 154.30], abs=0.01)
    t = r['tabla_din']
    assert t['Fu'] == pytest.approx([2.13, 0.88, 1.14, 1.75, 0.78], abs=0.01)           # niveles 5 a 1
    assert t['Vu'][-1] == pytest.approx(6.68, abs=0.01) and t['Mvu'].sum() == pytest.approx(88.64, abs=0.02)
    qr = [f for f in r['normativa']['factores'] if f[0].startswith('Q·R′ (distorsiones)')][0]
    assert float(qr[1]) == pytest.approx(2.8)
    # E tecleado 10 veces mayor y bases articuladas capturadas como empotradas: la K de la memoria equivocada
    mal = copy.deepcopy(EJ16)
    mal['E'] = '21000000'
    for f in mal['columnas']:
        f['apoyo'] = 'Empotrado'
    assert correr(mal)['p']['K'][0, 0] == pytest.approx(40782.0, abs=0.1)
    assert any('poco usual para acero' in tx for _, tx in correr(mal)['avisos'])
