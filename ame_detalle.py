"""Procedimiento paso a paso del análisis modal espectral y del análisis estático.

Cada análisis se arma como una lista de bloques (dicts con la clave 't'): título, paso con su explicación de dos líneas,
fórmula, matriz, tabla, texto y nota. La página, el Excel y el PDF recorren la misma lista, así que lo que ves en pantalla es
lo mismo que se exporta.
"""
import numpy as np

import ame_sismico as ame

G = ame.G


# ----------------------------------------------------------------------------
# Constructores de bloques
# ----------------------------------------------------------------------------
def titulo(texto, sub=''):
    return dict(t='titulo', texto=texto, sub=sub)


def paso(n, texto, linea1, linea2):
    return dict(t='paso', n=n, texto=texto, expl=[linea1, linea2])


def formula(latex, texto=None):
    return dict(t='formula', latex=latex, texto=texto or latex)


def matriz(nombre, datos, dec=4, unidad='', filas=None, cols=None, simbolica=None, latex_nombre=None):
    """Matriz numérica (o simbólica si se da 'simbolica': lista de listas de cadenas LaTeX)."""
    datos = np.asarray(datos, float)
    return dict(t='matriz', nombre=nombre, latex_nombre=latex_nombre or nombre, datos=datos, dec=dec, unidad=unidad,
                filas=filas, cols=cols, simbolica=simbolica)


def tabla(titulo_, enc, filas, dec=4, anchos=None):
    return dict(t='tabla', titulo=titulo_, enc=list(enc), filas=[list(f) for f in filas], dec=dec)


def texto(txt):
    return dict(t='texto', texto=txt)


def destacado(titulo_, items):
    """Cifras grandes y evidentes: items = [(etiqueta, valor, detalle)] con valor y detalle ya en texto."""
    return dict(t='destacado', titulo=titulo_, items=[list(i) for i in items])


def seccion(letra, texto_):
    return dict(t='seccion', letra=letra, texto=texto_)


def sub(texto_):
    return dict(t='sub', texto=texto_)


def figura(titulo_, fig, ancho):
    return dict(t='figura', titulo=titulo_, fig=fig, ancho=ancho)


def nota(txt):
    return dict(t='nota', texto=nota_limpia(txt))


def nota_limpia(txt):
    return str(txt).strip()


def _g(x, dec=5):
    """Número con cifras significativas razonables (sin ceros de relleno)."""
    x = float(x)
    if x == 0:
        return '0'
    t = f'{x:.{dec}g}'
    return f'{x:.2f}' if 'e' in t else t


def _niveles(n):
    return [f'N{i}' for i in range(1, n + 1)]


# ----------------------------------------------------------------------------
# Matriz de rigidez: columna por columna
# ----------------------------------------------------------------------------
def _columnas_modelo(p):
    """Lista de columnas con k propio (las continuaciones de una columna que cruza niveles no se repiten)."""
    cols = []
    for l in p['geom']['niveles']:
        for c in l['cols']:
            if c.get('k_cont'):
                continue
            cols.append(c)
    return cols


def _tipo_y_expresion(c, alturas):
    """Tipo de columna y la expresión numérica de su rigidez lateral."""
    i0, i1 = c['nivel'] - 1, c.get('hasta', c['nivel']) - 1
    L = float(sum(alturas[i0:i1 + 1])) * 100.0
    E, I = float(c['E']), float(c['I'])
    if c.get('rodillo'):
        return 'Rodillo en la base', 'no transmite cortante', 0.0, L
    if c.get('artic'):
        coef, tipo = 3, 'Articulada en un extremo (3EI/L³)'
    elif c.get('artic_sup'):
        coef, tipo = 3, 'Articulada arriba (3EI/L³)'
    else:
        coef, tipo = 12, 'Empotrada (12EI/L³)'
    k = coef * E * I / L ** 3
    exp = f'{coef}·{_g(E, 7)}·{_g(I, 7)} / {_g(L, 5)}³'
    if abs(k - c['k']) > 1e-6 * max(abs(k), 1):
        return tipo, 'según apoyos del modelo', float(c['k']), L
    return tipo + (' · cruza niveles' if i1 > i0 else ''), exp, float(c['k']), L


def _k_simbolica(p):
    """Matriz K con símbolos por nivel (k_i) y por columna que cruza (k_c). None si K no se arma así (vigas flexibles)."""
    n = p['n']
    cols = _columnas_modelo(p)
    por_nivel = {}
    cruzan = []
    for c in cols:
        if c.get('hasta', c['nivel']) > c['nivel']:
            cruzan.append(c)
        else:
            por_nivel[c['nivel']] = por_nivel.get(c['nivel'], 0.0) + c['k']
    term = {}                       # (i, j) -> lista de (signo, símbolo, valor)

    def suma(i, j, signo, simbolo, valor):
        term.setdefault((i, j), []).append((signo, simbolo, valor))

    for nivel, k in por_nivel.items():
        hi, lo = nivel - 1, nivel - 2
        s = f'k_{{{nivel}}}'
        suma(hi, hi, 1, s, k)
        if lo >= 0:
            suma(lo, lo, 1, s, k)
            suma(lo, hi, -1, s, k)
            suma(hi, lo, -1, s, k)
    for c in cruzan:
        hi, lo = c['hasta'] - 1, c['nivel'] - 2
        s = f'k_{{c{c["eje"]}}}'
        suma(hi, hi, 1, s, c['k'])
        if lo >= 0:
            suma(lo, lo, 1, s, c['k'])
            suma(lo, hi, -1, s, c['k'])
            suma(hi, lo, -1, s, c['k'])
    Kr = np.zeros((n, n))
    for (i, j), lst in term.items():
        Kr[i, j] = sum(sg * v for sg, _, v in lst)
    if not np.allclose(Kr, p['K'], rtol=1e-7, atol=1e-6):
        return None, por_nivel, cruzan
    sim = [['0'] * n for _ in range(n)]
    for (i, j), lst in term.items():
        txt = ''
        for sg, s, _ in lst:
            txt += ('-' if sg < 0 else ('+' if txt else '')) + s
        sim[i][j] = txt or '0'
    return sim, por_nivel, cruzan


# ----------------------------------------------------------------------------
# Análisis modal espectral
# ----------------------------------------------------------------------------
def items_aceleraciones_modal(res):
    """(etiqueta, 'x g', 'y cm/s²') de la ordenada espectral de diseño de cada modo."""
    T = res['T']
    Sa = np.asarray(res['Sa_modal'][:, 0], float)
    nombre = 'Sa·Ks' if res['clave'] == 'DL' else 'Sa de diseño'
    return [(f'Modo {j + 1} · T = {float(T[j]):.4f} s ({nombre})', f'{Sa[j] / G:.5f} g', f'{Sa[j]:.3f} cm/s²')
            for j in range(len(T))]


def items_aceleraciones_estatico(res):
    """Coeficiente sísmico C del método estático, en g y en cm/s²."""
    se = res['estatico']
    items = [('Coeficiente sísmico C (gobierna: ' + str(se['gobierna']) + ')', f'{se["C"]:.5f} g', f'{se["C"] * G:.3f} cm/s²')]
    if 'C_pl' in se and 'C_0' in se:
        items += [('En el plato del espectro', f'{se["C_pl"]:.5f} g', f'{se["C_pl"] * G:.3f} cm/s²'),
                  ('En T = 0', f'{se["C_0"]:.5f} g', f'{se["C_0"] * G:.3f} cm/s²')]
    return items


def bloques_aceleraciones(res):
    """Las aceleraciones de todos los análisis, para el resumen."""
    B = [destacado('Aceleraciones espectrales de diseño · modal espectral', items_aceleraciones_modal(res))]
    if 'tabla_est' in res:
        B.append(destacado('Aceleración del método estático', items_aceleraciones_estatico(res)))
    return B


def _bloques_masas(p, M, nv):
    n = p['n']
    B = []
    if p['cargas'] is not None:
        filas = [[i + 1, float(p['cargas'][i]), f'{_g(p["cargas"][i])} t · 1000 / 981', float(M[i, i])] for i in range(n)]
        B.append(tabla('Masa de cada nivel', ['Nivel', 'W (t)', 'm = W·1000 / g', 'm (kg·s²/cm)'], filas, 5))
    else:
        B.append(nota('Las masas vienen dadas directamente en la matriz M del problema.'))

    B.append(matriz('M', M, 4, 'kg·s²/cm', nv, nv, latex_nombre=r'\mathbf{M}'))
    return B


def _bloques_rigidez(p, K, nv):
    n = p['n']
    B = []
    if 'geom' in p:
        filas = []
        for c in _columnas_modelo(p):
            tipo, exp, k, L = _tipo_y_expresion(c, p['alturas'])
            nivel_txt = str(c['nivel']) + (f'–{c["hasta"]}' if c.get('hasta', c['nivel']) > c['nivel'] else '')
            filas.append([nivel_txt, c['eje'], c['texto'], float(c['I']), float(c['E']), L / 100.0, tipo, exp, k])
        B.append(tabla('Rigidez de cada columna  (k en kg/cm, I en cm⁴, E en kg/cm², L en m)',
                       ['Nivel', 'Eje', 'Sección', 'I', 'E', 'L (m)', 'Tipo', 'Cálculo', 'k'], filas, 3))
        sim, por_nivel, cruzan = _k_simbolica(p)
        if sim is not None:
            lineas = []
            for nivel in sorted(por_nivel):
                cs = [c for c in _columnas_modelo(p) if c['nivel'] == nivel and c.get('hasta', nivel) == nivel]
                lineas.append(f'k_{nivel} = ' + ' + '.join(_g(c['k']) for c in cs) + f' = {_g(por_nivel[nivel])} kg/cm')
            for c in cruzan:
                lineas.append(f'k_c{c["eje"]} = {_g(c["k"])} kg/cm  (columna del eje {c["eje"]} que cruza del nivel {c["nivel"]} al {c["hasta"]})')
            B.append(tabla('Rigidez de cada nivel (suma de sus columnas)', ['Rigidez del nivel'], [[l] for l in lineas]))
            B.append(texto('Planteamiento de K con las rigideces de nivel:'))
            B.append(matriz('K (simbólica)', np.zeros((n, n)), 0, '', nv, nv, simbolica=sim, latex_nombre=r'\mathbf{K}'))
        else:
            B.append(nota('Este modelo no es un marco de cortante simple (vigas flexibles o apoyos especiales): K se obtuvo '
                          'por condensación estática del marco completo, no sumando k por nivel.'))
        B.append(texto('Matriz K con números:'))
    else:
        B.append(nota('La matriz de rigidez viene dada directamente en el problema.'))
    B.append(matriz('K', K, 4, 'kg/cm', nv, nv, latex_nombre=r'\mathbf{K}'))

    return B


def detalle_modal(res):
    p, n = res['p'], res['p']['n']
    clave = res['clave']
    dl = clave == 'DL'
    nv = _niveles(n)
    B = [titulo('ANÁLISIS MODAL ESPECTRAL',
                f'{p["estado"]} · combinación {p["combinacion"]}' + (f' (ζ = {p["zeta"]:g})' if p['combinacion'] == 'CQC' else ''))]

    # ---- 1 datos
    filas = [['Estado límite de diseño', p['estado']], ['Combinación modal', p['combinacion']],
             ['E (kg/cm²)', _g(p['E'], 8) if 'E' in p and p['E'] else '—'],
             ['a0 · c (cm/s²)', f'{_g(p["a0"])} · {_g(p["c"])}'], ['Ta · Tb (s)', f'{_g(p["Ta"])} · {_g(p["Tb"])}'],
             ['k · Ts', f'{_g(p["k"])} · {_g(p["Ts"])}'],
             ['Q (nominal)', ', '.join(_g(q) for q in p['Q'])], ['k1', _g(p['k1'])],
             ['Factor de carga FC', _g(p['factor_Fu']) if not dl else '1 (limitación de daño: sin factor de carga)']]
    if dl:
        filas.append(['Ks', f'{_g(p["Ks"])}' + (f' (ec. 3.1.1, zona {p["Ks_zona"]}, Ts = {_g(p["Ks_Ts"])} s)' if p['Ks_auto'] else ' (dado)')])
    B += [paso(1, 'Datos del problema',
               'Voy a reunir los datos que gobiernan el cálculo: material, espectro SASID, Q, k1 y el factor de carga.',
               'Todo lo que sigue se obtiene de estos valores; si alguno está mal, se arrastra a todos los pasos.'),
          tabla('', ['Dato', 'Valor'], filas)]

    # ---- 2 masas
    B_base, B = B, []
    M = np.asarray(p['M'], float)
    B.append(paso(2, 'Matriz de masas M',
                  'Voy a convertir el peso de cada nivel en masa, m = W / g, con g = 981 cm/s², y a colocarla en la diagonal de M.',
                  'M es diagonal porque la masa de cada piso se concentra en su nivel (modelo de cortante).'))
    B += _bloques_masas(p, M, nv)

    B_masa, B = B, []
    # ---- 3 rigidez
    K = np.asarray(p['K'], float)
    B.append(paso(3, 'Matriz de rigidez K',
                  'Voy a calcular la rigidez lateral de cada columna (12EI/L³ si está empotrada, 3EI/L³ si tiene una articulación, 0 con rodillo).',
                  'Después sumo las columnas de cada nivel (k_i) y ensamblo K: en la diagonal k_i + k_(i+1) y fuera de ella −k_(i+1).'))
    B += _bloques_rigidez(p, K, nv)

    B_rig = B
    B = B_base
    aj = p.get('ajuste_periodo')
    if aj:      # el peso no es un dato: sale del periodo medido (problema "al revés")
        B += B_rig + _bloques_ajuste(p, aj, M, K) + []
    else:
        B += B_masa + B_rig

    # ---- 4 valores propios
    lam, w, f, T, fi = res['lam'], res['w'], res['f'], res['T'], np.asarray(res['fi'], float)
    Minv = np.linalg.inv(M)
    A = Minv @ K
    B.append(paso(4, 'Periodos y formas modales',
                  'Voy a resolver el problema de valores propios de A = M⁻¹K: sus valores propios son λ = ω² y sus vectores las formas modales φ.',
                  'De λ obtengo ω = √λ, la frecuencia f = ω/2π y el periodo T = 1/f; ordeno los modos del más flexible (T mayor) al más rígido.'))
    B += [formula(r'\mathbf{A}=\mathbf{M}^{-1}\mathbf{K}\qquad(\mathbf{A}-\lambda\mathbf{I})\,\boldsymbol{\varphi}=\mathbf{0}',
                  'A = M⁻¹ K ;  (A − λ I) φ = 0'),
          matriz('M⁻¹', Minv, 6, '(kg·s²/cm)⁻¹', nv, nv, latex_nombre=r'\mathbf{M}^{-1}'),
          matriz('A = M⁻¹K', A, 4, '1/s²', nv, nv, latex_nombre=r'\mathbf{A}=\mathbf{M}^{-1}\mathbf{K}'),
          tabla('Valores propios, frecuencias y periodos',
                ['Modo', 'λ = ω² (rad/s)²', 'ω (rad/s)', 'f (Hz)', 'T (s)'],
                [[j + 1, float(lam[j]), float(w[j]), float(f[j]), float(T[j])] for j in range(n)], 5),
          matriz('Φ (formas modales, cada columna es un modo)', fi, 5, '', nv, [f'Modo {j + 1}' for j in range(n)],
                 latex_nombre=r'\boldsymbol{\Phi}')]
    if p.get('norm_modal') and p['norm_modal'] != 'primero':
        B.append(nota(f'Normalización de las formas modales: {p["norm_modal"]}.'))

    # ---- 5 participación
    one = np.ones(n)
    Lj = np.array([fi[:, j] @ M @ one for j in range(n)])
    Mj = np.array([fi[:, j] @ M @ fi[:, j] for j in range(n)])
    gam = res['gam']
    mo = res['normativa']['modos']
    B += [paso(5, 'Factores de participación modal',
               'Voy a medir cuánto del movimiento del suelo excita cada modo con γ = (φᵀ M 1) / (φᵀ M φ), donde 1 es el vector de unos.',
               'Los γ de todos los modos suman 1; con ellos y la masa efectiva reviso que los modos usados representen casi toda la masa.'),
          formula(r'\gamma_j=\dfrac{\boldsymbol{\varphi}_j^{T}\mathbf{M}\,\mathbf{1}}{\boldsymbol{\varphi}_j^{T}\mathbf{M}\,\boldsymbol{\varphi}_j}',
                  'γ_j = (φ_jᵀ M 1) / (φ_jᵀ M φ_j)'),
          tabla('', ['Modo', 'φᵀ M 1', 'φᵀ M φ', 'γ', 'Masa efectiva', 'Acumulada'],
                [[j + 1, float(Lj[j]), float(Mj[j]), float(gam[j]), float(mo['m_ef'][j]), float(mo['acum'][j])] for j in range(n)], 6),
          texto(f'Σγ = {float(np.sum(gam)):.6f}   (debe ser 1)')]

    # ---- 6 espectro
    clave_q = ame.CLAVE_ESTADO[p['estado']]
    pq = p if clave_q == 'SV' else dict(p, Q=np.array([1.0]))
    Q0 = float(pq['Q'][0])
    Sa_el = ame._interp(T, res['serie'], ame.espectro(dict(p, Ks=1.0), 1.0, 1.0, res['serie'], dl=True)[0])
    f_sa = formula(
        r'S_a(T)=\begin{cases}a_0+(c-a_0)\dfrac{T}{T_a} & T<T_a\\ c & T_a\leq T<T_b\\ c\,p\left(\dfrac{T_b}{T}\right)^{2} & T\geq T_b\end{cases}'
        r'\qquad p=k+(1-k)\left(\dfrac{T_b}{T}\right)^{2}',
        'Sa(T) = a0 + (c − a0)·T/Tₐ (T < Tₐ) · c (Tₐ ≤ T < T_b) · c·p·(T_b/T)² (T ≥ T_b),  con p = k + (1 − k)(T_b/T)²')
    Sa_d = np.asarray(res['Sa_modal'][:, 0], float)
    if dl:
        B += [paso(6, 'Ordenadas espectrales de cada modo',
                   'Voy a leer en el espectro de base de diseño la ordenada Sa de cada periodo y a multiplicarla por Ks.',
                   'En limitación de daño no se divide entre Q′ ni R′ y no hay factor de carga: la ordenada de cálculo es Sa(BD)·Ks.'),
              f_sa,
              formula(r'S_a^{\,diseño}(T_j)=S_a(T_j)\cdot K_s', 'Sa_diseño(T_j) = Sa(T_j) · Ks'),
              tabla('', ['Modo', 'T (s)', 'Sa(T) elástica (cm/s²)', 'Ks', 'Sa·Ks (cm/s²)', 'Sa·Ks (g)'],
                    [[j + 1, float(T[j]), float(Sa_el[j]), float(p['Ks']), float(Sa_d[j]), float(Sa_d[j] / G)] for j in range(n)], 5),
              destacado('Aceleraciones espectrales de diseño', items_aceleraciones_modal(res))]
    else:
        q_, r_ = ame._q_r(pq, Q0, T, clave_q)
        B += [paso(6, 'Ordenadas espectrales de cada modo',
                   'Voy a leer Sa en el espectro para cada T, calcular Q′ y R′ de ese periodo (R = k1·R0 + k2, R′ = R o 0.75R) y reducir: Sa/(Q′R′).',
                   'Ese cociente es la ordenada de diseño de cada modo; si queda por debajo de la ordenada mínima a_min se sustituye por ella.'),
              f_sa,
              formula(r"Q^{\prime}=\begin{cases}1+(Q-1)\sqrt{1/k}\,\dfrac{T}{T_a} & T<T_a\\ 1+(Q-1)\sqrt{1/k} & T_a\leq T<T_b\\ 1+(Q-1)\sqrt{p/k} & T\geq T_b\end{cases}",
                      "Q′ = 1 + (Q−1)·√(1/k)·T/Tₐ (T < Tₐ) · 1 + (Q−1)·√(1/k) (Tₐ ≤ T < T_b) · 1 + (Q−1)·√(p/k) (T ≥ T_b)"),
              formula(r"R=k_1R_0+k_2\quad k_2=\max\left\{0,\ 0.5\left(1-\sqrt{T/T_a}\right)\right\}\quad R^{\prime}=R\ \text{ó}\ 0.75R",
                      "R = k1·R0 + k2 ,  k2 = máx{0, 0.5(1 − √(T/Tₐ))} ;  R′ = R (seguridad de vida) ó 0.75·R (ocupación inmediata)"),
              formula(r'a_j=\dfrac{S_a(T_j)}{Q^{\prime}(T_j)\,R^{\prime}(T_j)}', "a_j = Sa(T_j) / (Q'(T_j) · R'(T_j))"),
              tabla('', ['Modo', 'T (s)', 'Sa(T) elástica (cm/s²)', "Q'", "R'", "Q'·R'", "Sa/(Q'R') (cm/s²)", 'en g'],
                    [[j + 1, float(T[j]), float(Sa_el[j]), float(q_[j]), float(r_[j]), float(q_[j] * r_[j]), float(Sa_d[j]),
                      float(Sa_d[j] / G)] for j in range(n)], 5),
              destacado('Aceleraciones espectrales de diseño', items_aceleraciones_modal(res))]
        if res.get('a_min'):
            B.append(nota(f'Ordenada mínima a_min = {res["a_min"]:.4f} cm/s² ({res["a_min_g"]:.5f} g).'))
    if p.get('Sa_manual') is not None:
        B.append(nota('Las ordenadas por modo se tomaron del enunciado (campo «Ordenadas Sa por modo»), no del espectro.'))

    # ---- 7 desplazamientos modales
    d, F = np.asarray(res['d'], float), np.asarray(res['F'], float)
    modos = [f'Modo {j + 1}' for j in range(n)]
    B += [paso(7, 'Desplazamientos modales',
               'Voy a calcular el desplazamiento de cada nivel en cada modo: δ_j = φ_j · γ_j · a_j / ω_j².',
               'Es la respuesta máxima de cada modo por separado; todavía no se combinan.'),
          formula(r'\boldsymbol{\delta}_j=\boldsymbol{\varphi}_j\,\gamma_j\,\dfrac{a_j}{\omega_j^{2}}', 'δ_j = φ_j · γ_j · a_j / ω_j²'),
          matriz('δ (cada columna es un modo)', d, 6, 'cm', nv, modos, latex_nombre=r'\boldsymbol{\delta}')]

    # ---- 8 fuerzas modales
    B += [paso(8, 'Fuerzas modales',
               'Voy a convertir los desplazamientos en fuerzas laterales con F_j = K · δ_j (equivale a M·φ_j·γ_j·a_j).',
               'Cada columna es el juego de fuerzas que actuaría en los pisos si solo vibrara ese modo.'),
          formula(r'\mathbf{F}_j=\mathbf{K}\,\boldsymbol{\delta}_j', 'F_j = K δ_j'),
          matriz('F (cada columna es un modo)', F, 3, 'kg', nv, modos, latex_nombre=r'\mathbf{F}')]

    # ---- 9 combinación
    Ff = np.asarray(res['F_final'], float)
    m_comb = res['n_modos']
    comb = p['combinacion']
    B.append(paso(9, 'Combinación modal',
                  f'Voy a combinar las fuerzas de los {m_comb} modos usados por nivel con {comb}, porque los máximos de cada modo no ocurren a la vez.',
                  'Después acumulo desde arriba para obtener el cortante en cada entrepiso.'))
    if comb.startswith('SRSS'):
        B.append(formula(r'F_i=\sqrt{\sum_{j}F_{ij}^{2}}', 'F_i = √( Σ_j F_ij² )'))
        lineas = []
        for i in range(n - 1, -1, -1):
            lineas.append(f'F_{i + 1} = √(' + ' + '.join((f'({_g(F[i, j], 5)})²' if F[i, j] < 0 else f'{_g(F[i, j], 5)}²')
                                                         for j in range(m_comb)) + f') = {_g(Ff[i], 6)} kg')
        B.append(tabla('', ['Combinación por nivel (de arriba hacia abajo)'], [[l] for l in lineas]))
    elif comb.startswith('CQC'):
        B.append(formula(r'F_i=\sqrt{\sum_{j}\sum_{k}\rho_{jk}F_{ij}F_{ik}}', 'F_i = √( Σ_j Σ_k ρ_jk F_ij F_ik )'))
    else:
        B.append(formula(r'F_i=\sum_{j}\left|F_{ij}\right|', 'F_i = Σ_j |F_ij|'))
    V = np.cumsum(Ff[::-1])[::-1]
    B.append(tabla('', ['Nivel', 'F combinada (kg)', 'F combinada (t)', 'Cortante V (kg)'],
                   [[i + 1, float(Ff[i]), float(Ff[i] / 1000), float(V[i])] for i in range(n - 1, -1, -1)], 4))

    # ---- 10 fuerzas de diseño
    t = res['tabla_din']
    FC = 1.0 if dl else p['factor_Fu']
    B += [paso(10, 'Fuerzas sísmicas de diseño',
               f'Voy a multiplicar cada fuerza por el factor de carga FC = {FC:g}, acumular los cortantes Vu desde arriba y calcular Mvu = Fu · H.',
               'La suma de los Mvu es el momento de volteo de diseño en la base.'),
          formula(r'F_u=FC\cdot F\qquad V_{u,i}=\sum_{k\geq i}F_{u,k}\qquad M_{vu}=F_u\cdot H_i', 'Fu = FC·F ;  Vu_i = Σ_(k≥i) Fu_k ;  Mvu = Fu·H_i'),
          tabla('', ['Nivel', 'H (m)', 'F (t)', 'Fu (t)', 'Vu (t)', 'Mvu (t·m)'],
                [[int(t['nivel'][i]), float(t['h'][i]), float(t['F'][i]), float(t['Fu'][i]), float(t['Vu'][i]), float(t['Mvu'][i])]
                 for i in range(n)], 4),
          texto(f'Momento de volteo de diseño  Σ Mvu = {float(t["Mvu"].sum()):.4f} t·m = {float(t["Mvu"].sum()) * 1e5:,.1f} kg·cm'),
          tabla('Los mismos resultados en kg y kg·cm', ['Nivel', 'F (kg)', 'Fu (kg)', 'Vu (kg)', 'Mvu (kg·cm)'],
                [[int(t['nivel'][i]), float(t['F'][i] * 1000), float(t['Fu'][i] * 1000), float(t['Vu'][i] * 1000),
                  float(t['Mvu'][i] * 1e5)] for i in range(n)], 1)]
    n_paso = 0
    for b in B:
        if b['t'] == 'paso':
            n_paso += 1
            b['n'] = n_paso
    return B


def _bloques_ajuste(p, aj, M, K):
    """Paso que obtiene la masa (y el peso) a partir del periodo medido T1."""
    n = p['n']
    nv = _niveles(n)
    g = ame.G
    Tm = float(aj['T_medido'])
    if n == 1:
        k = float(K[0, 0])
        m = k * Tm ** 2 / (4 * np.pi ** 2)
        B = [paso(3, 'Peso del marco a partir del periodo medido',
                  f'Voy a obtener la masa del marco con el periodo medido T = {Tm:g} s y la rigidez lateral total k del paso anterior.',
                  'Como T = 2π√(m/k), la masa debe ser m = k·T²/(4π²); el peso es W = m·g y con él se arma la matriz M.'),
             formula(r'T=2\pi\sqrt{m/k}\ \Rightarrow\ m=\dfrac{k\,T^{2}}{4\pi^{2}}\qquad W=m\,g',
                     'T = 2π√(m/k)  →  m = k·T² / (4π²) ;  W = m·g'),
             tabla('', ['Concepto', 'Cálculo', 'Resultado'],
                   [['Rigidez lateral total k', 'Σ de las columnas (paso anterior)', f'{k:.4f} kg/cm'],
                    ['Periodo medido T', 'dato', f'{Tm:g} s'],
                    ['Masa m = k·T² / (4π²)', f'{k:.4f} · {Tm:g}² / (4π²)', f'{m:.5f} kg·s²/cm'],
                    ['Peso W = m·g', f'{m:.5f} · {g:g} / 1000', f'{m * g / 1000:.5f} t']])]
    else:
        B = [paso(3, 'Peso del marco a partir del periodo medido',
                  f'Voy a escalar las masas para que el primer periodo calculado coincida con el medido T₁ = {Tm:g} s.',
                  'Los pesos dados solo fijan cómo se reparte la masa entre niveles; el factor de escala es (T medido / T calculado)².'),
             formula(r'T\propto\sqrt{m}\ \Rightarrow\ f=\left(\dfrac{T_{medido}}{T_{calculado}}\right)^{2}\qquad W_{total}=f\cdot W_{prov}',
                     'T ∝ √m  →  f = (T_medido / T_calculado)² ;  W_total = f · W_provisional'),
             tabla('', ['Concepto', 'Resultado'],
                   [['T₁ calculado con los pesos relativos', f'{float(aj["T_calculado"]):.5f} s'],
                    ['T₁ medido', f'{Tm:g} s'],
                    ['Factor f = (T medido / T calculado)²', f'{float(aj["factor"]):.6f}'],
                    ['Masa total', f'{float(aj["m_total"]):.5f} kg·s²/cm'],
                    ['Peso total W', f'{float(aj["W_total"]):.5f} t']])]
    filas = [[i + 1, float(p['cargas'][i]), f'{p["cargas"][i]:.4f} t · 1000 / {g:g}', float(M[i, i])] for i in range(n)]
    B += [tabla('Pesos y masas resultantes', ['Nivel', 'W (t)', 'm = W·1000 / g', 'm (kg·s²/cm)'], filas, 5),
          matriz('M', M, 4, 'kg·s²/cm', nv, nv, latex_nombre=r'\mathbf{M}')]
    return B


# ----------------------------------------------------------------------------
# Análisis estático
# ----------------------------------------------------------------------------
def detalle_estatico(res):
    p, n = res['p'], res['p']['n']
    B = [titulo('ANÁLISIS ESTÁTICO', f'{p["estado"]}')]
    if 'tabla_est' not in res:
        B.append(nota('Para el método estático hacen falta los pesos y las alturas de los niveles; captura las cargas en la pestaña Estructura.'))
        return B
    e, se = res['tabla_est'], res['estatico']
    nm = res['normativa']['estatico']
    dl = res['clave'] == 'DL'
    FC = se['FC']

    # ---- 1 aplicabilidad
    B += [paso(1, 'Aplicabilidad del método',
               'Voy a revisar si el método estático puede usarse: depende de la altura total, de la zona y de la regularidad.',
               'Si no aplica, el análisis se calcula igual pero solo sirve de referencia (y para el factor de escala).'),
          tabla('', ['Concepto', 'Valor'],
                [['Altura total H (m)', f'{nm["H"]:.2f}'],
                 ['Límite de altura (m)', f'{nm["limite"]:g}' if nm['limite'] else 'no definido (falta la zona geotécnica)'],
                 ['¿Aplica?', 'Sí' if nm['aplica'] else (
                     'No: el método estático no aplica a estructuras del grupo A' if p.get('grupo') == 'A' else
                     ('No: la altura excede el límite' if nm['limite'] else 'No se puede confirmar: captura la zona geotécnica'))]])]

    # ---- 2 coeficiente
    enc, filas = ame.tabla_derivacion_estatica(res)
    if dl:
        l1 = 'Voy a calcular C como la ordenada elástica por Ks, en el plato y en T = 0, y tomo la mayor (limitación de daño: sin Q′ ni R′).'
    else:
        l1 = "Voy a calcular C = Sa/(Q'R') en el plato del espectro (c) y en T = 0 (a0) con las mismas reglas de Q′ y R′, y tomo la mayor."
    B += [paso(2, 'Coeficiente sísmico C',
               l1, f'Para este problema gobierna {se["gobierna"]}: C = {se["C"]:.6f} g.'),
          formula(r'C=\max\left\{\dfrac{c}{Q^{\prime}R^{\prime}}\ (\text{plato}),\ \dfrac{a_0}{Q^{\prime}R^{\prime}}\ (T=0)\right\}',
                  "C = máx { c/(Q'R') en el plato ; a0/(Q'R') en T = 0 }"),
          tabla('', enc, filas, 5),
          destacado('Aceleración del método estático', items_aceleraciones_estatico(res))]

    # ---- 3 cortante basal
    WT = float(se['W_T'])
    B += [paso(3, 'Cortante basal V',
               'Voy a multiplicar el coeficiente por el peso total del edificio para obtener la fuerza lateral total.',
               'Ese cortante se reparte después entre los niveles.'),
          formula(r'V=C\cdot\sum W_i', 'V = C · ΣW'),
          texto(f'V = {se["C"]:.6f} · {WT:.4f} t = {se["C"] * WT:.4f} t')]

    # ---- 4 distribución
    B += [paso(4, 'Distribución de la fuerza por nivel',
               'Voy a repartir V en proporción a W_i·H_i: los pisos más pesados y más altos reciben más fuerza.',
               'La suma de las F_i debe dar otra vez V; es la comprobación de este paso.'),
          formula(r'F_i=C\,\dfrac{W_iH_i}{\sum W_jH_j}\,\sum W_j', 'F_i = C · (W_i·H_i / Σ W_j·H_j) · ΣW'),
          tabla('', ['Nivel', 'W_i (t)', 'H_i (m)', 'W_i·H_i', 'F_i (t)'],
                [[int(e['nivel'][i]), float(e['Wi'][i]), float(e['h'][i]), float(e['hWi'][i]), float(e['Fi'][i])] for i in range(n)], 5),
          texto(f'Σ W_i = {float(e["Wi"].sum()):.4f} t     Σ W_i·H_i = {float(e["hWi"].sum()):.4f} t·m     '
                f'Σ F_i = {float(e["Fi"].sum()):.5f} t  =  C·ΣW = {se["C"] * WT:.5f} t')]

    # ---- 5 diseño
    Vi = np.cumsum(e['Fi'])
    B += [paso(5, 'Fuerzas de diseño',
               f'Voy a aplicar el factor de carga FC = {FC:g}, acumular cortantes desde arriba y calcular los momentos Mv = F·H y Mvu = Fu·H.',
               'Con ello queda la tabla de fuerzas estáticas de diseño: Fu, Vu y Mvu por nivel.'),
          tabla('', ['Nivel', 'F_i (t)', 'V_i (t)', 'Mv = F·H (t·m)', 'Fu (t)', 'Vu (t)', 'Mvu (t·m)'],
                [[int(e['nivel'][i]), float(e['Fi'][i]), float(Vi[i]), float(e['Fi'][i] * e['h'][i]), float(e['Fu'][i]),
                  float(e['Vu'][i]), float(e['Mvu'][i])] for i in range(n)], 5),
          texto(f'Σ Mv = {float((e["Fi"] * e["h"]).sum()):.4f} t·m     Σ Mvu = {float(e["Mvu"].sum()):.4f} t·m     '
                f'V estático de diseño Vu = {float(e["Vu"][-1]):.4f} t')]

    # ---- 6 distorsiones
    enc, filas = ame.tabla_distorsiones_estatico(res)
    B += [paso(6, 'Distorsiones por el método estático',
               'Voy a calcular los desplazamientos con K⁻¹·F_i, a amplificarlos si corresponde y a dividir la deriva entre la altura de entrepiso.',
               'Comparo con la distorsión del análisis modal espectral y con el límite del estado de diseño.'),
          tabla('', enc[:-1], [f[:-1] for f in filas], 5),
          texto('Cumple (estático / dinámico): ' + '; '.join(f'N{f[0]}: {f[-1]}' for f in filas))]

    # ---- 7 factor de escala
    te = res['tabla_din_esc']
    Vi_est, Vi_din = float(e['Fi'].sum()), float(res['tabla_din']['F'].sum())
    B += [paso(7, 'Factor de escala del análisis modal espectral',
               'Voy a comparar el cortante basal del modal espectral con el del estático: si el dinámico es menor, sus fuerzas se escalan hacia arriba.',
               'El factor es V_estático / V_dinámico y solo se aplica cuando es mayor que 1.'),
          formula(r'\text{Factor}=\dfrac{V_{est}}{V_{din}}', 'Factor = V_est / V_din'),
          texto(f'Factor = {Vi_est:.4f} / {Vi_din:.4f} = {res["factor_escala"]:.6f}'
                + ('   → se escalan las fuerzas dinámicas' if res['factor_aplicado'] > 1 else '   → no hace falta escalar (el dinámico ya cumple)'))]
    if res['factor_aplicado'] > 1:
        B.append(tabla(f'Fuerzas dinámicas escaladas (× {res["factor_aplicado"]:.4f})', ['Nivel', 'H (m)', 'F (t)', 'Fu (t)', 'Vu (t)', 'Mvu (t·m)'],
                       [[int(te['nivel'][i]), float(te['h'][i]), float(te['F'][i]), float(te['Fu'][i]), float(te['Vu'][i]),
                         float(te['Mvu'][i])] for i in range(n)], 4))
    return B


# ----------------------------------------------------------------------------
# Texto plano (consola, pruebas)
# ----------------------------------------------------------------------------
def _fmt(v, dec):
    if isinstance(v, (float, np.floating)):
        return f'{v:,.{dec}f}'.replace(',', '')
    return str(v)


def bloques_a_texto(bloques):
    s = []
    for b in bloques:
        t = b['t']
        if t == 'titulo':
            s += ['', '=' * 78, b['texto'] + (f'   ({b["sub"]})' if b['sub'] else ''), '=' * 78, '']
        elif t == 'paso':
            s += ['', '', '-' * 78, f'PASO {b["n"]} · {b["texto"].upper()}', '-' * 78] + list(b['expl']) + ['']
        elif t == 'formula':
            s += ['    ' + b['texto'], '']
        elif t == 'matriz':
            s.append(f'{b["nombre"]}' + (f'   [{b["unidad"]}]' if b['unidad'] else '') + ' =')
            if b['simbolica'] is not None:
                for fila in b['simbolica']:
                    s.append('   [ ' + '  '.join(f'{x:>14}' for x in fila) + ' ]')
            else:
                for fila in b['datos']:
                    s.append('   [ ' + '  '.join(f'{_fmt(float(x), b["dec"]):>14}' for x in fila) + ' ]')
            s.append('')
        elif t == 'tabla':
            if b['titulo']:
                s.append(b['titulo'])
            s.append('  '.join(f'{h:>16}' for h in b['enc']))
            for fila in b['filas']:
                s.append('  '.join(f'{_fmt(v, b["dec"]):>16}' for v in fila))
            s.append('')
        elif t == 'destacado':
            s += [b['titulo'].upper()] + [f'   {e}:  {v}   ({dt})' for e, v, dt in b['items']] + ['']
        elif t in ('texto', 'nota'):
            s += [b['texto'], '']
    return '\n'.join(s)


# ----------------------------------------------------------------------------
# Resumen de los dos análisis (sin elástico)
# ----------------------------------------------------------------------------
def resumen_metodos(res):
    """(encabezados, filas) con el cortante basal y el Σ Mvu del modal espectral, del estático y del modal escalado."""
    t = res['tabla_din']
    enc = ['Análisis', 'V base (t)', 'Σ Mvu (t·m)', 'Observación']
    filas = [['Modal espectral', float(res['V_din']), float(t['Mvu'].sum()),
              f'{res["n_modos"]} de {res["p"]["n"]} modos · {res["p"]["combinacion"]}']]
    if 'tabla_est' in res:
        e = res['tabla_est']
        filas.append(['Estático', float(e['Vu'][-1]), float(e['Mvu'].sum()),
                      f'C = {res["estatico"]["C"]:.5f} g' + ('' if res['normativa']['estatico']['aplica'] else ' · solo referencia')])
        if res['factor_aplicado'] > 1:
            te = res['tabla_din_esc']
            filas.append(['Modal espectral escalado', float(te['Vu'][-1]), float(te['Mvu'].sum()),
                          f'× {res["factor_aplicado"]:.4f} (V estático / V dinámico)'])
    return enc, filas


def niveles_metodos(res):
    """Fu y Vu por nivel (arriba → abajo) del modal, del estático y del modal escalado."""
    n = res['p']['n']
    t = res['tabla_din']
    enc = ['Nivel', 'Fu modal (t)', 'Vu modal (t)']
    cols = [t['Fu'], t['Vu']]
    if 'tabla_est' in res:
        e = res['tabla_est']
        enc += ['Fu estático (t)', 'Vu estático (t)']
        cols += [e['Fu'], e['Vu']]
        if res['factor_aplicado'] > 1:
            te = res['tabla_din_esc']
            enc += ['Fu modal escalado (t)', 'Vu modal escalado (t)']
            cols += [te['Fu'], te['Vu']]
    filas = [[int(t['nivel'][i])] + [float(c[i]) for c in cols] for i in range(n)]
    return enc, filas


# ----------------------------------------------------------------------------
# Excel
# ----------------------------------------------------------------------------
def bloques_a_hoja(ws, bloques):
    """Escribe los bloques en una hoja de openpyxl (con dos filas en blanco entre pasos)."""
    from openpyxl.styles import Alignment, Font, PatternFill
    azul = PatternFill('solid', fgColor='1D4E89')
    claro = PatternFill('solid', fgColor='DDEBF7')
    fila = [1]

    def esc(valor, col=1, **kw):
        c = ws.cell(fila[0], col, valor)
        for k, v in kw.items():
            setattr(c, k, v)
        return c

    def sig(n=1):
        fila[0] += n

    def fmt(dec):
        return '#,##0' + ('.' + '0' * dec if dec else '')

    for b in bloques:
        t = b['t']
        if t == 'titulo':
            for col in range(1, 11):
                ws.cell(fila[0], col).fill = PatternFill('solid', fgColor='0B2545')
            esc(b['texto'], font=Font(bold=True, size=18, color='FFFFFF'))
            ws.cell(fila[0], 1).fill = PatternFill('solid', fgColor='0B2545')
            sig()
            if b['sub']:
                esc(b['sub'], font=Font(italic=True, size=11, color='1D4E89'))
                sig()
            sig(2)
        elif t == 'paso':
            sig(2)
            for col in range(1, 11):
                ws.cell(fila[0], col).fill = azul
            esc(f'PASO {b["n"]} · {b["texto"]}', font=Font(bold=True, size=14, color='FFFFFF'), fill=azul)
            sig()
            for linea in b['expl']:
                esc(linea, font=Font(italic=True, size=11, color='13315C'))
                sig()
            sig()
        elif t == 'formula':
            esc(b['texto'], font=Font(italic=True, size=11, color='1D4E89'))
            sig(2)
        elif t == 'matriz':
            esc(b['nombre'] + (f'   [{b["unidad"]}]' if b['unidad'] else ''), font=Font(bold=True, size=11))
            sig()
            n_f, n_c = b['datos'].shape
            cols = b['cols'] or [str(j + 1) for j in range(n_c)]
            filas_l = b['filas'] or [str(i + 1) for i in range(n_f)]
            esc('', 1)
            for j, c_ in enumerate(cols):
                esc(c_, 2 + j, font=Font(bold=True), fill=claro, alignment=Alignment(horizontal='center'))
            sig()
            for i in range(n_f):
                esc(filas_l[i], 1, font=Font(bold=True), fill=claro)
                for j in range(n_c):
                    if b['simbolica'] is not None:
                        esc(b['simbolica'][i][j].replace('{', '').replace('}', ''), 2 + j, alignment=Alignment(horizontal='center'))
                    else:
                        c_ = esc(float(b['datos'][i, j]), 2 + j)
                        c_.number_format = fmt(b['dec'])
                sig()
            sig(2)
        elif t == 'tabla':
            if b['titulo']:
                esc(b['titulo'], font=Font(bold=True, size=11))
                sig()
            for j, h in enumerate(b['enc']):
                esc(h, 1 + j, font=Font(bold=True), fill=claro, alignment=Alignment(horizontal='center', wrap_text=True))
            sig()
            for f_ in b['filas']:
                for j, v in enumerate(f_):
                    v = v.item() if isinstance(v, np.generic) else v
                    c_ = esc(v, 1 + j)
                    if isinstance(v, float):
                        c_.number_format = fmt(b['dec'])
                sig()
            sig(2)
        elif t == 'destacado':
            amarillo = PatternFill('solid', fgColor='FFF2CC')
            esc(b['titulo'], font=Font(bold=True, size=12, color='13315C'))
            sig()
            for e_, v_, dt_ in b['items']:
                esc(e_, 1, font=Font(bold=True, size=11), fill=amarillo)
                for col in (2, 3):
                    ws.cell(fila[0], col).fill = amarillo
                esc(v_, 4, font=Font(bold=True, size=14, color='1D4E89'), fill=amarillo, alignment=Alignment(horizontal='center'))
                esc(dt_, 5, font=Font(size=11), fill=amarillo, alignment=Alignment(horizontal='center'))
                sig()
            sig(2)
        elif t in ('texto', 'nota'):
            esc(b['texto'], font=Font(italic=(t == 'nota'), size=10))
            sig(2)
    for col in range(1, 12):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = 19
    ws.column_dimensions['A'].width = 24


# ----------------------------------------------------------------------------
# Formato compacto (PDF): como las memorias de los exámenes
# ----------------------------------------------------------------------------
def _sin_explicaciones(bloques):
    """Quita el título grande y convierte cada paso en un subtítulo (sin sus dos líneas de explicación)."""
    out = []
    for b in bloques:
        if b['t'] == 'titulo':
            continue
        out.append(sub(f'{b["texto"]}') if b['t'] == 'paso' else b)
    return out


def _compacto_modal(res):
    p, n = res['p'], res['p']['n']
    nv = _niveles(n)
    M, K = np.asarray(p['M'], float), np.asarray(p['K'], float)
    modos = [f'Modo {j + 1}' for j in range(n)]
    B = [seccion('A', 'Matrices de masa y de rigidez, periodos y formas modales')]
    aj = p.get('ajuste_periodo')
    if aj:
        B.append(sub('Peso obtenido del periodo medido'))
        B += [b for b in _bloques_ajuste(p, aj, M, K) if b['t'] != 'paso']
    else:
        B.append(sub('Matriz de masas M'))
        B += _bloques_masas(p, M, nv)
    B.append(sub('Matriz de rigidez K'))
    B += [b for b in _bloques_rigidez(p, K, nv) if b['t'] != 'texto']
    lam, w, f, T, fi = res['lam'], res['w'], res['f'], res['T'], np.asarray(res['fi'], float)
    mo = res['normativa']['modos']
    B += [sub('Periodos, frecuencias, participación y formas modales'),
          tabla('', ['Modo', 'λ = ω² (rad/s)²', 'ω (rad/s)', 'f (Hz)', 'T (s)', 'γ', 'Masa efectiva'],
                [[j + 1, float(lam[j]), float(w[j]), float(f[j]), float(T[j]), float(res['gam'][j]), float(mo['m_ef'][j])]
                 for j in range(n)], 5),
          matriz('Φ (formas modales; cada columna es un modo)', fi, 5, '', nv, modos),
          figura('Formas modales', ame.fig_modal(res), 17)]

    # ---- B. espectro y fuerzas (mismo orden que la solución del profesor)
    clave_q = ame.CLAVE_ESTADO[p['estado']]
    dl = res['clave'] == 'DL'
    pq = p if clave_q == 'SV' else dict(p, Q=np.array([1.0]))
    Sa_el = ame._interp(T, res['serie'], ame.espectro(dict(p, Ks=1.0), 1.0, 1.0, res['serie'], dl=True)[0])
    Sa_d = np.asarray(res['Sa_modal'][:, 0], float)
    B.append(seccion('B', 'Fuerzas sísmicas de diseño · modal espectral'))
    quiero = ('Q', 'k1', 'R0', 'k2', 'R = ', "R′", 'Q·R′ (distorsiones)', 'FC·Q·R′', 'γc', 'Ks', 'Amortiguamiento')
    filas_n = [[a_, b_] for a_, b_, _ in res['normativa']['factores'] if str(a_).startswith(quiero)]
    pares = []
    for i in range(0, len(filas_n), 2):
        pares.append(filas_n[i] + (filas_n[i + 1] if i + 1 < len(filas_n) else ['', '']))
    B.append(tabla('Parámetros del espectro y de la reducción', ['Concepto', 'Valor', 'Concepto', 'Valor'], pares))
    gam = [float(x) for x in res['gam']]
    if dl:
        enc = ['Modo', 'T (s)', 'Sa elástica (cm/s²)', 'Ks', 'Sa·Ks (cm/s²)', 'Sa (g)', 'Partic.']
        filas = [[j + 1, float(T[j]), float(Sa_el[j]), float(p['Ks']), float(Sa_d[j]), float(Sa_d[j] / G), gam[j]] for j in range(n)]
    else:
        q_, r_ = ame._q_r(pq, float(pq['Q'][0]), T, clave_q)
        enc = ['Modo', 'T (s)', 'Sa elástica (cm/s²)', "Q'·R'", 'Sa diseño (cm/s²)', 'Sa diseño (g)', 'Partic.']
        filas = [[j + 1, float(T[j]), float(Sa_el[j]), float(q_[j] * r_[j]), float(Sa_d[j]), float(Sa_d[j] / G), gam[j]]
                 for j in range(n)]
    B.append(destacado('Aceleraciones espectrales de diseño', items_aceleraciones_modal(res)))
    B.append(tabla('Aceleraciones espectrales y factores de participación modal', enc, filas, 5))
    B.append(texto(f'Suma de los factores de participación = {sum(gam):.5f}'
                   + (f'     Ordenada mínima a_min = {res["a_min"]:.4f} cm/s² ({res["a_min_g"]:.5f} g)' if res.get('a_min') else '')))
    B.append(figura('Espectro elástico y de diseño', ame.fig_espectros_comparados(res), 14))

    F = np.asarray(res['F'], float)
    t_ = res['tabla_din']
    FC = 1.0 if dl else p['factor_Fu']
    Fkg = np.asarray(t_['F'], float) * 1000                 # arriba -> abajo, kg
    Vkg = np.cumsum(Fkg)
    B += [sub('Fuerzas sísmicas modales (cada columna es un modo)'),
          matriz('F', F, 2, 'kg', nv, modos),
          sub(f'Combinación {p["combinacion"]} sin factor de carga'),
          tabla('', ['Nivel', 'Fuerza (kg)', 'Cortante (kg)'],
                [[int(t_['nivel'][i]), float(Fkg[i]), float(Vkg[i])] for i in range(n)], 1),
          sub(f'Fuerzas sísmicas con el análisis modal espectral NTC-Sismo (FC = {FC:g})'),
          tabla('', ['Nivel', 'Fuerza (kg)', 'Cortante (kg)', 'Fu (kg)', 'Vu (kg)', 'Mvu (kg·cm)'],
                [[int(t_['nivel'][i]), float(Fkg[i]), float(Vkg[i]), float(t_['Fu'][i] * 1000), float(t_['Vu'][i] * 1000),
                  float(t_['Mvu'][i] * 1e5)] for i in range(n)], 1),
          texto(f'Momento de volteo de diseño = Σ Mvu = {float(t_["Mvu"].sum() * 1e5):,.1f} kg·cm '
                f'= {float(t_["Mvu"].sum()):.4f} t·m     Vu en la base = {float(t_["Vu"][-1]) * 1000:,.1f} kg'),
          figura('Diagramas de Fu, Vu y Mvu', ame.fig_diagramas(res), 12.5)]
    enc, filas = ame.tabla_distorsiones(res)
    B.append(tabla('Distorsiones de entrepiso por estado límite', enc, filas, 5))
    return B


def _compacto_estatico(res):
    """Sección C como la solución del profesor: cs, tabla estática, factor de escala y fuerzas dinámicas escaladas."""
    p, n = res['p'], res['p']['n']
    e, se = res['tabla_est'], res['estatico']
    Fi = np.asarray(e['Fi'], float)
    Vi = np.cumsum(Fi)
    Mv = Fi * np.asarray(e['h'], float)
    Wi = np.asarray(e['Wi'], float)
    H = np.asarray(e['h'], float)
    h_ent = [float(H[i] - (H[i + 1] if i + 1 < n else 0.0)) for i in range(n)]      # altura de entrepiso, arriba -> abajo
    B = [seccion('C', 'Método estático, factor de escala y fuerzas escaladas'),
         destacado('Aceleración del método estático', items_aceleraciones_estatico(res)),
         texto(f'V = cₛ·ΣW = {se["C"]:.5f} · {float(se["W_T"]):.4f} t = {se["C"] * float(se["W_T"]):.4f} t'),
         tabla('Fuerzas estáticas', ['Niv.', 'Peso Wi (t)', 'Alt. entrepiso (m)', 'Hi (m)', 'Wi·Hi', 'Fi (t)', 'Vi (t)',
                                     'Mv (t·m)', 'Fu (t)', 'Vu (t)', 'Mvu (t·m)'],
               [[int(e['nivel'][i]), float(Wi[i]), h_ent[i], float(H[i]), float(e['hWi'][i]), float(Fi[i]),
                 float(Vi[i]), float(Mv[i]), float(e['Fu'][i]), float(e['Vu'][i]), float(e['Mvu'][i])] for i in range(n)]
               + [['Σ', float(Wi.sum()), '', '', float(np.asarray(e['hWi']).sum()), float(Fi.sum()), '', float(Mv.sum()),
                   float(np.asarray(e['Fu']).sum()), '', float(np.asarray(e['Mvu']).sum())]], 4)]
    Vi_est, Vi_din = float(Fi.sum()), float(res['tabla_din']['F'].sum())
    B.append(texto(f'Factor de escala: V_est / V_din = {Vi_est:.4f} / {Vi_din:.4f} = {res["factor_escala"]:.4f}'
                   + ('' if res['factor_aplicado'] > 1 else '   (no se escala: el dinámico ya cumple)')))
    if res['factor_aplicado'] > 1:
        te = res['tabla_din_esc']
        Fe = np.asarray(te['F'], float)
        Ve = np.cumsum(Fe)
        Mve = Fe * np.asarray(te['h'], float)
        B.append(tabla('Fuerzas dinámicas escaladas al 100 % del estático',
                       ['Nivel', 'Fi (t)', 'Vi (t)', 'Mv (t·m)', 'Fu (t)', 'Vu (t)', 'Mvu (t·m)'],
                       [[int(te['nivel'][i]), float(Fe[i]), float(Ve[i]), float(Mve[i]), float(te['Fu'][i]), float(te['Vu'][i]),
                         float(te['Mvu'][i])] for i in range(n)]
                       + [['Σ', float(Fe.sum()), '', float(Mve.sum()), float(np.asarray(te['Fu']).sum()), '',
                           float(np.asarray(te['Mvu']).sum())]], 4))
    B.append(figura('Diagramas de Fu, Vu y Mvu (estático)', ame.fig_diagramas(res, 'estatico'), 12.5))
    return B


def compacto(res, metodo='ambos'):
    """Bloques de la memoria compacta: A (matrices y modos), B (espectro y fuerzas) y C (estático)."""
    B = []
    if metodo in ('ambos', 'modal'):
        B += _compacto_modal(res)
    if metodo in ('ambos', 'estatico') and 'tabla_est' in res:
        B += _compacto_estatico(res)
    return B
