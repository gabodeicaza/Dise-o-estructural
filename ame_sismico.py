"""
Análisis Modal Espectral (AME) - marcos planos
==============================================
Equivalente al script de MATLAB, con captura de datos flexible, vigas flexibles opcionales, revisiones y exportación.

Uso:
    python ame_app.py                           -> abre la interfaz gráfica (recomendado)
    python ame_sismico.py proyecto.json         -> calcula sin interfaz y exporta a Excel
    python ame_sismico.py proyecto.json -o salida.xlsx

A partir de secciones, alturas, cargas, Q, k1 y tipo de edificación arma K y M y hace el AME.
(El proyecto también puede traer K y M escritas a mano: ver preparar_datos.)
Unidades: cm, kg, s  (a0, c y Sa en cm/s^2; cargas en t/m, longitudes y alturas en m).
"""
import json
import re
import sys
from pathlib import Path

import numpy as np
from matplotlib.figure import Figure

import ame_ntc as ntc
from ame_marco import rigidez_marco
from ame_revision import (COMBINACIONES, avisos as _avisos, combinar, derivas_modales, irregularidades,
                          theta_pdelta)
from ame_secciones import AYUDA_SECCIONES, ancho_en_direccion, dibujar_seccion, inercia, parse_secciones

G = 981.0  # cm/s^2
AZULES = ['#0B2545', '#3E7CB1', '#1D4E89', '#81A4CD', '#13315C']
ESTADOS = ('Seguridad de vida', 'Ocupación inmediata', 'Limitación de daños')
CLAVE_ESTADO = {'Seguridad de vida': 'SV', 'Ocupación inmediata': 'OI', 'Limitación de daños': 'DL'}
MATERIALES = ('Concreto', 'Acero', 'Mampostería', 'Otro')

# ----------------------------------------------------------------------------
# Lectura flexible de datos
# ----------------------------------------------------------------------------
_NUM = re.compile(r'[-+]?(?:\d+\.?\d*|\.\d+)(?:[eE][-+]?\d+)?')
_MILES = re.compile(r'(?<![\d.])\d{1,3}(?:,\d{3})+(?=\.\d|\s|$|;|\])')


def parse_filas(texto):
    """Devuelve lista de filas (listas de float). Filas separadas por salto de línea o ';'."""
    t = texto.replace('−', '-').replace('[', ' ').replace(']', ' ')
    t = _MILES.sub(lambda m: m.group().replace(',', ''), t)
    filas = []
    for linea in re.split(r'[\n;]', t):
        nums = [float(x) for x in _NUM.findall(linea)]
        if nums:
            filas.append(nums)
    return filas


def parse_matriz(texto, nombre):
    filas = parse_filas(texto)
    if not filas:
        raise ValueError(f'La matriz {nombre} está vacía.')
    if len({len(f) for f in filas}) != 1:
        raise ValueError(f'La matriz {nombre} no es rectangular (filas con distinto número de datos).')
    return np.array(filas, dtype=float)


def parse_vector(texto, nombre):
    filas = parse_filas(texto)
    v = [x for f in filas for x in f]
    if not v:
        raise ValueError(f'Falta el dato: {nombre}.')
    return np.array(v, dtype=float)


def parse_num(texto, nombre):
    try:
        return float(str(texto).strip().replace(',', '.'))
    except ValueError:
        raise ValueError(f'Valor numérico no válido en "{nombre}": {texto!r}')


def preparar_datos(d):
    """K y M escritas a mano (GUI/JSON) -> datos numéricos validados.

    Modo nuevo (hay 'subgrupo' o 'intensidad'): además de K y M se dan las alturas de entrepiso 'alturas' (m, de abajo
    hacia arriba) y, opcionalmente, los pesos 'pesos' (t); si faltan, W = m·g. 'unidades': 'kg, cm' (por defecto) o 't, cm'.
    Modo anterior: 'grupo', 'sismo' y 'niveles' (una fila por nivel: h, carga, longitud).
    """
    K = parse_matriz(d['K'], 'K')
    if K.shape[0] != K.shape[1]:
        raise ValueError(f'K debe ser cuadrada (es {K.shape[0]}x{K.shape[1]}).')
    n = K.shape[0]
    M = parse_matriz(d['M'], 'M')
    if M.shape == (n, n):
        pass
    elif M.size == n and 1 in M.shape:
        M = np.diag(M.ravel())
    else:
        raise ValueError(f'M debe ser {n}x{n} o un vector de {n} masas (es {M.shape[0]}x{M.shape[1]}).')
    factor = 1000.0 if str(d.get('unidades', '')).strip().lower().startswith('t') else 1.0  # t -> kg
    K, M = K * factor, M * factor

    if str(d.get('alturas', '')).strip():
        alturas = parse_vector(d['alturas'], 'Alturas de entrepiso')
        if len(alturas) != n:
            raise ValueError(f'Hay {len(alturas)} alturas y K es de {n} grados de libertad.')
        if (alturas <= 0).any():
            raise ValueError('Las alturas de entrepiso deben ser mayores que cero.')
        if str(d.get('pesos', '')).strip():
            W = parse_vector(d['pesos'], 'Pesos por nivel')
            if len(W) != n:
                raise ValueError(f'Hay {len(W)} pesos y K es de {n} grados de libertad.')
        else:
            W = np.diag(M) * G / 1000  # peso (t) a partir de la masa (kg·s²/cm)
        cargas, dist_x = W, np.ones(n)
    else:
        niv = parse_filas(d['niveles'])
        if len(niv) != n:
            raise ValueError(f'En "Niveles" hay {len(niv)} filas y K es de {n} grados de libertad.')
        alturas = np.array([f[0] for f in niv])
        if all(len(f) >= 3 for f in niv):
            cargas = np.array([f[1] for f in niv])
            dist_x = np.array([f[2] for f in niv])
        else:
            cargas = dist_x = None  # sin estático

    if d.get('subgrupo') or d.get('intensidad'):
        mat = d.get('material') or 'Concreto'
        obj = _objetivo_de(d)
        p = dict(M=M, K=K, n=n, alturas=alturas, cargas=cargas, dist_x=dist_x, grupo=obj['grupo'], sismo=obj['sismo'],
                 subgrupo=obj['subgrupo'], intensidad=obj['intensidad'], objetivo=obj['objetivo'], material=mat,
                 R_legacy=False, manual=True, **_params_norma(d, mat),
                 **_params_espectro(d, obj['estado_auto'], mat, obj['intensidad']))
        T_med = _periodo_medido(d)
        if T_med:
            _ajustar_periodo(p, T_med)
    else:
        grupo = d['grupo'].strip().upper()[:1]
        sismo = d['sismo'].strip().upper()[:1]
        if grupo not in ('A', 'B'):
            raise ValueError('El grupo de estructura debe ser A o B.')
        if grupo == 'A' and sismo not in ('B', 'I'):
            raise ValueError('Para el grupo A indica el sismo: B (base) o I (infrecuente).')
        p = dict(M=M, K=K, n=n, alturas=alturas, cargas=cargas, dist_x=dist_x,
                 grupo=grupo, sismo=sismo, **_params_espectro(d))
        if cargas is not None:  # el estático usa el peso Wi = carga * longitud
            p['cargas'], p['dist_x'] = cargas * dist_x, np.ones(n)
    return p


def preparar(d):
    """Arma el modelo según el modo del proyecto: 'matrices' (K y M dadas) o secciones y cargas."""
    return preparar_datos(d) if d.get('modo') == 'matrices' else preparar_modelo(d)


def _num_opc(d, clave, defecto):
    """Número opcional: si la clave no existe -> defecto; si está vacía -> None."""
    if clave not in d:
        return defecto
    return parse_num(d[clave], clave) if str(d[clave]).strip() else None


def _spec_dict(d, intensidad):
    """Parámetros del espectro de la intensidad dada: d['espectros'][intensidad] o los campos sueltos."""
    esp = d.get('espectros')
    if not esp:
        return d
    fuente = intensidad
    if fuente not in esp and ntc.espectro_fuente(intensidad) in esp:
        fuente = ntc.espectro_fuente(intensidad)  # frecuente y ocasional salen del espectro de base de diseño × Ks
    if fuente not in esp:
        nombre = ntc.espectro_fuente(intensidad) if intensidad in ('Frecuente', 'Ocasional') else intensidad
        extra = ' (de él salen la intensidad frecuente y la ocasional)' if nombre != intensidad else ''
        raise ValueError(f'Captura el espectro SASID de la intensidad "{nombre}"{extra} (faltan: a0, c, Ta, Tb, k, Ts).')
    src = dict(d, **esp[fuente])
    faltan = [k for k in ('a0', 'c', 'Ta', 'Tb', 'k', 'Ts') if not str(src.get(k, '')).strip()]
    if faltan:
        raise ValueError(f'Captura el espectro SASID de la intensidad "{intensidad}" (faltan: {", ".join(faltan)}).')
    return src


FACTOR_ESFUERZO = {'kg/cm²': 1.0, 't/cm²': 1000.0, 'MPa': 10.19716, 't/m²': 0.1, 'GPa': 10197.16}  # a kg/cm²


def _a_kgcm2(valor, unidad, nombre):
    unidad = unidad or 'kg/cm²'
    if unidad not in FACTOR_ESFUERZO:
        raise ValueError(f'Unidad no válida para {nombre}: {unidad}. Usa {", ".join(FACTOR_ESFUERZO)}.')
    return valor * FACTOR_ESFUERZO[unidad]


def _factor_espectro(d):
    """Las ordenadas a0 y c se capturan en cm/s² (por defecto) o en g."""
    u = str(d.get('unidad_espectro') or 'cm/s²').strip().lower()
    if u.startswith('g'):
        return G
    if not u.startswith('cm'):
        raise ValueError('Las unidades del espectro deben ser cm/s² o g.')
    return 1.0


def _params_espectro(d, estado_auto=None, material='Concreto', intensidad='Base de diseño'):
    pedido = d.get('estado')
    estado = (estado_auto or ESTADOS[0]) if (not pedido or str(pedido).startswith('Autom')) else pedido
    comb = d.get('combinacion') or COMBINACIONES[0]
    if estado not in ESTADOS:
        raise ValueError(f'Estado límite no válido: {estado}.')
    if comb not in COMBINACIONES:
        raise ValueError(f'Combinación modal no válida: {comb}.')
    zeta_defecto = ntc.ZETA_MATERIAL.get(material, 0.05)
    zeta = _num_opc(d, 'zeta', zeta_defecto)
    sp = _spec_dict(d, intensidad)
    fu = _factor_espectro(d)
    ks_base = _num_opc(d, 'Ks', 0.25) or 0.25
    sa_txt = str(d.get('Sa_modos') or '').strip()
    return dict(Q=parse_vector(d['Q'], 'Q'), unidad_espectro='g' if fu == G else 'cm/s²',
                Sa_manual=parse_vector(sa_txt, 'Ordenadas Sa por modo') * fu if sa_txt else None,
                k1=parse_num(d['k1'], 'k1'), a0=parse_num(sp['a0'], 'a0') * fu, c=parse_num(sp['c'], 'c') * fu,
                Ta=parse_num(sp['Ta'], 'Ta'), Tb=parse_num(sp['Tb'], 'Tb'), k=parse_num(sp['k'], 'k'),
                Ts=parse_num(sp['Ts'], 'Ts'),
                factor_Fu=parse_num(d.get('factor_Fu', '1.1') or '1.1', 'Factor Fu'),
                mult_V=parse_num(d.get('mult_V', '1') or '1', 'Multiplicador V estático'),
                estado=estado, Ks=ks_base * (ntc.FACTOR_KS_OCASIONAL if intensidad == 'Ocasional' else 1.0), Ks_base=ks_base,
                combinacion=comb,
                zeta=zeta if zeta is not None else zeta_defecto,
                lim_dl=_num_opc(d, 'lim_dl', 0.004), lim_sv=_num_opc(d, 'lim_sv', 0.03),
                lim_oi=_num_opc(d, 'lim_oi', 0.005))


NORMA_DEFECTO = dict(irreg=[], fuerte_torsion=False, fuerte_elev=False, gamma_c=1.0,
                     gamma_c_txt='Estructura regular: γc = 1.0.', b_planta=None, zona='', regular=True,
                     R_unitaria=False, R1_otros=True, intensidad='Base de diseño', subgrupo=None, objetivo=None,
                     material='Concreto', R_legacy=True, criterio_est='apuntes', R_dist='plato', e_s=None,
                     modos_modo='todos', modos_n=None, Sa_manual=None, unidad_espectro='cm/s²', norm_modal='primero')


def _flag(v):
    return v is True or str(v).strip().lower() in ('true', 'sí', 'si', '1', 'yes')


def _objetivo_de(d):
    """Grupo, subgrupo, intensidad y objetivo de diseño (NTC-S 2023) a partir de subgrupo/tipo e intensidad."""
    sub = str(d.get('subgrupo') or '').strip().upper()
    tipo = str(d.get('tipo') or 'Grupo B')
    if sub in ntc.SUBGRUPOS:
        grupo = ntc.SUBGRUPOS[sub]
    elif sub:
        raise ValueError(f'Subgrupo no válido: {sub} (usa B2, B1, A2 o A1).')
    else:
        grupo, _ = grupo_sismo(tipo)
        sub = grupo
    intensidad = ntc.normalizar_intensidad(d.get('intensidad') or ('Infrecuente' if 'infrec' in tipo.lower() else 'Base de diseño'))
    if intensidad not in ntc.INTENSIDADES:
        raise ValueError(f'Intensidad sísmica no válida: {intensidad}.')
    objetivo = ntc.objetivo_diseno(grupo, intensidad, sub)
    nombre = {v: k for k, v in CLAVE_ESTADO.items()}[objetivo['clave']]
    sismo = {'Frecuente': 'F', 'Ocasional': 'O', 'Base de diseño': 'B', 'Infrecuente': 'I'}[intensidad]
    return dict(grupo=grupo, subgrupo=sub, intensidad=intensidad, sismo=sismo, objetivo=objetivo, estado_auto=nombre)


def _norm_modal(texto):
    t = str(texto or '').lower()
    if 'últim' in t or 'ultim' in t or 'azotea' in t:
        return 'ultimo'
    if 'máx' in t or 'max' in t:
        return 'maximo'
    if 'masa' in t:
        return 'masa'
    return 'primero'


def _normalizar_modos(p, fi):
    """Normaliza los modos: primer nivel = 1 (apuntes), azotea = 1, valor máximo = 1 o masa modal unitaria."""
    modo = p.get('norm_modal', 'primero')
    n, j = fi.shape[0], np.arange(fi.shape[1])
    if modo == 'ultimo':
        div = fi[-1, :]
    elif modo == 'maximo':
        div = fi[np.argmax(np.abs(fi), axis=0), j]
    elif modo == 'masa':
        div = np.sqrt(np.einsum('ij,ik,kj->j', fi, p['M'], fi)) * np.where(fi[0, :] < 0, -1.0, 1.0)
    else:
        div = fi[0, :]
    return fi / np.where(np.abs(div) < 1e-12, 1.0, div)


def _volteo_entrepiso(Fu_arriba, alturas):
    """Momento de volteo en la base de cada entrepiso (t·m), de arriba hacia abajo: Σ Fu_j·(z_j − z_base)."""
    zs = np.concatenate([[0.0], np.cumsum(alturas)])
    Fu = np.asarray(Fu_arriba, float)[::-1]
    Mv = np.array([np.sum(Fu[i:] * (zs[i + 1:] - zs[i])) for i in range(len(Fu))])
    return Mv[::-1]


def apoyo_txt(c):
    """Texto del apoyo de una columna (para tablas)."""
    if c.get('k_cont'):
        return 'Continuación de la columna del nivel inferior'
    base = 'Rodillo' if c.get('rodillo') else ('Articulado' if c['artic'] else 'Empotrado')
    extra = f' · cruza hasta el piso {c["hasta"]}' if c.get('hasta') else ''
    return base + (' · articulada arriba' if c.get('artic_sup') else '') + extra


def _periodo_medido(d):
    """Periodo medido opcional (s). Si se escribe, debe ser positivo."""
    T = _num_opc(d, 'T_medido', None)
    if T is not None and T <= 0:
        raise ValueError('El periodo medido debe ser mayor que cero.')
    return T


def _ajustar_periodo(p, T_med):
    """Escala masas y pesos para que el periodo fundamental sea el medido (T ∝ √m). Los pesos dados son solo la
    distribución relativa entre niveles."""
    if T_med <= 0:
        raise ValueError('El periodo medido debe ser mayor que cero.')
    lam = np.linalg.eigvals(np.linalg.solve(p['M'], p['K'])).real
    if lam.min() <= 0:
        raise ValueError('No se puede ajustar al periodo medido: la estructura es inestable (revisa apoyos y articulaciones).')
    T1 = float(2 * np.pi / np.sqrt(lam.min()))
    f = (T_med / T1) ** 2
    p['M'] = p['M'] * f
    if p['cargas'] is not None:
        p['cargas'] = p['cargas'] * f
    for r in p.get('info', []):
        r['W'] *= f
        r['m'] *= f
    for l in p.get('geom', {}).get('niveles', []):
        l['W'] *= f
        l['w_bays'] = None
    p['ajuste_periodo'] = dict(T_medido=T_med, T_calculado=T1, factor=float(f), m_total=float(np.trace(p['M'])),
                               W_total=float(np.sum(p['cargas'])) if p['cargas'] is not None else None)
    return p


def _params_norma(d, material):
    """Irregularidades (γc), torsión accidental, método estático y R = 1 en materiales distintos del concreto."""
    irreg = d.get('irreg') or []
    if isinstance(irreg, str):
        irreg = re.findall(r'5\.[23]\.\d', irreg)
    ft, fe = _flag(d.get('fuerte_torsion')), _flag(d.get('fuerte_elev'))
    gc, txt = ntc.gamma_c(irreg, ft, fe)
    zona = str(d.get('zona') or '').strip().upper()
    if zona and zona not in ntc.LIMITES_ALTURA_ESTATICO:
        raise ValueError('La zona geotécnica debe ser I, II o III.')
    r1 = d.get('R1_otros', True)
    r1 = True if r1 == '' else _flag(r1)
    crit = str(d.get('criterio_estatico') or '').lower()
    criterio_est = 'maximo' if ('máx' in crit or 'max' in crit or 'matlab' in crit) else 'apuntes'
    rd = str(d.get('R_dist') or '').lower()
    modos = str(d.get('modos') or '').lower()
    modos_modo = 'primeros' if 'primeros' in modos else ('apuntes' if 'apuntes' in modos or '95' in modos else 'todos')
    nm_ = _num_opc(d, 'modos_n', None)
    return dict(irreg=sorted(set(irreg)), fuerte_torsion=ft, fuerte_elev=fe, gamma_c=gc, gamma_c_txt=txt,
                criterio_est=criterio_est, R_dist='T1' if ('fund' in rd or rd == 't1') else 'plato',
                e_s=_num_opc(d, 'e_s', None), modos_modo=modos_modo, modos_n=int(nm_) if nm_ else None,
                norm_modal=_norm_modal(d.get('norm_modal')),
                b_planta=_num_opc(d, 'b_planta', None), zona=zona,
                regular=str(d.get('regularidad') or 'Regular').lower().startswith('reg'),
                R_unitaria=(material != 'Concreto') and r1, R1_otros=r1)


def grupo_sismo(tipo):
    """'Grupo B' | 'Grupo A - sismo base' | 'Grupo A - sismo infrecuente' -> (grupo, sismo)."""
    t = tipo.lower()
    if 'grupo a' in t:
        return 'A', ('I' if 'infrec' in t else 'B')
    return 'B', 'B'


def preparar_estructura(d):
    """Arma K, M y alturas a partir de secciones de columnas, niveles y cargas (marco de cortante).

    d['niveles'] = lista (de abajo hacia arriba) de dicts con: h (m), ncol, b (cm), hs (cm, en la
    dirección del sismo), nart (columnas articuladas en la base), w (t/m), L (m), W (t, opcional).
    """
    niv = d['niveles']
    n = len(niv)
    if n < 1:
        raise ValueError('Agrega al menos un nivel.')
    if str(d.get('E', '')).strip():
        E = parse_num(d['E'], 'E')
    elif str(d.get('fc', '')).strip():
        E = 14000 * np.sqrt(parse_num(d['fc'], "f'c"))
    else:
        raise ValueError("Indica E o f'c del concreto.")

    filas, ks, m, W, alturas = [], [], [], [], []
    for i, r in enumerate(niv, 1):
        def num(clave, et, defecto=None):
            txt = str(r.get(clave, '')).strip()
            if not txt and defecto is not None:
                return defecto
            return parse_num(txt, f'Nivel {i} - {et}')
        h = num('h', 'h entrepiso')
        ncol, b, hs = int(num('ncol', 'Nº columnas')), num('b', 'b'), num('hs', 'h sección')
        nart = int(num('nart', 'Nº articuladas', 0))
        if not 0 <= nart <= ncol:
            raise ValueError(f'Nivel {i}: las columnas articuladas deben estar entre 0 y {ncol}.')
        if str(r.get('W', '')).strip():
            Wi = num('W', 'W')
        else:
            Wi = num('w', 'carga') * num('L', 'longitud')
        I = b * hs ** 3 / 12
        hcm = h * 100
        k_emp, k_art = 12 * E * I / hcm ** 3, 3 * E * I / hcm ** 3
        k = (ncol - nart) * k_emp + nart * k_art
        ks.append(k); alturas.append(h); W.append(Wi); m.append(Wi * 1000 / G)
        filas.append(dict(nivel=i, h=h, ncol=ncol, nart=nart, seccion=f'{b:g}x{hs:g}', I=I * ncol, k=k,
                          W=Wi, m=Wi * 1000 / G))

    K = np.zeros((n, n))
    for i in range(n):
        K[i, i] = ks[i] + (ks[i + 1] if i + 1 < n else 0)
        if i + 1 < n:
            K[i, i + 1] = K[i + 1, i] = -ks[i + 1]
    grupo, sismo = grupo_sismo(d.get('tipo', 'Grupo B'))
    return dict(M=np.diag(m), K=K, n=n, alturas=np.array(alturas), cargas=np.array(W), dist_x=np.ones(n),
                grupo=grupo, sismo=sismo, E=E, info=filas, **_params_espectro(d))


# ----------------------------------------------------------------------------
# Modelo por ejes: columnas faltantes, cargas por crujía, secciones rectangulares y circulares
# ----------------------------------------------------------------------------
def _tokens(texto):
    return [t for t in re.split(r'[\s,;]+', str(texto).strip()) if t]


def parse_ejes(texto, nejes, nombre, defecto):
    """'todos' | '1 3' | '1-3' -> lista ordenada de ejes (1..nejes). Vacío -> defecto."""
    t = str(texto).strip().lower()
    if not t:
        return list(defecto)
    if t in ('todos', 'todas', 'all', '*'):
        return list(range(1, nejes + 1))
    if t in ('0', 'no', '-', 'ninguno', 'ninguna', 'none'):
        return []  # "ninguna" (útil en articuladas: todas empotradas)
    ejes = set()
    for tok in _tokens(t):
        m = re.fullmatch(r'(\d+)(?:-(\d+))?', tok)
        if not m:
            raise ValueError(f'{nombre}: "{tok}" no es un eje válido. Escribe, por ejemplo, "todos", "1 3" o "1-3".')
        a, b = int(m.group(1)), int(m.group(2) or m.group(1))
        for e in range(a, b + 1):
            if not 1 <= e <= nejes:
                raise ValueError(f'{nombre}: el eje {e} no existe (hay {nejes} ejes, del 1 al {nejes}).')
            ejes.add(e)
    return sorted(ejes)


def _modulo_E(d):
    """E (kg/cm²) según el material: valor dado, f'c del concreto (clase 1 o 2) o 2,040,000 del acero."""
    mat = d.get('material') or 'Concreto'
    if str(d.get('E', '')).strip():
        return _a_kgcm2(parse_num(d['E'], 'E'), d.get('unidad_E'), 'E')
    if mat == 'Acero':
        return 2_040_000.0
    if mat == 'Concreto':
        if not str(d.get('fc', '')).strip():
            raise ValueError("Indica E o f'c del concreto.")
        coef = 8000 if 'clase 2' in str(d.get('clase', '')).lower() else 14000
        return coef * np.sqrt(_a_kgcm2(parse_num(d['fc'], "f'c"), d.get('unidad_fc'), "f'c"))
    raise ValueError(f'Indica el módulo de elasticidad E de la {mat.lower()}.')


def _escalar_texto(texto, factor, nombre):
    """Convierte las dimensiones de una sección escrita en metros a centímetros (factor = 100)."""
    if factor == 1.0:
        return texto
    if re.search(r'(?i)\bi:', str(texto)):
        raise ValueError(f'{nombre}: las propiedades I:/A: deben darse en cm⁴ y cm², no en metros.')
    return re.sub(r'(?<![@\d.])\d+(?:\.\d+)?', lambda m: f'{float(m.group()) * factor:g}', str(texto))


def _E_nivel(r, d, E, nom):
    """E de las columnas de un nivel: campo 'E_col' (en la unidad de E, o la palabra 'acero'); vacío = E general."""
    txt = str(r.get('E_col', '')).strip()
    if not txt:
        return E
    if txt.lower().startswith('acero'):
        return 2_040_000.0
    return _a_kgcm2(parse_num(txt, f'{nom} - E de columnas'), d.get('unidad_E'), f'{nom} - E de columnas')


def _E_columnas(r, d, E, nom, ncol):
    """E de cada columna de un nivel: vacío = E general; un valor para todas o uno por columna.
    Cada valor va en la unidad de E; las palabras 'acero' (2,040,000 kg/cm²) y 'general' también valen."""
    txt = str(r.get('E_col', '')).strip()
    if not txt:
        return [E] * ncol
    vals = []
    for t in _tokens(txt):
        tl = t.lower()
        if tl.startswith('acero'):
            vals.append(2_040_000.0)
        elif tl in ('general', 'concreto', '-'):
            vals.append(E)
        else:
            vals.append(_a_kgcm2(parse_num(t, f'{nom} - E de columnas'), d.get('unidad_E'), f'{nom} - E de columnas'))
    if len(vals) == 1:
        return vals * ncol
    if len(vals) != ncol:
        raise ValueError(f'{nom}: tiene {ncol} columna(s) y escribiste {len(vals)} valores de E. '
                         'Escribe uno solo o uno por columna.')
    return vals


def _E_viga(r, d, E, nom):
    """E de las vigas de un piso (solo importa con vigas flexibles): vacío = E general."""
    txt = str(r.get('E_viga', '')).strip()
    if not txt:
        return E
    if txt.lower().startswith('acero'):
        return 2_040_000.0
    return _a_kgcm2(parse_num(txt, f'{nom} - E de la viga'), d.get('unidad_E'), f'{nom} - E de la viga')


def _tramos(texto, nb, nom):
    """Tramos (crujías 1..nb, de izquierda a derecha) sin viga ni losa en un piso: "3", "1 3", "todos"."""
    try:
        return parse_ejes(texto, nb, f'{nom} - tramos sin losa', [])
    except ValueError as e:
        raise ValueError(str(e).replace('ejes', 'tramos').replace('eje', 'tramo'))


APOYOS = ('Empotrado', 'Articulado', 'Rodillo')


def estados_apoyo(artic, rod, nejes):
    """Apoyo de cada eje en la base ('Empotrado' por defecto) a partir de las listas de ejes articulados y con rodillos."""
    art = set(parse_ejes(artic, nejes, 'apoyos articulados', []))
    rodillos = set(parse_ejes(rod, nejes, 'apoyos con rodillos', []))
    return ['Rodillo' if e in rodillos else ('Articulado' if e in art else 'Empotrado') for e in range(1, nejes + 1)]


def apoyos_a_texto(estados):
    """Inverso de estados_apoyo: ejes articulados y ejes con rodillos (listas de enteros)."""
    for e in estados:
        if e not in APOYOS:
            raise ValueError(f'Apoyo no válido: {e}. Usa {", ".join(APOYOS)}.')
    return ([i for i, e in enumerate(estados, 1) if e == 'Articulado'], [i for i, e in enumerate(estados, 1) if e == 'Rodillo'])


def propiedades_miembros(texto, E, unidad_E='kg/cm²', direccion='X', unidad_seccion='cm', h_m=None):
    """Calculadora de EI: propiedades de una o varias secciones con un E dado.

    texto: secciones con la misma sintaxis de la tabla de niveles (25x50, Ø60, I:3688/A:26.8/I2:98, H:..., T:...).
    Devuelve una fila por sección con I (en la dirección del análisis), A, radio de giro, EI y, si se da la altura,
    las rigideces laterales 12EI/h³ (empotrada-empotrada) y 3EI/h³ (articulada en un extremo)."""
    esc = {'cm': 1.0, 'm': 100.0}.get(str(unidad_seccion).strip().lower())
    if esc is None:
        raise ValueError('Las unidades de las secciones deben ser cm o m.')
    secs = parse_secciones(_escalar_texto(texto, esc, 'Calculadora'), 'Calculadora')
    E_kg = _a_kgcm2(parse_num(E, 'E'), unidad_E, 'E')
    if E_kg <= 0:
        raise ValueError('E debe ser mayor que cero.')
    otra = 'Y' if direccion == 'X' else 'X'
    filas = []
    for sec in secs:
        I = inercia(sec, direccion)
        A = sec['A'] if sec['A'] < 1e8 else None            # sin área capturada
        EI = E_kg * I
        fila = dict(seccion=sec['texto'], E=E_kg, I=I, I_otra=inercia(sec, otra), A=A,
                    r=float(np.sqrt(I / A)) if A else None, EI=EI, EI_tcm2=EI / 1000.0, EI_tm2=EI / 1e7, EA=E_kg * A if A else None)
        if h_m:
            hcm = h_m * 100
            fila.update(k12=12 * EI / hcm ** 3, k3=3 * EI / hcm ** 3, k12_tcm=12 * EI / hcm ** 3 / 1000.0,
                        k3_tcm=3 * EI / hcm ** 3 / 1000.0)
        filas.append(fila)
    return filas


def preparar_modelo(d):
    """Arma K, M y la geometría a partir de ejes y niveles.

    d['crujias']: anchos de crujía (m) de izquierda a derecha, p. ej. "4 5" -> ejes 1, 2 y 3.
    d['volado_izq'], d['volado_der']: volados (m) más allá del primer y del último eje; "cargas" admite entonces un valor
        por tramo de izquierda a derecha (volado izq., crujías, volado der.), con 0 donde no hay losa.
    d['direccion']: 'X' (por defecto) o 'Y': dirección del sismo. En una sección b x h, h va en X y b en Y.
    d['vigas']: 'rigidas' (marco de cortante) o 'flexibles' (marco plano con vigas y columnas).
    d['niveles'] (de abajo hacia arriba), dicts con:
        h       altura de entrepiso (m)
        ejes    ejes con columna: "todos" (por defecto), "1 3", "1-3"
        seccion una sola sección para todas las columnas o una por columna (ver ame_secciones)
        artic   ejes con base articulada (opcional): "2"
        rod     ejes con apoyo de rodillos en la base (solo nivel 1): no transmiten cortante, rigidez lateral = 0
        art_sup ejes con la columna articulada en su parte superior (opcional)
        viga    sección de la viga de ese piso (solo con vigas flexibles): "30x60"
        cargas  carga por crujía (t/m): un valor para todas o uno por crujía, p. ej. "4.7 3.2"
        E_col   E de las columnas del nivel (en la unidad de E; o 'acero'): un valor para todas o uno por columna;
                vacío = el E general
        E_viga  E de las vigas de ese piso (vigas flexibles); vacío = el E general
        sin_losa tramos (crujías 1, 2, … de izquierda a derecha) sin viga ni losa en ese piso: su carga se ignora; si un
                eje queda sin viga en ese piso, su columna lo cruza (una sola columna de varios entrepisos)
        W       peso del nivel (t), opcional: sustituye a la carga
    d['T_medido']: periodo fundamental medido (s), opcional. Los pesos pasan a ser solo la distribución relativa entre
        niveles (si se omiten, todos iguales) y se escalan para que T1 sea el medido: así se obtiene el peso total.
    """
    bays = parse_vector(d['crujias'], 'Crujías')
    if (bays <= 0).any():
        raise ValueError('Las crujías deben ser mayores que cero.')
    nb = len(bays)
    nejes = nb + 1
    xs = np.concatenate([[0.0], np.cumsum(bays)])
    vi = _num_opc(d, 'volado_izq', None) or 0.0      # volados (voladizos) más allá del primer y último eje
    vd = _num_opc(d, 'volado_der', None) or 0.0
    if vi < 0 or vd < 0:
        raise ValueError('Los volados no pueden ser negativos.')
    spans = ([(-vi, 0.0)] if vi > 0 else []) + [(float(xs[j]), float(xs[j + 1])) for j in range(nb)] \
        + ([(float(xs[-1]), float(xs[-1] + vd))] if vd > 0 else [])
    lon_spans = np.array([b_ - a_ for a_, b_ in spans])
    direccion = (str(d.get('direccion') or 'X').strip().upper() or 'X')[0]
    if direccion not in 'XY':
        raise ValueError('La dirección del sismo debe ser X o Y.')
    flexibles = str(d.get('vigas') or '').lower().startswith('flex')
    niv = d['niveles']
    n = len(niv)
    if n < 1:
        raise ValueError('Agrega al menos un nivel.')
    E = _modulo_E(d)
    T_med = _periodo_medido(d)
    esc_sec = {'cm': 1.0, 'm': 100.0}.get(str(d.get('unidad_seccion') or 'cm').strip().lower())
    if esc_sec is None:
        raise ValueError('Las unidades de las secciones deben ser cm o m.')

    filas, detalle, geom_niv, ks, m, W, alturas = [], [], [], [], [], [], []
    notas = []
    for i, r in enumerate(niv, 1):
        nom = f'Nivel {i}'
        h = parse_num(r.get('h', ''), f'{nom} - h')
        if h <= 0:
            raise ValueError(f'{nom}: la altura debe ser mayor que cero.')
        ejes = parse_ejes(r.get('ejes', ''), nejes, f'{nom} - ejes', range(1, nejes + 1))
        secs = parse_secciones(_escalar_texto(r.get('seccion', ''), esc_sec, f'{nom} - sección'), f'{nom} - sección')
        if len(secs) == 1:
            secs = secs * len(ejes)
        elif len(secs) != len(ejes):
            raise ValueError(f'{nom}: tiene {len(ejes)} columna(s) (ejes {", ".join(map(str, ejes))}) y escribiste '
                             f'{len(secs)} secciones. Escribe una sola o una por columna.')
        artic = parse_ejes(r.get('artic', ''), nejes, f'{nom} - articuladas', [])
        rod = parse_ejes(r.get('rod', ''), nejes, f'{nom} - rodillos', [])
        art_sup = parse_ejes(r.get('art_sup', ''), nejes, f'{nom} - articuladas arriba', [])
        if rod and i != 1:
            raise ValueError(f'{nom}: los apoyos de rodillos solo aplican en la base (nivel 1).')
        for lista, et in ((artic, 'articulado'), (rod, 'con rodillos'), (art_sup, 'articulado arriba')):
            fuera = [e for e in lista if e not in ejes]
            if fuera:
                raise ValueError(f'{nom}: el eje {fuera[0]} está marcado como {et} pero no tiene columna en ese nivel.')
        doble = sorted(set(artic) & set(rod))
        if doble:
            raise ValueError(f'{nom}: el eje {doble[0]} está marcado como articulado y como rodillo; elige uno.')

        hcm = h * 100
        E_cols = _E_columnas(r, d, E, nom, len(ejes))
        E_v = _E_viga(r, d, E, nom)
        sin = [j - 1 for j in _tramos(r.get('sin_losa', ''), nb, nom)]          # tramos sin viga (0-based)
        cols, kn = [], 0.0
        for (e, sec), E_n in zip(zip(ejes, secs), E_cols):
            art, roll, arriba = e in artic, e in rod, e in art_sup
            I = inercia(sec, direccion)
            if roll or (art and arriba):  # sin restricción horizontal o articulada en ambos extremos: no resiste cortante
                k = 0.0
            else:
                k = (3 if (art or arriba) else 12) * E_n * I / hcm ** 3
            kn += k
            cols.append(dict(sec, nivel=i, eje=e, x=float(xs[e - 1]), artic=art, artic_sup=arriba, rodillo=roll, k=k, I=I, E=E_n,
                             eq=ancho_en_direccion(sec, direccion)))
        detalle.append(cols)

        viga = dict(I=None, A=None, texto='', alto=30.0, E=E_v)
        if flexibles:
            vs = parse_secciones(_escalar_texto(r.get('viga', ''), esc_sec, f'{nom} - viga'), f'{nom} - viga')
            if len(vs) != 1:
                raise ValueError(f'{nom} - viga: escribe una sola sección de viga (por ejemplo 30x60).')
            viga = dict(I=vs[0]['Ix'], A=vs[0]['A'], texto=vs[0]['texto'], alto=vs[0]['dx'], E=E_v)

        if str(r.get('W', '')).strip():
            Wi, w_bays = parse_num(r['W'], f'{nom} - W'), None
        elif T_med and not str(r.get('cargas', '')).strip():
            Wi, w_bays = 1.0, None  # peso relativo: se escala para igualar el periodo medido
        else:
            v = parse_vector(r.get('cargas', ''), f'{nom} - carga por crujía')
            if len(v) == 1:
                v = np.repeat(v, len(spans))
            elif len(v) != len(spans):
                desc = ('volado izq. + ' if vi > 0 else '') + f'{nb} crujía(s)' + (' + volado der.' if vd > 0 else '')
                raise ValueError(f'{nom}: hay {len(spans)} tramo(s) con carga ({desc}, de izquierda a derecha) y escribiste '
                                 f'{len(v)} cargas. Escribe una sola o una por tramo (0 si no hay carga).')
            v = np.array(v, float)
            for j in sin:                                   # sin losa: no hay carga en ese tramo
                jj = j + (1 if vi > 0 else 0)
                if v[jj] != 0:
                    notas.append(f'{nom}: el tramo {j + 1} no tiene losa; se ignoró su carga de {v[jj]:g} t/m.')
                v[jj] = 0.0
            Wi, w_bays = float(v @ lon_spans), v
        ks.append(kn); alturas.append(h); W.append(Wi); m.append(Wi * 1000 / G)
        textos = list(dict.fromkeys(c['texto'] for c in cols))
        filas.append(dict(nivel=i, h=h, ncol=len(cols), nart=len(artic) + len(rod), seccion=', '.join(textos),
                          I=sum(c['I'] for c in cols), k=kn, W=Wi, m=Wi * 1000 / G,
                          viga=viga['texto'] if flexibles else ''))
        w_spans = None if w_bays is None else [(a_, b_, float(w_)) for (a_, b_), w_ in zip(spans, w_bays)]
        ext = [spans[0][0] if vi > 0 else float(xs[0]), spans[-1][1] if vd > 0 else float(xs[-1])]
        if w_spans is not None:  # un volado sin carga en este nivel se dibuja sin losa
            if vi > 0 and w_spans[0][2] == 0:
                ext[0] = float(xs[0])
            if vd > 0 and w_spans[-1][2] == 0:
                ext[1] = float(xs[-1])
        segs = []                                           # tramos de losa/viga que se dibujan (y unen los ejes)

        def _unir(a_, b_):
            if segs and abs(segs[-1][1] - a_) < 1e-9:
                segs[-1] = (segs[-1][0], b_)
            else:
                segs.append((a_, b_))
        conectado = set()                                   # ejes cuyo nodo está unido a la losa de este piso
        if vi > 0 and (w_spans is None or w_spans[0][2] != 0):
            _unir(-vi, 0.0)
            conectado.add(1)
        for j in range(nb):
            if j not in sin:
                _unir(float(xs[j]), float(xs[j + 1]))
                conectado |= {j + 1, j + 2}
        if vd > 0 and (w_spans is None or w_spans[-1][2] != 0):
            _unir(float(xs[-1]), float(xs[-1] + vd))
            conectado.add(nejes)
        geom_niv.append(dict(h=h, W=Wi, w_bays=w_bays, w_spans=w_spans, ext=ext, cols=cols, viga=viga, sin=sin, segs=segs,
                             conectado=sorted(conectado), viga_h=viga['alto'] if flexibles else 30.0))

    # columnas que cruzan un piso sin viga ni losa en su eje: se vuelven una sola columna de varios entrepisos
    for k in range(n - 1):
        cn = set(geom_niv[k]['conectado'])
        for e in range(1, nejes + 1):
            if e in cn:
                continue
            cb = next((c for c in detalle[k] if c['eje'] == e), None)
            ca = next((c for c in detalle[k + 1] if c['eje'] == e), None)
            if cb and ca:
                cb['cont_sup'], ca['cont_inf'] = True, True
            elif cb or ca:
                raise ValueError(f'Eje {e}: en el piso del nivel {k + 1} no hay viga ni losa y la columna no continúa '
                                 f'{"hacia arriba" if cb else "hacia abajo"}: quedaría sin arriostrar. '
                                 'Agrega la columna del otro nivel o la viga.')
    for i in range(n):
        for c in detalle[i]:
            if not c.get('cont_sup') or c.get('cont_inf'):
                continue
            j, ultimo = i, c
            while ultimo.get('cont_sup'):
                j += 1
                ultimo = next(x for x in detalle[j] if x['eje'] == c['eje'])
                ultimo['k_cont'], ultimo['k'] = True, 0.0
            L_cm = sum(alturas[i:j + 1]) * 100
            c['hasta'], c['fin_nivel'], c['artic_sup'] = j + 1, j, ultimo['artic_sup']
            if c['rodillo'] or (c['artic'] and c['artic_sup']):
                c['k'] = 0.0
            else:
                c['k'] = (3 if (c['artic'] or c['artic_sup']) else 12) * c['E'] * c['I'] / L_cm ** 3
            notas.append(f'Eje {c["eje"]}: no hay viga ni losa en el piso del nivel {i + 1}'
                         + (f' (hasta el {j})' if j > i + 1 else '') + f'; la columna cruza esos pisos como una sola de '
                         f'{L_cm / 100:g} m (secc. {c["texto"]}, E = {c["E"]:,.0f} kg/cm² del nivel {i + 1}).')
            if ultimo['texto'] != c['texto'] or abs(ultimo['E'] - c['E']) > 1e-6:
                notas.append(f'Eje {c["eje"]}: la columna que cruza tiene secciones o E distintos en cada nivel; '
                             f'se usó el del nivel {i + 1}.')
    for i, f in enumerate(filas):
        f['k'] = sum(c['k'] for c in detalle[i])

    if flexibles:
        K = rigidez_marco(E, [x * 100 for x in alturas], xs * 100,
                          [[dict(eje=c['eje'], A=c['A'], I=c['I'], E=c['E'], artic=c['artic'], artic_sup=c['artic_sup'],
                                 rodillo=c['rodillo'], hasta=c.get('hasta', i + 1), omitir=c.get('k_cont', False))
                            for c in cols] for i, cols in enumerate(detalle)],
                          [dict(I=l['viga']['I'], A=l['viga']['A'], E=l['viga']['E'], sin=set(l['sin'])) for l in geom_niv])
    else:
        K = np.zeros((n, n))
        for i, cols in enumerate(detalle):
            for c in cols:
                if c.get('k_cont'):
                    continue
                a_, b_ = i, c.get('hasta', i + 1)          # pisos (0 = suelo)
                k_ = c['k']
                if a_ >= 1:
                    K[a_ - 1, a_ - 1] += k_
                    K[a_ - 1, b_ - 1] -= k_
                    K[b_ - 1, a_ - 1] -= k_
                K[b_ - 1, b_ - 1] += k_
    mat = d.get('material') or 'Concreto'
    obj = _objetivo_de(d)
    modelo = dict(M=np.diag(m), K=K, n=n, alturas=np.array(alturas), cargas=np.array(W), dist_x=np.ones(n),
                  grupo=obj['grupo'], sismo=obj['sismo'], subgrupo=obj['subgrupo'], intensidad=obj['intensidad'],
                  objetivo=obj['objetivo'], E=E, info=filas, material=mat, direccion=direccion,
                  vigas='flexibles' if flexibles else 'rigidas', R_legacy=False,
                  geom=dict(bays=bays, xs=xs, niveles=geom_niv, volados=(vi, vd), spans=spans), notas_modelo=notas,
                  **_params_norma(d, mat), **_params_espectro(d, obj['estado_auto'], mat, obj['intensidad']))
    return _ajustar_periodo(modelo, T_med) if T_med else modelo


def convertir_legacy(d):
    """Convierte un proyecto del formato anterior (ncol, b, hs, nart, w, L) al formato por ejes."""
    niv = d.get('niveles')
    if 'crujias' in d or not niv or 'ncol' not in niv[0]:
        return d
    ncol_max = max(int(float(r['ncol'])) for r in niv)
    nb = max(ncol_max - 1, 1)
    L = next((float(r['L']) for r in niv if str(r.get('L', '')).strip()), 6.0)
    nuevo = dict(d, crujias=' '.join([f'{L / nb:g}'] * nb), niveles=[])
    for r in niv:
        nc = int(float(r['ncol']))
        na = int(float(r.get('nart') or 0))
        ejes = list(range(1, nc + 1))
        nuevo['niveles'].append(dict(
            h=r['h'], ejes=' '.join(map(str, ejes)), seccion=f"{r['b']}x{r['hs']}",
            artic=' '.join(map(str, ejes[nc - na:])) if na else '', cargas=r.get('w', ''), W=r.get('W', '')))
    return nuevo


# ----------------------------------------------------------------------------
# Cálculo
# ----------------------------------------------------------------------------
def _r0(Q):
    return 1.75 if Q < 3 else 2.0


def _frp(p, clave='SV'):
    """R′/R: 0.75 en ocupación inmediata y 1 en seguridad de vida (apuntes, p. 10).

    En proyectos del formato anterior (K y M a mano) se conserva la regla vieja: grupo A con sismo base -> 0.75."""
    if clave == 'OI':
        return ntc.FACTOR_R_OI
    if p.get('R_legacy', True) and p['grupo'] == 'A' and p['sismo'] == 'B':
        return ntc.FACTOR_R_OI
    return 1.0


def _R(p, Q, T):
    """Sobrerresistencia R = k1·R0 + k2 (R = 1 en materiales distintos del concreto, apuntes p. 11)."""
    T = np.asarray(T, float)
    if p.get('R_unitaria'):
        return np.ones_like(T)
    return p['k1'] * _r0(Q) + _k2(T, p['Ta'])


def _R_dist(p, Q, T1):
    """R para amplificar los desplazamientos: la del plato (k2 = 0, como en las hojas) o la del periodo fundamental."""
    return _R(p, Q, p['Ta'] if p.get('R_dist', 'plato') == 'plato' else T1)


def _m_ef(p, fi):
    """Masa efectiva modal como fracción de la masa total."""
    M, n = p['M'], p['n']
    return np.array([(fi[:, j] @ M @ np.ones(n)) ** 2 / (fi[:, j] @ M @ fi[:, j]) for j in range(n)]) / np.trace(M)


def _n_modos(p, res):
    """Modos que se combinan: todos, los primeros N o el criterio de los apuntes (95 % de masa o T ≥ 0.4 s, mínimo 3)."""
    n = p['n']
    modo = p.get('modos_modo', 'todos')
    if modo == 'primeros':
        return int(min(max(p.get('modos_n') or 1, 1), n))
    if modo == 'apuntes':
        acum = np.cumsum(_m_ef(p, res['fi']))
        n95 = int(np.argmax(acum >= 0.95)) + 1 if acum[-1] >= 0.95 else n
        return int(min(n, max(n95, int(np.sum(res['T'] >= 0.4)), min(3, n))))
    return n


def _comb(res, X):
    """Combina las respuestas modales (columnas) solo con los modos considerados."""
    p, m = res['p'], res['n_modos']
    return combinar(X[:, :m], res['w'][:m], p['zeta'], p['combinacion'])


def _k2(T, Ta):
    return np.maximum(0.5 * (1 - np.sqrt(np.asarray(T, float) / Ta)), 0)


def _a_minima(p, Rprima_prom, esc):
    """a_min del espectro; esc=1 -> cm/s^2, esc=1/981 -> g."""
    Ts = p['Ts']
    a1, a2 = 0.04 / Rprima_prom, 0.06 / Rprima_prom
    if Ts < 0.5:
        a = a1
    elif Ts < 1:
        a = a1 + (a2 - a1) * (Ts - 0.5) / 0.5
    else:
        a = a2
    return a * G * esc


def espectro(p, Q, esc, serie, dl=False, clave='SV'):
    """Sa/(Q'R') en la malla 'serie' (esc=1 cm/s^2; esc=1/981 g). Devuelve (curva, a_minima)."""
    a0, c, Ta, Tb, k = p['a0'] * esc, p['c'] * esc, p['Ta'], p['Tb'], p['k']
    Sa = np.zeros_like(serie)
    q = np.zeros_like(serie)
    i1 = serie < Ta
    i2 = (serie >= Ta) & (serie < Tb)
    i3 = serie >= Tb
    Sa[i1] = a0 + (c - a0) * serie[i1] / Ta
    q[i1] = 1 + (Q - 1) * np.sqrt(1 / k) * serie[i1] / Ta
    Sa[i2] = c
    q[i2] = 1 + (Q - 1) * np.sqrt(1 / k)
    pp = k + (1 - k) * (Tb / serie[i3]) ** 2
    Sa[i3] = c * pp * (Tb / serie[i3]) ** 2
    q[i3] = 1 + (Q - 1) * np.sqrt(np.maximum(pp / k, 0))
    if dl:  # limitación de daños: espectro elástico por Ks, sin Q' ni R'
        return Sa * p['Ks'], 0.0

    f = _frp(p, clave)
    # R' promedio (igual que el script original: promedio sobre toda la malla, con r0 de Q(1))
    Rprom = f * np.mean(_R(p, p['Q'][0], serie))
    a_min = _a_minima(p, Rprom, esc)
    r = f * _R(p, Q, serie)

    curva = Sa / (q * r)
    idx = np.flatnonzero(curva <= a_min)
    if idx.size:
        curva[idx[0]:] = a_min
    return curva, a_min


def _q_r(p, Q, T, clave='SV'):
    """Q' y R' del espectro de diseño para los periodos T."""
    T = np.maximum(np.atleast_1d(np.asarray(T, float)), 1e-9)
    k, Ta, Tb = p['k'], p['Ta'], p['Tb']
    pp = k + (1 - k) * (Tb / T) ** 2
    q = np.where(T < Ta, 1 + (Q - 1) * np.sqrt(1 / k) * T / Ta,
                 np.where(T < Tb, 1 + (Q - 1) * np.sqrt(1 / k), 1 + (Q - 1) * np.sqrt(np.maximum(pp / k, 0))))
    return q, _frp(p, clave) * _R(p, Q, T)


def _interp(x, xp, fp):
    """Interpolación lineal con extrapolación lineal."""
    x = np.atleast_1d(x)
    y = np.interp(x, xp, fp)
    lo, hi = x < xp[0], x > xp[-1]
    if lo.any():
        y[lo] = fp[0] + (fp[1] - fp[0]) / (xp[1] - xp[0]) * (x[lo] - xp[0])
    if hi.any():
        y[hi] = fp[-1] + (fp[-1] - fp[-2]) / (xp[-1] - xp[-2]) * (x[hi] - xp[-1])
    return y


def _revisiones(p, res):
    """Distorsiones en los tres estados límite, P-Δ e irregularidades (todo de abajo hacia arriba)."""
    n, T, fi, lam, gam, w, K = p['n'], res['T'], res['fi'], res['lam'], res['gam'], res['w'], p['K']
    h_cm = p['alturas'] * 100
    FC, zeta, comb = p['factor_Fu'], p['zeta'], p['combinacion']
    estados = {}
    for clave, nombre, limite in (('DL', 'Limitación de daños', p['lim_dl']),
                                  ('SV', 'Seguridad de vida', p['lim_sv']),
                                  ('OI', 'Ocupación inmediata', p['lim_oi'])):
        gc = 1.0 if clave == 'DL' else p['gamma_c']  # γc no aplica a limitación de daños (apuntes p. 12)
        if clave == 'DL':
            curva, _ = espectro(p, 1.0, 1.0, res['serie'], dl=True)
            amp = np.ones(n)  # fuerzas sin FC y sin Q·R′
        else:
            pq = p if clave == 'SV' else dict(p, Q=np.array([1.0]))
            Q = float(pq['Q'][0])
            curva, _ = espectro(pq, Q, 1.0, res['serie'], clave=clave)
            R1 = float(_frp(p, clave) * _R_dist(pq, Q, T[0]))  # R′ del plato (hojas) o del periodo fundamental
            amp = np.full(n, FC * Q * R1)  # desplazamientos × FC·Q·R′ (apuntes p. 13 y 15)
        Sa = _interp(T, res['serie'], curva)
        d = fi * (Sa * gam / lam)
        deriva = _comb(res, derivas_modales(d * amp))
        estados[clave] = dict(nombre=nombre, limite=None if limite is None else limite * gc, limite_base=limite,
                              gamma_c=gc, Sa=Sa, amp=amp, d=d, deriva=deriva,
                              deriva_el=_comb(res, derivas_modales(d)), dist=deriva / h_cm)
    rev = dict(estados=estados, theta=np.zeros(n), P=None, V=None, irreg_masa=[], irreg_rigidez=[],
               k_ef=np.zeros(n))
    W = p['cargas']
    if W is not None:
        sv = estados['SV']
        F_t = _comb(res, K @ sv['d']) / 1000
        V = np.cumsum(F_t[::-1])[::-1]
        rev['theta'], rev['P'] = theta_pdelta(W, sv['deriva_el'], V, h_cm)
        rev['V'] = V
        rev['irreg_masa'], rev['irreg_rigidez'], rev['k_ef'] = irregularidades(W, V, sv['deriva_el'])
    return rev


def _normativa(p, res):
    """Factores aplicados, γc, cortante basal mínimo, cimentación, torsión accidental y método estático."""
    n, T, clave, FC = p['n'], res['T'], res['clave'], p['factor_Fu']
    T1 = float(T[0])
    pp = p if clave == 'SV' else dict(p, Q=np.array([1.0]))
    Q0 = float(pp['Q'][0])
    obj = p.get('objetivo')
    out = dict(clave=clave, T1=T1, objetivo=obj)
    filas = []

    def fila(nombre, valor, fuente=''):
        filas.append((nombre, valor, fuente))

    sub = p.get('subgrupo') or p['grupo']
    asce = ntc.RIESGO_ASCE.get(sub)
    fila('Grupo / subgrupo', f'{p["grupo"]} / {sub}' + (f' (ASCE 7 categoría {asce})' if asce else ''), 'Apuntes p. 10')
    fila('Intensidad sísmica', p['intensidad'], 'SASID')
    if obj:
        fila('Nivel de desempeño', obj['desempeno'], 'Tabla 1.1a')
        fila('Espectro de diseño', obj['espectro'], 'Tabla 3.1.1')
    fila('Estado de cálculo', p['estado'], '')
    fila('Amortiguamiento ζ', f'{p["zeta"]:g}', 'Concreto 0.05 · acero 0.03 (p. 14)')
    if clave == 'DL':
        fila('Ks', f'{p["Ks_base"]:g}' + (f' × {ntc.FACTOR_KS_OCASIONAL:g} (sismo ocasional) = {p["Ks"]:g}'
                                          if p['Ks'] != p['Ks_base'] else ''), 'Sa(BD)·Ks, sin factor de carga')
    else:
        r0 = _r0(Q0)
        plato = p.get('R_dist', 'plato') == 'plato'
        k2 = 0.0 if (p['R_unitaria'] or plato) else float(_k2(T1, p['Ta']))
        R_ = float(_R_dist(pp, Q0, T1))
        frp = _frp(p, clave)
        q1 = float(_q_r(pp, Q0, T1, clave)[0][0])
        fila('Q', f'{Q0:g}', ntc.GUIA_Q.get(Q0, 'Q fuera de los valores usuales'))
        fila('k1', f'{p["k1"]:g}', 'Hiperestaticidad: 0.8 / 1.0 / 1.25 (p. 23)')
        if p['R_unitaria']:
            fila('R = k1·R0 + k2', '1', 'R = 1 en materiales distintos del concreto (p. 11)')
        else:
            fila('R0', f'{r0:g}', '1.75 si Q < 3 · 2 si Q ≥ 3')
            fila('k2 ' + ('(plato)' if plato else f'(T1 = {T1:.4f} s)'), f'{k2:.4f}', 'k2 = 0.5·(1 − √(T/Ta)) ≥ 0')
            fila('R = k1·R0 + k2', f'{R_:.4f}', 'Del plato, como en las hojas' if plato else 'Con el periodo fundamental (p. 12)')
        fila("R′ = " + ('0.75·R' if frp != 1 else 'R'), f'{frp * R_:.4f}', 'OI: 0.75R · SV: R (p. 10)')
        fila("Q′ (T1)", f'{q1:.4f}', 'Q′ = 1 + (Q − 1)·√(p/k)…')
        fila("Q′·R′ (T1)", f'{q1 * frp * R_:.4f}', 'Reduce las ordenadas del espectro')
        fila("FC·Q·R′ (distorsiones)", f'{FC * Q0 * frp * R_:.4f}', 'Desplazamientos × FC·Q·R′ (p. 13, 15)')
    fila('γc por irregularidad', f'{p["gamma_c"]:g}', p['gamma_c_txt'])
    out['factores'] = filas

    # --- cortante basal mínimo: V ≥ FC·a_min·W (apuntes p. 10, inciso 7.5)
    W = p['cargas']
    if W is not None and res['a_min_g']:
        W_T = float(np.sum(W))
        v_min = FC * res['a_min_g'] * W_T
        out['vmin'] = dict(a_min_g=res['a_min_g'], W_T=W_T, V_min=v_min, V_din=res['V_din'],
                           cumple=res['V_din'] >= v_min * (1 - 1e-9), factor=max(1.0, v_min / res['V_din']))
    # --- cimentación: elementos mecánicos × 0.65·R′ (p. 9, inciso 3)
    if clave != 'DL':
        R1 = float(_frp(p, clave) * _R_dist(pp, Q0, T1))
        f_c = ntc.FACTOR_CIMENTACION * R1
        out['cimentacion'] = dict(R_prima=R1, factor=f_c, V=f_c * res['V_din'], M=f_c * float(res['tabla_din']['Mvu'].sum()))
    # --- torsión accidental y de diseño (p. 9): e = f·e_s ± e_a  (f = 1.0 dinámico, 1.5 estático)
    if p['b_planta']:
        e_a = ntc.excentricidad_accidental(n, p['b_planta'])
        Fu_arriba = res['tabla_din']['Fu']  # de arriba hacia abajo
        tor = dict(b=p['b_planta'], e_a=e_a, Mt=e_a * Fu_arriba[::-1], Fu=Fu_arriba[::-1])
        if p.get('e_s') is not None:
            es = p['e_s']
            tor.update(e_s=es, e1_din=1.0 * es + e_a, e2_din=es - e_a, Mt1_din=(es + e_a) * Fu_arriba[::-1],
                       Mt2_din=(es - e_a) * Fu_arriba[::-1])
            if 'tabla_est' in res:
                Fe = res['tabla_est']['Fu'][::-1]
                tor.update(e1_est=1.5 * es + e_a, e2_est=es - e_a, Mt1_est=(1.5 * es + e_a) * Fe, Mt2_est=(es - e_a) * Fe)
        out['torsion'] = tor
    # --- aplicabilidad del método estático (p. 20)
    H = float(np.sum(p['alturas']))
    lim = ntc.altura_maxima_estatico(p['zona'], p['regular']) if p['zona'] else None
    out['estatico'] = dict(H=H, limite=lim, zona=p['zona'], regular=p['regular'],
                           aplica=p['grupo'] == 'B' and (lim is None or H <= lim))
    # --- participación modal (p. 15): más de 3 modos, 95 % de masa efectiva o todos los modos con T ≥ 0.4 s
    m_ef = _m_ef(p, res['fi'])
    acum = np.cumsum(m_ef)
    out['modos'] = dict(m_ef=m_ef, acum=acum, n95=int(np.argmax(acum >= 0.95)) + 1 if acum[-1] >= 0.95 else n,
                        n_T04=int(np.sum(T >= 0.4)), n_usados=res['n_modos'], acum_usados=float(acum[res['n_modos'] - 1]))
    return out


def _estatico(p, res):
    """Método estático de análisis sísmico (apuntes p. 20-21).

    Fi = C · (Wi·hi / Σ Wi·hi) · ΣWi, con C = máx { c/(Q′R′) en el plato ; a0/(Q′R′) en T = 0 }.
    En el plato k2 = 0 y Q′ = 1 + (Q − 1)·√(1/k); en T = 0, k2 = 0.5 y Q′ = 1.
    Distorsiones: desplazamientos de K⁻¹·Fi × FC·Q·R′ del plato (limitación de daños: sin FC ni Q·R′).
    """
    n, clave = p['n'], res['clave']
    FC = 1.0 if clave == 'DL' else p['factor_Fu']
    h = res['alturas_acum'][::-1]  # de arriba hacia abajo
    W = np.asarray(p['cargas'] * p['dist_x'], float)[::-1]
    hW = h * W
    c_g, a0_g = p['c'] / G, p['a0'] / G
    Q = float(p['Q'][0]) if clave == 'SV' else 1.0
    frp = 1.0 if clave == 'DL' else _frp(p, clave)
    q_pl = 1.0 if clave == 'DL' else 1 + (Q - 1) * np.sqrt(1 / p['k'])
    R_pl = 1.0 if (clave == 'DL' or p['R_unitaria']) else p['k1'] * _r0(Q)           # k2 = 0 en el plato
    R_0 = 1.0 if (clave == 'DL' or p['R_unitaria']) else p['k1'] * _r0(Q) + 0.5      # k2 = 0.5 en T = 0
    if clave == 'DL':  # espectro elástico por Ks, sin Q′ ni R′
        C_pl, C_0 = c_g * p['Ks'], a0_g * p['Ks']
    else:
        C_pl, C_0 = c_g / (q_pl * frp * R_pl), a0_g / (frp * R_0)
    if p.get('criterio_est') == 'maximo':
        C, gob = float(res['curvas_g'][0].max()), 'máximo del espectro de diseño (como el código de MATLAB)'
    else:
        C, gob = max(C_pl, C_0), ('el plato' if C_pl >= C_0 else 'el inicio del espectro (T = 0)')
    a_min_g = res['a_min_g'] or 0.0
    if C < a_min_g:
        C, gob = a_min_g, 'el cortante basal mínimo (a_min)'
    Fi = C * hW / hW.sum() * W.sum()
    Fu = Fi * FC
    Vu = np.cumsum(Fu)
    Mvu = Fu * h

    u = np.linalg.solve(p['K'], Fi[::-1] * 1000)  # cm, de abajo hacia arriba (Fi en t -> kg)
    deriva = np.diff(np.concatenate([[0.0], u]))
    amp = 1.0 if clave == 'DL' else FC * Q * frp * R_pl
    dist = deriva * amp / (p['alturas'] * 100)
    lim_base = {'DL': p['lim_dl'], 'SV': p['lim_sv'], 'OI': p['lim_oi']}[clave]
    gc = 1.0 if clave == 'DL' else p['gamma_c']
    return dict(tabla=dict(nivel=np.arange(n, 0, -1), h=h, Wi=W, hWi=hW, Fi=Fi, Fu=Fu, Vu=Vu, Mvu=Mvu,
                           Mv=_volteo_entrepiso(Fu, p['alturas'])),
                C=C, C_pl=C_pl, C_0=C_0, gobierna=gob, criterio=p.get('criterio_est', 'apuntes'), Q=Q, FC=FC,
                q_pl=q_pl, R_pl=R_pl, R_0=R_0, frp=frp, Rp_pl=frp * R_pl, Rp_0=frp * R_0, W_T=float(W.sum()),
                V_est=float(Vu[-1]), u=u, deriva=deriva, amp=amp, dist=dist,
                limite=None if lim_base is None else lim_base * gc, limite_base=lim_base, gamma_c=gc)


def _completar_estatico(p, res):
    """Compara con el dinámico y escala las fuerzas dinámicas (factor = V estático / V dinámico, si es mayor que 1)."""
    est = res['estatico']
    est['dist_din'] = res['revision']['estados'][res['clave']]['dist']
    f = max(1.0, res['factor_escala'])
    t = res['tabla_din']
    res['factor_aplicado'] = f
    res['tabla_din_esc'] = dict(nivel=t['nivel'], h=t['h'], F=t['F'] * f, Fu=t['Fu'] * f, Vu=t['Vu'] * f, Mvu=t['Mvu'] * f,
                                Mv=t['Mv'] * f)


def _elastico(p, res):
    """Análisis elástico: el espectro SIN reducir (Q′ = 1, R′ = 1), el "espectro transparente" de los apuntes (p. 11).

    Se hace con el mismo modelo, la misma combinación modal y los mismos modos que el análisis de diseño, y se compara
    con él. Las fuerzas elásticas son las que resistiría una estructura que no se reduce por ductilidad ni sobrerresistencia."""
    n, K = p['n'], p['K']
    T, fi, lam, gam, serie = res['T'], res['fi'], res['lam'], res['gam'], res['serie']
    FC = p['factor_Fu']
    curva = espectro(dict(p, Ks=1.0), 1.0, 1.0, serie, dl=True)[0]  # Sa sin reducir (cm/s²)
    Sa = _interp(T, serie, curva)
    d = fi * (Sa * gam / lam)
    F = K @ d
    F_final = _comb(res, F)
    h = res['alturas_acum'][::-1]
    F_t = F_final[::-1] / 1000
    Fu = F_t * FC
    Vu = np.cumsum(Fu)
    deriva = _comb(res, derivas_modales(d))
    out = dict(curva=curva, Sa=Sa, d=d, F=F, F_final=F_final, d_comb=_comb(res, d),
               tabla=dict(nivel=np.arange(n, 0, -1), h=h, F=F_t, Fu=Fu, Vu=Vu, Mvu=h * Fu,
                          Mv=_volteo_entrepiso(Fu, p['alturas'])), V_el=float(Vu[-1]),
               deriva=deriva, dist=deriva / (p['alturas'] * 100))
    if p['cargas'] is not None:  # estático elástico: C = máx(c, a0) sin reducir
        W = np.asarray(p['cargas'] * p['dist_x'], float)[::-1]
        hW = h * W
        C = max(p['c'], p['a0']) / G
        Fi = C * hW / hW.sum() * W.sum()
        Fe = Fi * FC
        out['est'] = dict(C=C, tabla=dict(nivel=np.arange(n, 0, -1), h=h, Wi=W, hWi=hW, Fi=Fi, Fu=Fe, Vu=np.cumsum(Fe),
                                         Mvu=Fe * h, Mv=_volteo_entrepiso(Fe, p['alturas'])), V_est=float(np.sum(Fe)))
    return out


def calcular(p):
    for k_, v_ in NORMA_DEFECTO.items():
        p.setdefault(k_, v_)
    M, K, n = p['M'], p['K'], p['n']
    res = dict(p=p)

    # --- Modal ---
    A = np.linalg.solve(M, K)
    lam, fi = np.linalg.eig(A)
    lam, fi = lam.real, fi.real
    o = np.argsort(lam)
    lam, fi = lam[o], fi[:, o]
    if not p.get('manual') and lam.min() <= 1e-9 * np.abs(lam).max():
        raise ValueError('La estructura es inestable: algún nivel no tiene rigidez lateral (revisa articulaciones, '
                         'rodillos y columnas).')
    fi = _normalizar_modos(p, fi)
    w = np.sqrt(lam)
    f = w / (2 * np.pi)
    T = 1 / f
    uno = np.ones(n)
    gam = np.array([(fi[:, j] @ M @ uno) / (fi[:, j] @ M @ fi[:, j]) for j in range(n)])
    res.update(A=A, lam=lam, fi=fi, w=w, f=f, T=T, gam=gam, alturas_acum=np.cumsum(p['alturas']))
    res['n_modos'] = _n_modos(p, res)

    # --- Espectro de diseño según el estado límite (dinámico: cm/s^2) ---
    clave = CLAVE_ESTADO[p['estado']]
    pp = p if clave == 'SV' else dict(p, Q=np.array([1.0]))  # ocupación inmediata: Q' = 1
    Qs = pp['Q']
    serie = np.arange(0, 8.0005, 0.001)
    curvas = np.zeros((len(Qs), serie.size))
    curvas_g = np.zeros_like(curvas)
    Sa_modal = np.zeros((n, len(Qs)))
    a_min = a_min_g = None
    for i, q in enumerate(Qs):
        curvas[i], a_min = espectro(pp, q, 1.0, serie, dl=clave == 'DL', clave=clave)
        curvas_g[i], a_min_g = espectro(pp, q, 1 / G, serie, dl=clave == 'DL', clave=clave)
        Sa_modal[:, i] = _interp(T, serie, curvas[i])
    if p.get('Sa_manual') is not None:  # ordenadas Sa/(Q′R′) dadas por modo (p. ej. las de la hoja)
        if len(p['Sa_manual']) != n:
            raise ValueError(f'Hay {len(p["Sa_manual"])} ordenadas Sa por modo y la estructura tiene {n} modos.')
        Sa_modal[:, 0] = p['Sa_manual']
    res.update(serie=serie, curvas=curvas, curvas_g=curvas_g, Sa_modal=Sa_modal, a_min=a_min, a_min_g=a_min_g,
               Qs=Qs, clave=clave)

    # --- Desplazamientos y fuerzas modales (con la primera Q) y combinación modal ---
    d = np.zeros((n, n))
    F = np.zeros((n, n))
    for j in range(n):
        d[:, j] = fi[:, j] * (Sa_modal[j, 0] * gam[j] / lam[j])
        F[:, j] = K @ d[:, j]
    F_final = _comb(res, F)
    res.update(d=d, F=F, F_final=F_final, d_comb=_comb(res, d))

    # --- Tabla de fuerzas por nivel (dinámico) ---
    h_desc = res['alturas_acum'][::-1]
    F_desc = F_final[::-1] / 1000
    FC = 1.0 if clave == 'DL' else p['factor_Fu']  # limitación de daños: fuerzas sin factor de carga (apuntes p. 12)
    Fu = F_desc * FC
    Vu = np.cumsum(Fu)
    Mvu = h_desc * Fu
    res['tabla_din'] = dict(nivel=np.arange(n, 0, -1), h=h_desc, F=F_desc, Fu=Fu, Vu=Vu, Mvu=Mvu,
                            Mv=_volteo_entrepiso(Fu, p['alturas']))
    res['V_din'] = Vu[-1]

    # --- R' y QR' por modo ---
    filas = []
    for j in range(n):
        k2 = 0.0 if p['R_unitaria'] else float(_k2(T[j], p['Ta']))
        for q in Qs:
            Rp = float(_frp(p, clave) * _R(pp, q, T[j]))
            filas.append((j + 1, T[j], q, k2, Rp, q * Rp))
    res['QR'] = filas
    res['modo_Tmax'] = int(np.argmax(T)) + 1

    # --- Estático (apuntes p. 20-21) ---
    if p['cargas'] is not None:
        res['estatico'] = est = _estatico(p, res)
        res['tabla_est'] = est['tabla']
        res['Fmax_est'] = est['C']
        res['V_est'] = est['V_est']
        res['factor_escala'] = est['V_est'] * p['mult_V'] / res['V_din']
    res['revision'] = _revisiones(p, res)
    if 'estatico' in res:
        _completar_estatico(p, res)
    res['elastico'] = _elastico(p, res)
    res['normativa'] = _normativa(p, res)
    res['avisos'] = _avisos(res)
    return res


# ----------------------------------------------------------------------------
# Reporte en texto
# ----------------------------------------------------------------------------
def _tabla_txt(headers, filas, fmt='{:>14.6f}'):
    w = max(14, max(len(h) for h in headers) + 2)
    out = ''.join(f'{h:>{w}}' for h in headers) + '\n'
    for fila in filas:
        out += ''.join(f'{v:>{w}d}' if isinstance(v, (int, np.integer)) else
                       f'{v:>{w}}' if isinstance(v, str) else f'{v:>{w}.5f}' for v in fila) + '\n'
    return out


def tabla_derivacion_estatica(res):
    """Derivación del coeficiente sísmico estático (apuntes p. 20): (encabezados, filas)."""
    e, p = res['estatico'], res['p']
    c_g, a0_g = p['c'] / G, p['a0'] / G
    enc = ['Concepto', 'Plato (T ≥ Ta)', 'Inicio (T = 0)']
    if res['clave'] == 'DL':
        return enc, [['Ordenada elástica Sa (g)', c_g, a0_g], ['Ks', p['Ks'], p['Ks']],
                     ['C = Sa·Ks (g)', e['C_pl'], e['C_0']]]
    return enc, [['Q′', float(e['q_pl']), 1.0], ['k2', 0.0, 0.5], ['R = k1·R0 + k2', float(e['R_pl']), float(e['R_0'])],
                 ["R′", float(e['Rp_pl']), float(e['Rp_0'])],
                 ["Q′·R′", float(e['q_pl'] * e['Rp_pl']), float(e['Rp_0'])],
                 ['Ordenada Sa (g)', c_g, a0_g], ["C = Sa / (Q′·R′) (g)", float(e['C_pl']), float(e['C_0'])]]


def tabla_comparacion_metodos(res):
    """Cortante basal, volteo y coeficiente de cada análisis: (encabezados, filas)."""
    p, el = res['p'], res['elastico']
    W_T = float(np.sum(p['cargas'])) if p['cargas'] is not None else None
    t = res['tabla_din']

    def fila(nombre, V, M, obs=''):
        return [nombre, float(V), float(M), (float(V) / W_T if W_T else np.nan), obs]

    filas = [fila('Dinámico (modal espectral, diseño)', res['V_din'], t['Mvu'].sum(),
                  f'{res["n_modos"]} de {p["n"]} modos · {p["combinacion"]}')]
    if 'tabla_est' in res:
        filas.append(fila('Estático (apuntes)', res['V_est'], res['tabla_est']['Mvu'].sum(),
                          f'C = {res["Fmax_est"]:.4f} g'))
        f = res['factor_aplicado']
        filas.append(fila(f'Dinámico escalado (× {f:.3f})', res['V_din'] * f, t['Mvu'].sum() * f,
                          'el que se usa para diseñar' if f > 1 else 'sin escalar (el dinámico ya cumple)'))
    filas.append(fila('Elástico dinámico (sin reducir)', el['V_el'], el['tabla']['Mvu'].sum(), 'Q′ = 1, R′ = 1'))
    if 'est' in el:
        filas.append(fila('Elástico estático (sin reducir)', el['est']['V_est'], el['est']['tabla']['Mvu'].sum(),
                          f'C = {el["est"]["C"]:.4f} g'))
    return ['Análisis', 'V base (t)', 'Σ Mvu (t·m)', 'V/W (g)', 'Observación'], filas


def tabla_niveles_metodos(res):
    """Fu y Vu por nivel de cada análisis (arriba hacia abajo): (encabezados, filas)."""
    p, n, el = res['p'], res['p']['n'], res['elastico']
    t = res['tabla_din']
    cols = [('Fu dinámico', t['Fu']), ('Vu dinámico', t['Vu'])]
    if 'tabla_est' in res:
        te = res['tabla_est']
        cols += [('Fu estático', te['Fu']), ('Vu estático', te['Vu']), ('Fu dinámico escalado', res['tabla_din_esc']['Fu']),
                 ('Vu dinámico escalado', res['tabla_din_esc']['Vu'])]
    cols += [('Fu elástico', el['tabla']['Fu']), ('Vu elástico', el['tabla']['Vu'])]
    enc = ['Nivel', 'h (m)'] + [c[0] for c in cols]
    filas = [[int(t['nivel'][i]), float(t['h'][i])] + [float(c[1][i]) for c in cols] for i in range(n)]
    return enc, filas


def tabla_elastico_comparacion(res):
    """Fuerza y distorsión de diseño contra elásticas, por nivel (arriba hacia abajo): (encabezados, filas)."""
    p, n, el = res['p'], res['p']['n'], res['elastico']
    t = res['tabla_din']
    dis_d = res['revision']['estados'][res['clave']]['dist'][::-1]
    el_d = el['dist'][::-1]
    enc = ['Nivel', 'h (m)', 'F diseño (t)', 'F elástica (t)', 'F el. / F dis.', 'γ diseño', 'γ elástica', 'γ el. / γ dis.']
    filas = [[int(t['nivel'][i]), float(t['h'][i]), float(t['F'][i]), float(el['tabla']['F'][i]),
              float(el['tabla']['F'][i] / t['F'][i]), float(dis_d[i]), float(el_d[i]),
              float(el_d[i] / dis_d[i]) if dis_d[i] else np.nan] for i in range(n)]
    return enc, filas


def tabla_distorsiones_estatico(res):
    """Distorsión estática y dinámica del estado de cálculo por nivel (arriba hacia abajo): (encabezados, filas)."""
    e, p, n = res['estatico'], res['p'], res['p']['n']
    enc = ['Nivel', 'h (m)', 'γ estático', 'γ dinámico', 'Límite', 'Cumple (est. / din.)']
    filas = []
    for i in range(n - 1, -1, -1):
        lim = e['limite']
        cumple = '-' if lim is None else ('Sí' if e['dist'][i] <= lim else 'NO') + ' / ' + ('Sí' if e['dist_din'][i] <= lim else 'NO')
        filas.append([i + 1, float(p['alturas'][i]), float(e['dist'][i]), float(e['dist_din'][i]),
                      lim if lim is not None else 'sin definir', cumple])
    return enc, filas


def tabla_distorsiones(res):
    """Filas (de arriba hacia abajo): nivel, h, y distorsión de cada estado límite. Devuelve (encabezados, filas)."""
    p, n = res['p'], res['p']['n']
    est = res['revision']['estados']
    claves = ('DL', 'SV', 'OI')
    enc = ['Nivel', 'h (m)'] + [f'γ {est[c]["nombre"]}' + (' ★' if c == res['clave'] else '') for c in claves] + ['Cumple']
    filas = []
    for i in range(n - 1, -1, -1):
        cumple = []
        for c in claves:
            lim = est[c]['limite']
            cumple.append('-' if lim is None else ('Sí' if est[c]['dist'][i] <= lim else 'NO'))
        filas.append([i + 1, float(p['alturas'][i])] + [float(est[c]['dist'][i]) for c in claves]
                     + [' / '.join(cumple)])
    return enc, filas


def tabla_pdelta(res):
    """Nivel, P, V, δ, θ y estado por entrepiso (arriba hacia abajo). None si no hay pesos."""
    rev, n = res['revision'], res['p']['n']
    if rev['P'] is None:
        return None
    enc = ['Nivel', 'P (t)', 'V (t)', 'δ (cm)', 'θ', 'Estado', 'k ef. (t/cm)']
    filas = []
    for i in range(n - 1, -1, -1):
        th = rev['theta'][i]
        estado = 'Estable' if th <= 0.10 else ('Considerar P-Δ' if th <= 0.25 else 'Inestable')
        filas.append([i + 1, float(rev['P'][i]), float(rev['V'][i]), float(rev['estados']['SV']['deriva_el'][i]),
                      float(th), estado, float(rev['k_ef'][i])])
    return enc, filas


def texto_normativa(res):
    """Líneas de texto del marco normativo (NTC-S 2023, apuntes de clase)."""
    nm, p = res['normativa'], res['p']
    s = ['\n=== MARCO NORMATIVO (NTC-S 2023, apuntes de clase) ===']
    ancho = max(len(f[0]) for f in nm['factores']) + 2
    for nombre, valor, fuente in nm['factores']:
        s.append(f'  {nombre:<{ancho}}{valor:<34}{fuente}')
    if p.get('objetivo') and p['objetivo']['nota']:
        s.append('  ! ' + p['objetivo']['nota'])
    v = nm.get('vmin')
    if v:
        s.append(f'\nCortante basal mínimo: FC·a_min·W = {p["factor_Fu"]:g}·{v["a_min_g"]:.5f}·{v["W_T"]:.2f} = {v["V_min"]:.3f} t · '
                 f'V dinámico = {v["V_din"]:.3f} t → ' + ('cumple' if v['cumple'] else f'NO cumple (escalar × {v["factor"]:.3f})'))
    c = nm.get('cimentacion')
    if c:
        s.append(f'Cimentación (× 0.65·R′ = {c["factor"]:.3f}): V = {c["V"]:.3f} t · Mvolteo = {c["M"]:.3f} t·m')
    t = nm.get('torsion')
    if t:
        s.append(f'\nTorsión accidental (b = {t["b"]:g} m): e_a = [0.05 + 0.05(i−1)/(n−1)]·b; Mt = Fu·e_a')
        s.append(_tabla_txt(['Nivel', 'e_a (m)', 'Fu (t)', 'Mt (t·m)'],
                            [[i + 1, float(t['e_a'][i]), float(t['Fu'][i]), float(t['Mt'][i])] for i in range(p['n'] - 1, -1, -1)]))
    e = nm['estatico']
    s.append(f'Método estático: H = {e["H"]:.1f} m' + (f' (límite {e["limite"]:g} m, zona {e["zona"]}, '
             f'{"regular" if e["regular"] else "irregular"})' if e['limite'] else '') + ' → '
             + ('aplicable' if e['aplica'] else 'NO aplicable (grupo A o altura excedida)'))
    m = nm['modos']
    s.append(f'Participación modal: {m["n95"]} modo(s) alcanzan 95% de la masa efectiva; modos con T ≥ 0.4 s: {m["n_T04"]}.')
    return s


def texto_elastico(res):
    """Análisis elástico (sin reducir) y comparación de métodos, en texto."""
    el, p, n = res['elastico'], res['p'], res['p']['n']
    s = ['\n=== ANÁLISIS ELÁSTICO (sin reducir: Q′ = 1, R′ = 1) ===',
         f'Se combinan {res["n_modos"]} de {n} modos con {p["combinacion"]}.']
    s.append(_tabla_txt(['Modo', 'T(s)', 'Sa elást.', 'Sa diseño', 'Sa el./Sa dis.'],
                        [[j + 1, res['T'][j], el['Sa'][j], res['Sa_modal'][j, 0], el['Sa'][j] / res['Sa_modal'][j, 0]]
                         for j in range(n)]))
    t = el['tabla']
    s.append(_tabla_txt(['Nivel', 'h(m)', 'F(t)', 'Fu(t)', 'Vu(t)', 'Mvu(t-m)'],
                        zip(t['nivel'], t['h'], t['F'], t['Fu'], t['Vu'], t['Mvu'])))
    s.append(f'V elástico = {el["V_el"]:.4f} t · V de diseño = {res["V_din"]:.4f} t · '
             f'V elástico / V diseño = {el["V_el"] / res["V_din"]:.3f}')
    enc, filas = tabla_elastico_comparacion(res)
    s.append('\nFuerza y distorsión elásticas contra las de diseño:')
    s.append(_tabla_txt(enc, filas))
    enc, filas = tabla_comparacion_metodos(res)
    s.append('=== COMPARACIÓN DE MÉTODOS ===')
    s.append(f'{enc[0]:<38}{enc[1]:>12}{enc[2]:>14}{enc[3]:>10}   {enc[4]}')
    for nombre, V, M, c, obs in filas:
        s.append(f'{nombre:<38}{V:>12.3f}{M:>14.2f}{c:>10.4f}   {obs}')
    return s


def reporte_texto(res):
    p, n = res['p'], res['p']['n']
    np.set_printoptions(linewidth=200, suppress=True, precision=5)
    s = [f'Estado límite: {p["estado"]}'
         + (f' (Ks = {p["Ks"]:g})' if res['clave'] == 'DL' else '')
         + f' · Dirección del sismo: {p.get("direccion", "X")} · Combinación modal: {p["combinacion"]}'
         + (f' (ζ = {p["zeta"]:g})' if p['combinacion'] == 'CQC' else '')
         + f' · Vigas: {p.get("vigas", "rigidas")}\n']
    if 'info' in p:
        s.append('=== MODELO ESTRUCTURAL ===')
        s.append(f'E = {p["E"]:,.0f} kg/cm²\n')
        s.append(_tabla_txt(['Nivel', 'h(m)', 'Nº col', 'Nº art', 'Sección(cm)', 'ΣI(cm4)', 'K nivel', 'W(t)', 'm'],
                            [[r['nivel'], r['h'], r['ncol'], r['nart'], r['seccion'], r['I'], r['k'], r['W'], r['m']]
                             for r in p['info']]))
        s.append('Matriz K [kg/cm]:')
        for i in range(n):
            s.append('  ' + ''.join(f'{x:>14.2f}' for x in p['K'][i]))
        s.append('Matriz M [kg·s²/cm]:')
        for i in range(n):
            s.append('  ' + ''.join(f'{x:>14.4f}' for x in p['M'][i]))
        s.append('')
    s.append('=== RESULTADOS MODALES ===\n')
    s.append('Modo            ' + ''.join(f'{j + 1:>14d}' for j in range(n)))
    for nombre, v in (('λ [rad/s]²', res['lam']), ('ω [rad/s]', res['w']), ('f [Hz]', res['f']),
                      ('T [s]', res['T']), ('γ (partic.)', res['gam'])):
        s.append(f'{nombre:<16}' + ''.join(f'{x:>14.5f}' for x in v))
    s.append(f'Σγ = {res["gam"].sum():.5f}')
    s.append('\nEigenvectores normalizados (columnas = modos):')
    for i in range(n):
        s.append('  ' + ''.join(f'{x:>12.5f}' for x in res['fi'][i]))
    s.append('\nAceleraciones espectrales Sa/(Q\'R\') [cm/s²]:')
    s.append(_tabla_txt(['Modo', 'T(s)'] + [f'Q={q:g}' for q in res['Qs']],
                        [[j + 1, res['T'][j]] + list(res['Sa_modal'][j]) for j in range(n)]))
    s.append(f'a_mínima = {res["a_min"]:.6f} cm/s²  ({res["a_min_g"]:.6f} g)   Ts = {p["Ts"]:g}')
    s.append('\nDesplazamientos modales δ [cm] (columnas = modos):')
    for i in range(n):
        s.append('  ' + ''.join(f'{x:>14.6f}' for x in res['d'][i]))
    s.append('\nFuerzas modales F [kg] (columnas = modos):')
    for i in range(n):
        s.append('  ' + ''.join(f'{x:>14.4f}' for x in res['F'][i]))
    s.append('\nRaíz de la suma de cuadrados de F (SRSS) [kg]:')
    s.append('  ' + '  '.join(f'{x:.4f}' for x in res['F_final']))
    s.append('\n=== R\' y Q·R\' POR MODO ===')
    s.append(_tabla_txt(['Modo', 'T(s)', 'Q', 'k2', "R'", "Q·R'"], res['QR']))
    t = res['tabla_din']
    s.append('=== TABLA DE FUERZAS POR NIVEL (DINÁMICO) ===')
    s.append(_tabla_txt(['Nivel', 'h(m)', 'F(t)', 'Fu(t)', 'Vu(t)', 'Mvu(t-m)'],
                        zip(t['nivel'], t['h'], t['F'], t['Fu'], t['Vu'], t['Mvu'])))
    s.append(f'Σ Mvu = {t["Mvu"].sum():.6f} t-m')
    if 'tabla_est' in res:
        e, se = res['tabla_est'], res['estatico']
        s.append('\n=== MÉTODO ESTÁTICO (apuntes p. 20-21) ===')
        s.append('Fi = C · (Wi·hi / Σ Wi·hi) · ΣWi,  C = máx { c/(Q′R′) en el plato ; a0/(Q′R′) en T = 0 }\n')
        enc, filas = tabla_derivacion_estatica(res)
        s.append(_tabla_txt(enc, filas))
        s.append(f'C = {se["C"]:.6f} g · gobierna: {se["gobierna"]}')
        s.append(_tabla_txt(['Nivel', 'h(m)', 'Wi(t)', 'Wihi(t-m)', 'Fi(t)', 'Fu(t)', 'Vu(t)', 'Mvu(t-m)'],
                            zip(e['nivel'], e['h'], e['Wi'], e['hWi'], e['Fi'], e['Fu'], e['Vu'], e['Mvu'])))
        s.append(f'ΣWi = {e["Wi"].sum():.6f}   ΣWihi = {e["hWi"].sum():.6f}   ΣMvu = {e["Mvu"].sum():.6f}')
        s.append(f'Comprobación: ΣFi = C·ΣWi = {se["C"]:.6f}·{se["W_T"]:.4f} = {se["C"] * se["W_T"]:.6f} t '
                 f'(= {e["Fi"].sum():.6f})')
        s.append(f'V estático  = {res["V_est"]:.6f} t' + (f'  (× {p["mult_V"]:g} = {res["V_est"] * p["mult_V"]:.6f})'
                                                       if p['mult_V'] != 1 else ''))
        s.append(f'V dinámico  = {res["V_din"]:.6f} t')
        s.append(f'FACTOR DE ESCALA = {res["factor_escala"]:.6f}'
                 + ('   (V dinámico ya cumple, factor < 1)' if res['factor_escala'] < 1 else ''))
        enc, filas = tabla_distorsiones_estatico(res)
        s.append(f'\nDistorsiones del método estático ({p["estado"]}; desplazamientos × '
                 + ('1' if res['clave'] == 'DL' else f'FC·Q·R′ = {se["amp"]:.4f}') + '):')
        s.append(_tabla_txt(enc[:-1], [f[:-1] for f in filas]))
        s.append('Cumple (estático / dinámico): ' + '; '.join(f'N{f[0]}: {f[-1]}' for f in filas))
    else:
        s.append('\n(Sin cargas/longitudes: no se calculó el método estático ni el factor de escala.)')
    s.extend(texto_elastico(res))
    s.extend(texto_normativa(res))
    rev = res['revision']
    s.append('\n=== REVISIÓN DE DISTORSIONES (γ = deriva / altura de entrepiso) ===')
    for c in ('DL', 'SV', 'OI'):
        e = rev['estados'][c]
        s.append(f'  {e["nombre"]:<22} límite = ' + (f'{e["limite"]:g}' if e['limite'] is not None else 'sin definir')
                 + (f'  (γmáx {e["limite_base"]:g} × γc {e["gamma_c"]:g})' if e['gamma_c'] != 1.0 and e['limite'] is not None else ''))
    enc, filas = tabla_distorsiones(res)
    s.append(_tabla_txt(enc[:-1], [f[:-1] for f in filas]))
    s.append('Cumple (LD / SV / OI): ' + '; '.join(f'N{f[0]}: {f[-1]}' for f in filas))
    pd_ = tabla_pdelta(res)
    if pd_:
        s.append('\n=== EFECTOS P-Δ (θ = P·δ / (V·h); ≤ 0.10 se pueden ignorar, > 0.25 inestable) ===')
        s.append(_tabla_txt(pd_[0], pd_[1]))
        s.append(f'Irregularidad de masa: {rev["irreg_masa"] or "ninguna"} · piso blando: {rev["irreg_rigidez"] or "ninguno"}')
    s.append('\n=== AVISOS DE COHERENCIA ===')
    marcas = {'alerta': '[!!]', 'aviso': '[! ]', 'info': '[i ]'}
    s.extend(f'{marcas[sev]} {txt}' for sev, txt in res['avisos'])
    if not res['avisos']:
        s.append('Sin avisos.')
    return '\n'.join(s)


# ----------------------------------------------------------------------------
# Gráficas (Figure puro, sirve para GUI y para exportar)
# ----------------------------------------------------------------------------
def fig_modal(res):
    n = res['p']['n']
    fig = Figure(figsize=(max(6, 2.6 * n), 5.2), dpi=100)
    y = np.concatenate([[0], res['alturas_acum']])
    lim = 1.2 * np.abs(res['fi']).max()
    for j in range(n):
        ax = fig.add_subplot(1, n, j + 1)
        x = np.concatenate([[0], res['fi'][:, j]])
        ax.axvline(0, color='0.7', lw=1)
        ax.plot(x, y, '-o', color='#1f4e79', lw=1.8, ms=5)
        ax.set_xlim(-lim, lim)
        ax.set_ylim(0, y.max() * 1.05)
        ax.set_title(f'Modo {j + 1}\nT = {res["T"][j]:.4f} s', fontsize=10)
        ax.set_xlabel('Φ')
        if j == 0:
            ax.set_ylabel('Altura (m)')
        ax.grid(alpha=0.3)
    fig.tight_layout()
    return fig


def fig_espectro(res):
    p = res['p']
    fig = Figure(figsize=(8, 7), dpi=100)
    for k, (clave, unidad, Tcol) in enumerate((('curvas', 'cm/s²', 1.0), ('curvas_g', 'g', 1 / G))):
        ax = fig.add_subplot(2, 1, k + 1)
        for i, q in enumerate(res['Qs']):
            ln, = ax.plot(res['serie'], res[clave][i], lw=2, color=AZULES[i % len(AZULES)], label=f'Q = {q:g}')
            ax.plot(res['T'], _interp(res['T'], res['serie'], res[clave][i]), 'o', color=ln.get_color())
        ax.set_xlabel('T (s)')
        ax.set_ylabel(f"Sa / (Q'R')  [{unidad}]")
        ax.set_title('Espectro de diseño ' + ('(dinámico)' if k == 0 else '(estático)'))
        ax.grid(alpha=0.3)
        ax.legend(loc='upper right')
    fig.tight_layout()
    return fig


NOMBRE_SECCION = {'rect': 'Rect.', 'hrect': 'Rect. hueca', 'circ': 'Circ.', 'hcirc': 'Tubo', 'H': 'Perfil H',
                  'T': 'Sección T', 'gen': 'Propiedades'}


def fig_estructura(p):
    """Vista previa a escala (elevación): ejes, columnas con su sección, vigas, apoyos y cargas por crujía."""
    from matplotlib.patches import Circle, Polygon, Rectangle
    g = p['geom']
    xs, bays, niv = g['xs'], g['bays'], g['niveles']
    direccion = p.get('direccion', 'X')
    flexibles = p.get('vigas') == 'flexibles'
    zs = np.concatenate([[0.0], np.cumsum([l['h'] for l in niv])])
    tvs = [l_['viga_h'] / 100 for l_ in niv]
    L, H = xs[-1], zs[-1]
    fig = Figure(figsize=(10, min(11, max(5.2, 0.36 * (H + 4.4) + 1))), dpi=100)
    gs = fig.add_gridspec(1, 2, width_ratios=[3.4, 1], wspace=0.04)
    ax = fig.add_subplot(gs[0])
    ax.set_aspect('equal')
    ax.axis('off')
    ax.set_title(f'Vista previa de la estructura (elevación a escala, sismo en {direccion}) · '
                 + ('vigas flexibles' if flexibles else 'vigas rígidas'),
                 color=AZULES[2], fontsize=10.5, fontweight='bold', loc='left')

    # terreno
    vi, vd = g.get('volados', (0.0, 0.0))
    x0, x1 = xs[0] - 0.9, xs[-1] + 0.9
    ax.plot([x0, x1], [0, 0], color=AZULES[0], lw=1.6, zorder=2)
    for xh in np.arange(x0, x1, 0.25):
        ax.plot([xh, xh - 0.18], [0, -0.22], color=AZULES[3], lw=0.8, zorder=1)

    for i, l in enumerate(niv):
        z0, z1 = zs[i], zs[i + 1]
        tv = l['viga_h'] / 100  # peralte de la viga (m)
        xl, xr = l.get('ext', [xs[0], xs[-1]])
        for xa_, xb_ in l.get('segs', [(xl, xr)]):          # losa/viga solo donde existe (puede haber tramos sin losa)
            ax.add_patch(Rectangle((xa_ - 0.3, z1 - tv), (xb_ - xa_) + 0.6, tv, fc='#DCE9F7', ec=AZULES[1], lw=1, zorder=3))
        for c in l['cols']:
            if c.get('k_cont'):                              # continuación: ya se dibujó con la columna del nivel inferior
                continue
            fin = c.get('fin_nivel', i)
            zt = zs[fin + 1] - tvs[fin]                      # tope de la columna (cruza pisos si fin > i)
            w = max(c['eq'] / 100, 0.14)
            circular = c['tipo'] in ('circ', 'hcirc')
            ax.add_patch(Rectangle((c['x'] - w / 2, z0), w, zt - z0, fc=AZULES[1] if circular else AZULES[2],
                                   ec=AZULES[0], lw=0.8, zorder=4))
            cruza = fin > i                                  # la etiqueta de una columna que cruza va a su izquierda
            ax.text(c['x'] + (-1 if cruza else 1) * (w / 2 + 0.1), (z0 + zs[fin + 1]) / 2, c['texto'], fontsize=7.5,
                    color=AZULES[0], va='center', ha='right' if cruza else 'left', zorder=6)
            if c.get('rodillo'):  # apoyo de rodillos: triángulo sobre dos ruedas
                ax.add_patch(Polygon([(c['x'], 0), (c['x'] - 0.25, -0.32), (c['x'] + 0.25, -0.32)],
                                     fc='white', ec=AZULES[0], lw=1, zorder=5))
                for dx_ in (-0.13, 0.13):
                    ax.add_patch(Circle((c['x'] + dx_, -0.42), 0.085, fc='white', ec=AZULES[0], lw=1, zorder=5))
            elif c['artic']:
                if i == 0:
                    ax.add_patch(Polygon([(c['x'], 0), (c['x'] - 0.25, -0.42), (c['x'] + 0.25, -0.42)],
                                         fc='white', ec=AZULES[0], lw=1, zorder=5))
                else:
                    ax.add_patch(Circle((c['x'], z0), 0.13, fc='white', ec=AZULES[0], lw=1, zorder=6))
            elif i == 0:
                ax.add_patch(Rectangle((c['x'] - w / 2 - 0.1, -0.12), w + 0.2, 0.12, fc=AZULES[0], zorder=5))
            if c.get('artic_sup'):
                ax.add_patch(Circle((c['x'], zt), 0.13, fc='white', ec=AZULES[0], lw=1, zorder=6))
        # cargas
        if l['w_bays'] is None:
            ax.text(L / 2, z1 + 0.2, f'W = {l["W"]:g} t', ha='center', fontsize=8, color=AZULES[2], fontweight='bold')
        else:
            for xa0, xb0, w_ in l['w_spans']:
                if w_ == 0:
                    continue
                mg = 0.4 if (xb0 - xa0) > 1.2 else 0.15
                xa, xb, zt = xa0 + mg, xb0 - mg, z1 + 0.55
                ax.plot([xa, xb], [zt, zt], color=AZULES[2], lw=1.2, zorder=5)
                for xv in np.linspace(xa, xb, max(3, int((xb - xa) / 0.5))):
                    ax.annotate('', xy=(xv, z1 + 0.02), xytext=(xv, zt), zorder=5,
                                arrowprops=dict(arrowstyle='-|>', color=AZULES[2], lw=0.8, mutation_scale=7))
                ax.text((xa + xb) / 2, zt + 0.1, f'{w_:g} t/m', ha='center', va='bottom', fontsize=8, color=AZULES[0])
        ax.text(xs[-1] + vd + 0.7, z1 - 0.15, f'N{i + 1}', fontsize=10, fontweight='bold', color=AZULES[0], va='center')
        ax.text(xs[-1] + vd + 0.7, z1 - 0.7, f'W = {l["W"]:.2f} t', fontsize=8, color=AZULES[1], va='center')
        if flexibles:
            ax.text(xs[-1] + vd + 0.7, z1 - 1.2, f'Viga {l["viga"]["texto"]}', fontsize=8, color=AZULES[1], va='center')
        xd = xs[0] - vi - 1.2  # cota de altura
        ax.annotate('', xy=(xd, z1), xytext=(xd, z0), arrowprops=dict(arrowstyle='<->', color=AZULES[1], lw=1))
        ax.text(xd - 0.12, (z0 + z1) / 2, f'{l["h"]:g} m', rotation=90, ha='right', va='center', fontsize=8,
                color=AZULES[1])

    # ejes y cotas de crujías
    for j, x in enumerate(xs):
        ax.text(x, -1.85, str(j + 1), ha='center', va='center', fontsize=9, color=AZULES[0],
                bbox=dict(boxstyle='circle,pad=0.25', fc='white', ec=AZULES[2], lw=1))
    for j, b in enumerate(bays):
        ax.annotate('', xy=(xs[j + 1], -1.0), xytext=(xs[j], -1.0),
                    arrowprops=dict(arrowstyle='<->', color=AZULES[1], lw=1))
        ax.text((xs[j] + xs[j + 1]) / 2, -0.9, f'{b:g} m', ha='center', va='bottom', fontsize=8, color=AZULES[1])
    for a_, b_, lab in ((-vi, 0.0, vi), (xs[-1], xs[-1] + vd, vd)):
        if lab > 0:
            ax.annotate('', xy=(b_, -1.0), xytext=(a_, -1.0), arrowprops=dict(arrowstyle='<->', color=AZULES[1], lw=1))
            ax.text((a_ + b_) / 2, -0.9, f'{lab:g} m (volado)', ha='center', va='bottom', fontsize=8, color=AZULES[1])
    ax.annotate('', xy=(L / 2 + 1.6, H + 1.6), xytext=(L / 2 - 1.6, H + 1.6),
                arrowprops=dict(arrowstyle='-|>', color=AZULES[0], lw=2))
    ax.text(L / 2, H + 1.8, 'Dirección del análisis', ha='center', fontsize=9, color=AZULES[0])
    ax.set_xlim(xs[0] - vi - 2.4, xs[-1] + vd + 3.6)
    ax.set_ylim(-2.4, H + 2.6)

    # secciones en planta, a escala (X horizontal, Y vertical)
    axs = fig.add_subplot(gs[1])
    axs.set_aspect('equal')
    axs.axis('off')
    axs.set_title('Secciones en planta (cm)', color=AZULES[2], fontsize=10.5, fontweight='bold', loc='left')
    secs = {}
    for i, l in enumerate(niv, 1):
        for c in l['cols']:
            secs.setdefault(c['texto'], dict(sec=c, niveles=set()))['niveles'].add(i)
    mmax = max(max(v['sec']['dx'], v['sec']['dy']) for v in secs.values())
    y = 0.0
    for v in secs.values():
        c = v['sec']
        cy = y - c['dy'] / 2
        dibujar_seccion(axs, c, 0, cy)
        axs.text(mmax / 2 + 0.12 * mmax, cy,
                 f'{NOMBRE_SECCION[c["tipo"]]} {c["texto"]}\nNiveles: {", ".join(map(str, sorted(v["niveles"])))}',
                 fontsize=8, color=AZULES[0], va='center')
        y -= c['dy'] + 0.45 * mmax
    axs.annotate('', xy=(mmax * 0.55, 0.45 * mmax), xytext=(-mmax * 0.55, 0.45 * mmax),
                 arrowprops=dict(arrowstyle='-|>', color=AZULES[0], lw=1.5)) if direccion == 'X' else \
        axs.annotate('', xy=(-mmax * 0.65, 0.7 * mmax), xytext=(-mmax * 0.65, -0.1 * mmax),
                     arrowprops=dict(arrowstyle='-|>', color=AZULES[0], lw=1.5))
    axs.text(0, 0.6 * mmax, f'sismo en {direccion}', ha='center', fontsize=8, color=AZULES[0])
    axs.set_xlim(-mmax * 0.8, mmax * 3.4)
    axs.set_ylim(y + 0.2 * mmax, 0.85 * mmax)
    fig.subplots_adjust(left=0.01, right=0.99, top=0.93, bottom=0.02)
    return fig


def fig_elastico(res):
    """Espectro elástico contra el de diseño y fuerzas elásticas contra fuerzas de diseño por nivel."""
    p, n, el = res['p'], res['p']['n'], res['elastico']
    fig = Figure(figsize=(10, 4.2), dpi=100)
    ax = fig.add_subplot(1, 2, 1)
    ax.plot(res['serie'], el['curva'] / G, color=AZULES[3], lw=2, ls='--', label='Sa elástico (sin reducir)')
    ax.plot(res['serie'], res['curvas_g'][0], color=AZULES[2], lw=2, label="Sa/(Q′R′) de diseño")
    ax.plot(res['T'], el['Sa'] / G, 'o', color=AZULES[3])
    ax.plot(res['T'], res['Sa_modal'][:, 0] / G, 'o', color=AZULES[2])
    ax.set_xlabel('T (s)')
    ax.set_ylabel('Sa (g)')
    ax.set_xlim(0, min(8, max(4.0, 1.3 * p['Tb'] + 1)))
    ax.set_title('Espectros', color=AZULES[2], fontsize=10.5, fontweight='bold', loc='left')
    ax.grid(alpha=0.3)
    ax.legend(loc='upper right', fontsize=8)
    ax2 = fig.add_subplot(1, 2, 2)
    t = res['tabla_din']
    y = np.arange(n)
    ax2.barh(y + 0.2, el['tabla']['F'][::-1], height=0.38, color=AZULES[3], label='Elástica')
    ax2.barh(y - 0.2, t['F'][::-1], height=0.38, color=AZULES[2], label='Diseño')
    ax2.set_yticks(y)
    ax2.set_yticklabels([f'N{i}' for i in t['nivel'][::-1]])
    ax2.set_xlabel('Fuerza por nivel F (t)')
    ax2.set_title('Fuerzas sísmicas por nivel', color=AZULES[2], fontsize=10.5, fontweight='bold', loc='left')
    ax2.grid(axis='x', alpha=0.3)
    ax2.legend(fontsize=8)
    fig.tight_layout()
    return fig


def fig_comparacion(res):
    """Cortante por nivel (escalera) y Fu por nivel de los distintos análisis."""
    p, n, el = res['p'], res['p']['n'], res['elastico']
    zs = np.concatenate([[0.0], np.cumsum(p['alturas'])])
    t = res['tabla_din']
    series = [('Dinámico', t['Vu'][::-1], t['Fu'][::-1], AZULES[2], '-')]
    if 'tabla_est' in res:
        series.append(('Estático', res['tabla_est']['Vu'][::-1], res['tabla_est']['Fu'][::-1], AZULES[3], '--'))
        series.append(('Dinámico escalado', res['tabla_din_esc']['Vu'][::-1], res['tabla_din_esc']['Fu'][::-1], AZULES[0], ':'))
    series.append(('Elástico', el['tabla']['Vu'][::-1], el['tabla']['Fu'][::-1], '#7A9CC6', '-.'))
    fig = Figure(figsize=(10, 4.4), dpi=100)
    ax = fig.add_subplot(1, 2, 1)
    for nombre, V, _, col, ls in series:
        ax.step(np.concatenate([[V[0]], V]), zs, where='pre', color=col, ls=ls, lw=2, label=nombre)
    ax.set_xlabel('Cortante Vu (t)')
    ax.set_ylabel('Altura (m)')
    ax.set_title('Cortante por entrepiso', color=AZULES[2], fontsize=10.5, fontweight='bold', loc='left')
    ax.grid(alpha=0.3)
    ax.legend(fontsize=8)
    ax2 = fig.add_subplot(1, 2, 2)
    ancho = 0.8 / len(series)
    y = np.arange(n)
    for k, (nombre, _, F, col, _) in enumerate(series):
        ax2.barh(y + (k - (len(series) - 1) / 2) * ancho, F, height=ancho, color=col, label=nombre)
    ax2.set_yticks(y)
    ax2.set_yticklabels([f'N{i}' for i in range(1, n + 1)])
    ax2.set_xlabel('Fuerza Fu (t)')
    ax2.set_title('Fuerza por nivel', color=AZULES[2], fontsize=10.5, fontweight='bold', loc='left')
    ax2.grid(axis='x', alpha=0.3)
    fig.tight_layout()
    return fig


def fig_distorsiones(res):
    """Distorsión de entrepiso por estado límite contra su límite."""
    p, n = res['p'], res['p']['n']
    est = res['revision']['estados']
    zm = np.concatenate([[0], np.cumsum(p['alturas'])])
    zmid = (zm[:-1] + zm[1:]) / 2
    fig = Figure(figsize=(7.5, 5), dpi=100)
    ax = fig.add_subplot(111)
    estilos = {'DL': (AZULES[3], 'o'), 'SV': (AZULES[2], 's'), 'OI': (AZULES[0], '^')}
    for c, e in est.items():
        col, mk = estilos[c]
        ax.plot(e['dist'], zmid, marker=mk, color=col, lw=2, label=e['nombre'])
        if e['limite'] is not None:
            ax.axvline(e['limite'], color=col, ls='--', lw=1)
            ax.text(e['limite'], zm[-1] * 1.02, f'{e["limite"]:g}', color=col, ha='center', fontsize=8)
    ax.set_xlabel('Distorsión de entrepiso γ = Δ / h')
    ax.set_ylabel('Altura (m)')
    ax.set_title('Revisión de distorsiones por estado límite', color=AZULES[2], fontweight='bold', loc='left')
    ax.set_ylim(0, zm[-1] * 1.08)
    ax.grid(alpha=0.3)
    ax.legend(loc='lower right')
    fig.tight_layout()
    return fig


def fig_fuerzas(res):
    """Fuerza sísmica, cortante y momento de volteo por nivel (análisis dinámico)."""
    t = res['tabla_din']
    n = len(t['nivel'])
    fig = Figure(figsize=(9, 3.6), dpi=100)
    for k, (clave, titulo) in enumerate((('Fu', 'Fu (t)'), ('Vu', 'Vu (t)'), ('Mvu', 'Mvu (t·m)'))):
        ax = fig.add_subplot(1, 3, k + 1)
        y = np.arange(n)
        ax.barh(y, t[clave][::-1], color=AZULES[2 - (k % 2)], height=0.6)
        for yi, v in zip(y, t[clave][::-1]):
            ax.text(v, yi, f' {v:.2f}', va='center', fontsize=8, color=AZULES[0])
        ax.set_yticks(y)
        ax.set_yticklabels([f'N{i}' for i in t['nivel'][::-1]] if k == 0 else [])
        ax.set_title(titulo, color=AZULES[2], fontsize=10, fontweight='bold')
        ax.set_xlim(0, max(t[clave]) * 1.25)
        ax.grid(axis='x', alpha=0.3)
    fig.tight_layout()
    return fig


def fig_diagramas(res, metodo='dinamico'):
    """Diagramas Fu, Vu y Mvu como en las hojas de clase: marco con las fuerzas, escalera de cortantes y volteo."""
    from matplotlib.patches import FancyArrowPatch, Rectangle
    p, n = res['p'], res['p']['n']
    t = {'dinamico': res['tabla_din'], 'estatico': res.get('tabla_est'), 'elastico': res['elastico']['tabla'],
         'escalado': res['tabla_din_esc'] if 'tabla_din_esc' in res else res['tabla_din']}[metodo]
    sufijo = {'dinamico': '', 'estatico': ' · estático', 'elastico': ' · elástico', 'escalado': ' · dinámico escalado'}[metodo]
    zs = np.concatenate([[0.0], np.cumsum(p['alturas'])])
    H = zs[-1]
    Fu, Vu = t['Fu'][::-1], t['Vu'][::-1]   # de abajo hacia arriba (Vu: cortante en cada entrepiso)
    xs = p['geom']['xs'] if 'geom' in p else np.array([0.0, max(4.0, 0.45 * H)])
    L = float(xs[-1])
    fig = Figure(figsize=(10, min(8.5, max(4.4, 0.34 * H + 2.2))), dpi=100)
    gs = fig.add_gridspec(1, 3, width_ratios=[1.3, 1, 0.8], wspace=0.08)
    ylim = (-0.06 * H, 1.08 * H)

    # --- marco con las fuerzas sísmicas
    ax = fig.add_subplot(gs[0])
    ax.axis('off')
    ax.set_title('Fu (t)' + sufijo, color=AZULES[2], fontsize=11, fontweight='bold')
    for x in xs:
        ax.plot([x, x], [0, H], color=AZULES[0], lw=1.8, zorder=2)
    for z in zs[1:]:
        ax.plot([0, L], [z, z], color=AZULES[0], lw=2.4, zorder=2)
    ax.plot([-0.1 * L, 1.1 * L], [0, 0], color=AZULES[0], lw=1.6)
    for xh in np.arange(-0.1 * L, 1.1 * L, 0.04 * L):
        ax.plot([xh, xh - 0.03 * L], [0, -0.025 * H], color=AZULES[3], lw=0.8)
    fmax = max(Fu)
    for z, f in zip(zs[1:], Fu):
        largo = L * (0.22 + 0.5 * f / fmax)
        ax.add_patch(FancyArrowPatch((-largo, z), (0, z), arrowstyle='-|>', mutation_scale=14, color=AZULES[2], lw=2, zorder=4))
        ax.text(-largo, z + 0.012 * H, f'{f:.2f}', ha='left', va='bottom', fontsize=11, fontweight='bold', color=AZULES[0])
    ax.set_xlim(-0.95 * L, 1.15 * L)
    ax.set_ylim(*ylim)

    # --- escalera de cortantes
    ax2 = fig.add_subplot(gs[1])
    ax2.axis('off')
    ax2.set_title('Vu (t)', color=AZULES[2], fontsize=11, fontweight='bold')
    plt_hatch = {'hatch': '////'}
    for i in range(n):
        ax2.add_patch(Rectangle((0, zs[i]), Vu[i], zs[i + 1] - zs[i], fc='#DCE9F7', ec=AZULES[2], lw=1.4, **plt_hatch))
        ax2.text(Vu[i] * 1.02, zs[i + 1], f'{Vu[i]:.2f}', ha='left', va='bottom', fontsize=11, fontweight='bold',
                 color=AZULES[0])
    ax2.plot([0, 0], [0, H], color=AZULES[0], lw=1.2)
    ax2.plot([-0.05 * max(Vu), 1.05 * max(Vu)], [0, 0], color=AZULES[0], lw=1.6)
    ax2.set_xlim(-0.05 * max(Vu), 1.4 * max(Vu))
    ax2.set_ylim(*ylim)

    # --- momento de volteo
    ax3 = fig.add_subplot(gs[2])
    ax3.axis('off')
    ax3.set_title('Mvu (t·m)', color=AZULES[2], fontsize=11, fontweight='bold')
    ax3.set_xlim(0, 1)
    ax3.set_ylim(0, 1)
    ax3.text(0.5, 0.62, f'Mvu = {t["Mvu"].sum():.2f} t·m', ha='center', va='center', fontsize=12, fontweight='bold',
             color=AZULES[0], bbox=dict(boxstyle='round,pad=0.45', fc='white', ec=AZULES[2], lw=1.6))
    ax3.add_patch(FancyArrowPatch((0.2, 0.38), (0.8, 0.38), connectionstyle='arc3,rad=-0.9', arrowstyle='-|>',
                                  mutation_scale=18, color=AZULES[2], lw=2.4))
    fig.subplots_adjust(left=0.01, right=0.99, top=0.9, bottom=0.03)
    return fig


# ----------------------------------------------------------------------------
# Exportar a Excel
# ----------------------------------------------------------------------------
def exportar_excel(res, ruta):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill
    import io

    p, n = res['p'], res['p']['n']
    wb = Workbook()
    negrita = Font(bold=True)
    relleno = PatternFill('solid', fgColor='DDEBF7')

    def hoja(nombre):
        ws = wb.create_sheet(nombre)
        ws._fila = 1
        return ws

    def titulo(ws, txt):
        ws.cell(ws._fila, 1, txt).font = Font(bold=True, size=12)
        ws._fila += 1

    def tabla(ws, headers, filas, fmt='0.0000'):
        for j, h in enumerate(headers):
            c = ws.cell(ws._fila, 1 + j, h)
            c.font, c.fill = negrita, relleno
            c.alignment = Alignment(horizontal='center')
        ws._fila += 1
        for fila in filas:
            for j, v in enumerate(fila):
                v = v.item() if isinstance(v, np.generic) else v
                c = ws.cell(ws._fila, 1 + j, v)
                if isinstance(v, float):
                    c.number_format = fmt
            ws._fila += 1
        ws._fila += 1

    def ancho(ws, w=16):
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = w

    wb.remove(wb.active)

    # Datos
    ws = hoja('Datos')
    titulo(ws, 'Parámetros')
    tabla(ws, ['Parámetro', 'Valor'], [
        ['Grupo', p['grupo']], ['Sismo', p['sismo'] if p['grupo'] == 'A' else '-'],
        ['Estado límite', p['estado']], ['Combinación modal', p['combinacion']], ['ζ', p['zeta']],
        ['Dirección del sismo', p.get('direccion', 'X')], ['Vigas', p.get('vigas', 'rigidas')],
        ['Material', p.get('material', 'Concreto')], ['E (kg/cm²)', p.get('E', '')],
        ['k1', p['k1']], ['Q', ', '.join(f'{q:g}' for q in p['Q'])],
        ['a0 [cm/s²]', p['a0']], ['c [cm/s²]', p['c']], ['Ta [s]', p['Ta']], ['Tb [s]', p['Tb']],
        ['k', p['k']], ['Ts [s]', p['Ts']], ['Factor Fu', p['factor_Fu']],
        ['Multiplicador V estático', p['mult_V']]], '0.######')
    titulo(ws, 'Matriz K [kg/cm]')
    tabla(ws, [f'{j + 1}' for j in range(n)], p['K'].tolist(), '#,##0.0000')
    titulo(ws, 'Matriz M [kg s²/cm]')
    tabla(ws, [f'{j + 1}' for j in range(n)], p['M'].tolist(), '0.0000')
    if 'info' in p:
        titulo(ws, f'Modelo estructural (E = {p["E"]:,.0f} kg/cm²)')
        tabla(ws, ['Nivel', 'h (m)', 'Nº col.', 'Nº art.', 'Sección (cm)', 'ΣI (cm4)', 'K nivel (kg/cm)',
                   'W (t)', 'm (kg s²/cm)', 'Viga'],
              [[r['nivel'], r['h'], r['ncol'], r['nart'], r['seccion'], r['I'], r['k'], r['W'], r['m'], r.get('viga', '')]
               for r in p['info']], '#,##0.0000')
        if 'geom' in p:
            titulo(ws, 'Columnas (eje, posición y rigidez)')
            tabla(ws, ['Nivel', 'Eje', 'x (m)', 'Sección (cm)', 'I (cm4)', 'Apoyo', 'k (kg/cm)'],
                  [[c['nivel'], c['eje'], c['x'], c['texto'], c['I'], apoyo_txt(c), c['k']]
                   for l in p['geom']['niveles'] for c in l['cols']], '#,##0.0000')
    else:
        titulo(ws, 'Niveles (1 = planta baja)')
        tabla(ws, ['Nivel', 'h entrepiso (m)'] + (['Peso Wi (t)'] if p['cargas'] is not None else []),
              [[i + 1, p['alturas'][i]] + ([p['cargas'][i]] if p['cargas'] is not None else []) for i in range(n)])
    ancho(ws, 20)

    # Modal
    ws = hoja('Modal')
    titulo(ws, 'Propiedades modales')
    tabla(ws, ['Modo', 'λ (rad/s)²', 'ω (rad/s)', 'f (Hz)', 'T (s)', 'γ'],
          [[j + 1, res['lam'][j], res['w'][j], res['f'][j], res['T'][j], res['gam'][j]] for j in range(n)],
          '0.00000')
    titulo(ws, 'Eigenvectores normalizados (columnas = modos)')
    tabla(ws, [f'Modo {j + 1}' for j in range(n)], res['fi'].tolist(), '0.00000')
    ancho(ws)

    # Espectro
    ws = hoja('Espectro')
    titulo(ws, "Sa/(Q'R') por modo [cm/s²]")
    tabla(ws, ['Modo', 'T (s)'] + [f'Q = {q:g}' for q in res['Qs']],
          [[j + 1, res['T'][j]] + list(res['Sa_modal'][j]) for j in range(n)], '0.0000')
    tabla(ws, ['a_mínima [cm/s²]', 'a_mínima [g]'], [[res['a_min'], res['a_min_g']]], '0.000000')
    titulo(ws, "R' y Q·R' por modo")
    tabla(ws, ['Modo', 'T (s)', 'Q', 'k2', "R'", "Q·R'"], res['QR'], '0.00000')
    ancho(ws)

    # Fuerzas dinámicas
    ws = hoja('Dinámico')
    titulo(ws, 'Desplazamientos modales δ [cm] (columnas = modos)')
    tabla(ws, [f'Modo {j + 1}' for j in range(n)], res['d'].tolist(), '0.000000')
    titulo(ws, 'Fuerzas modales F [kg] (columnas = modos)')
    tabla(ws, [f'Modo {j + 1}' for j in range(n)], res['F'].tolist(), '#,##0.00')
    titulo(ws, 'Fuerzas por nivel')
    t = res['tabla_din']
    tabla(ws, ['Nivel', 'h (m)', 'F (t)', 'Fu (t)', 'Vu (t)', 'Mvu (t-m)'],
          zip(t['nivel'], t['h'], t['F'], t['Fu'], t['Vu'], t['Mvu']), '0.0000')
    tabla(ws, ['Σ Mvu (t-m)'], [[float(t['Mvu'].sum())]], '0.0000')
    ancho(ws)

    # Estático y factor de escala
    if 'tabla_est' in res:
        ws = hoja('Estático')
        e, se = res['tabla_est'], res['estatico']
        titulo(ws, 'Método estático: derivación del coeficiente sísmico (apuntes p. 20)')
        enc, filas = tabla_derivacion_estatica(res)
        tabla(ws, enc, filas, '0.0000')
        tabla(ws, ['Coeficiente C (g)', 'Gobierna', 'Criterio'],
              [[se['C'], se['gobierna'], 'Apuntes' if se['criterio'] == 'apuntes' else 'Máximo del espectro']], '0.000000')
        titulo(ws, 'Fuerzas estáticas por nivel')
        tabla(ws, ['Nivel', 'h (m)', 'Wi (t)', 'Wihi (t-m)', 'Fi (t)', 'Fu (t)', 'Vu (t)', 'Mvu (t-m)'],
              zip(e['nivel'], e['h'], e['Wi'], e['hWi'], e['Fi'], e['Fu'], e['Vu'], e['Mvu']), '0.0000')
        titulo(ws, 'Distorsiones: estático contra dinámico')
        enc, filas = tabla_distorsiones_estatico(res)
        tabla(ws, enc, filas, '0.00000')
        titulo(ws, 'Factor de escala')
        tabla(ws, ['Concepto', 'Valor'], [
            ['Coeficiente sísmico estático C [g]', res['Fmax_est']],
            ['ΣWi (t)', float(e['Wi'].sum())], ['ΣWihi (t-m)', float(e['hWi'].sum())],
            ['ΣMvu (t-m)', float(e['Mvu'].sum())],
            ['V estático (t)', res['V_est']], ['Multiplicador', p['mult_V']],
            ['V dinámico (t)', res['V_din']], ['FACTOR DE ESCALA', res['factor_escala']],
            ['Factor aplicado a las fuerzas dinámicas', res['factor_aplicado']]], '0.000000')
        titulo(ws, f'Fuerzas dinámicas escaladas (× {res["factor_aplicado"]:.4f})')
        te = res['tabla_din_esc']
        tabla(ws, ['Nivel', 'h (m)', 'F (t)', 'Fu (t)', 'Vu (t)', 'Mvu (t-m)'],
              zip(te['nivel'], te['h'], te['F'], te['Fu'], te['Vu'], te['Mvu']), '0.0000')
        ancho(ws, 30)

    # Elástico y comparación
    ws = hoja('Elástico')
    el = res['elastico']
    titulo(ws, 'Análisis elástico: espectro sin reducir (Q′ = 1, R′ = 1)')
    tabla(ws, ['Modo', 'T (s)', 'Sa elástico (cm/s²)', 'Sa de diseño (cm/s²)', 'Relación'],
          [[j + 1, res['T'][j], el['Sa'][j], res['Sa_modal'][j, 0], el['Sa'][j] / res['Sa_modal'][j, 0]] for j in range(n)],
          '0.0000')
    te = el['tabla']
    tabla(ws, ['Nivel', 'h (m)', 'F (t)', 'Fu (t)', 'Vu (t)', 'Mvu (t-m)'],
          zip(te['nivel'], te['h'], te['F'], te['Fu'], te['Vu'], te['Mvu']), '0.0000')
    titulo(ws, 'Fuerza y distorsión elásticas contra las de diseño')
    enc, filas = tabla_elastico_comparacion(res)
    tabla(ws, enc, filas, '0.00000')
    ancho(ws, 22)
    ws = hoja('Comparación')
    titulo(ws, 'Comparación de métodos')
    enc, filas = tabla_comparacion_metodos(res)
    tabla(ws, enc, filas, '0.0000')
    titulo(ws, 'Fu y Vu por nivel')
    enc, filas = tabla_niveles_metodos(res)
    tabla(ws, enc, filas, '0.0000')
    ancho(ws, 26)

    # Normativa
    ws = hoja('Normativa')
    nm = res['normativa']
    titulo(ws, 'Factores y reglas aplicadas (NTC-S 2023, apuntes de clase)')
    tabla(ws, ['Concepto', 'Valor', 'Fuente / regla'], [list(f) for f in nm['factores']])
    v = nm.get('vmin')
    if v:
        titulo(ws, 'Cortante basal mínimo')
        tabla(ws, ['a_min (g)', 'W total (t)', 'V mínimo (t)', 'V dinámico (t)', 'Cumple'],
              [[v['a_min_g'], v['W_T'], v['V_min'], v['V_din'], 'Sí' if v['cumple'] else f'No (× {v["factor"]:.3f})']], '0.0000')
    c = nm.get('cimentacion')
    if c:
        titulo(ws, "Cimentación: elementos mecánicos × 0.65·R′")
        tabla(ws, ["R′ (T1)", '0.65·R′', 'V (t)', 'Mvolteo (t·m)'], [[c['R_prima'], c['factor'], c['V'], c['M']]], '0.0000')
    t = nm.get('torsion')
    if t:
        titulo(ws, f'Torsión accidental (b = {t["b"]:g} m)')
        tabla(ws, ['Nivel', 'e_a (m)', 'Fu (t)', 'Mt (t·m)'],
              [[i + 1, float(t['e_a'][i]), float(t['Fu'][i]), float(t['Mt'][i])] for i in range(n - 1, -1, -1)], '0.0000')
    e = nm['estatico']
    titulo(ws, 'Método estático')
    tabla(ws, ['H (m)', 'Límite (m)', 'Aplicable'], [[e['H'], e['limite'] if e['limite'] else '-', 'Sí' if e['aplica'] else 'No']], '0.00')
    titulo(ws, 'Participación modal')
    tabla(ws, ['Modo', 'Masa efectiva', 'Acumulada'], [[j + 1, float(nm['modos']['m_ef'][j]), float(nm['modos']['acum'][j])] for j in range(n)], '0.0000')
    ancho(ws, 24)
    ws.column_dimensions['B'].width = 46
    ws.column_dimensions['C'].width = 60

    # Revisiones y avisos
    ws = hoja('Revisiones')
    rev = res['revision']
    titulo(ws, 'Distorsiones de entrepiso por estado límite')
    tabla(ws, ['Estado límite', 'Límite γ'],
          [[e['nombre'], e['limite'] if e['limite'] is not None else 'sin definir'] for e in rev['estados'].values()],
          '0.0000')
    enc, filas = tabla_distorsiones(res)
    tabla(ws, enc, filas, '0.00000')
    pd_ = tabla_pdelta(res)
    if pd_:
        titulo(ws, 'Efectos P-Δ (θ = P·δ / (V·h))')
        tabla(ws, pd_[0], pd_[1], '0.0000')
        tabla(ws, ['Irregularidad de masa (niveles)', 'Piso blando (niveles)'],
              [[', '.join(map(str, rev['irreg_masa'])) or 'ninguna', ', '.join(map(str, rev['irreg_rigidez'])) or 'ninguno']])
    titulo(ws, 'Avisos de coherencia')
    tabla(ws, ['Tipo', 'Aviso'], [[sev.upper(), txt] for sev, txt in res['avisos']] or [['-', 'Sin avisos']])
    ancho(ws, 22)
    ws.column_dimensions['B'].width = 40

    # Gráficas
    ws = hoja('Gráficas')
    fila = 1
    figs = (([fig_estructura(p)] if 'geom' in p else []) + [fig_modal(res), fig_espectro(res), fig_diagramas(res)]
            + ([fig_diagramas(res, 'estatico')] if 'tabla_est' in res else [])
            + [fig_elastico(res), fig_comparacion(res), fig_fuerzas(res), fig_distorsiones(res)])
    for fig in figs:
        buf = io.BytesIO()
        fig.savefig(buf, format='png')
        buf.seek(0)
        img = XLImage(buf)
        img.width, img.height = img.width * 0.7, img.height * 0.7
        ws.add_image(img, f'A{fila}')
        fila += int(img.height / 20) + 3

    wb.save(ruta)


# ----------------------------------------------------------------------------
# Ejemplo (el de las fotos)
# ----------------------------------------------------------------------------
# Ejemplo (el de las fotos)
# ----------------------------------------------------------------------------
EJEMPLO = dict(
    niveles=[dict(h='4.5', ncol='2', b='25', hs='50', nart='1', w='2.0', L='6', W=''),
             dict(h='3.7', ncol='2', b='25', hs='45', nart='0', w='2.0', L='6', W=''),
             dict(h='5.2', ncol='2', b='20', hs='40', nart='0', w='1.5', L='6', W='')],
    E='158000', fc='', tipo='Grupo B', Q='4', k1='1.0',
    a0='153', c='636', Ta='0.6', Tb='1.8', k='0.505', Ts='0.9', factor_Fu='1.1', mult_V='1')


EJEMPLOS_MODELO = {
    'Ejemplo 1 (hoja 14: 1 crujía)': dict(
        crujias='6',
        niveles=[dict(h='4.5', ejes='todos', seccion='25x50', artic='2', cargas='2.0', W=''),
                 dict(h='3.7', ejes='todos', seccion='25x45', artic='', cargas='2.0', W=''),
                 dict(h='5.2', ejes='todos', seccion='20x40', artic='', cargas='1.5', W='')],
        E='158000', fc='', tipo='Grupo B', Q='4', k1='1.0', a0='153', c='636', Ta='0.6', Tb='1.8', k='0.505',
        Ts='0.9', factor_Fu='1.1', mult_V='1'),
    'Ejemplo 2 (hoja 18: 2 crujías)': dict(
        crujias='4 5',
        niveles=[dict(h='5.0', ejes='todos', seccion='25x50', artic='', cargas='4.6917', W=''),
                 dict(h='3.5', ejes='todos', seccion='25x50', artic='', cargas='4.6667', W=''),
                 dict(h='3.5', ejes='todos', seccion='25x50', artic='', cargas='3.475', W='')],
        E='158000', fc='', tipo='Grupo B', Q='4', k1='1.0', a0='224', c='975', Ta='0.8', Tb='1.7', k='0.445',
        Ts='1.0', factor_Fu='1.1', mult_V='1'),
    'Ejemplo 3 (columna faltante, circulares, cargas por crujía)': dict(
        crujias='4 5',
        niveles=[dict(h='5.0', ejes='todos', seccion='25x50', artic='', cargas='4.7 4.7', W=''),
                 dict(h='3.5', ejes='1 3', seccion='25x50', artic='', cargas='4.0 5.5', W=''),
                 dict(h='3.5', ejes='todos', seccion='Ø50', artic='', cargas='3.0 3.8', W='')],
        E='158000', fc='', tipo='Grupo B', Q='4', k1='1.0', a0='224', c='975', Ta='0.8', Tb='1.7', k='0.445',
        Ts='1.0', factor_Fu='1.1', mult_V='1'),
    'Ejemplo 4 (vigas flexibles, concreto)': dict(
        crujias='4 5', vigas='Flexibles (marco plano)', material='Concreto', direccion='X',
        niveles=[dict(h='5.0', ejes='todos', seccion='40x40', artic='', viga='30x60', cargas='4.7', W=''),
                 dict(h='3.5', ejes='todos', seccion='40x40', artic='', viga='30x60', cargas='4.7', W=''),
                 dict(h='3.5', ejes='todos', seccion='35x35', artic='', viga='25x50', cargas='3.5', W='')],
        E='158000', fc='', tipo='Grupo B', Q='4', k1='1.0', a0='224', c='975', Ta='0.8', Tb='1.7', k='0.445',
        Ts='1.0', factor_Fu='1.1', mult_V='1'),
    'Ejemplo 5 (acero: columnas H, vigas H)': dict(
        crujias='6 6', vigas='Flexibles (marco plano)', material='Acero', direccion='X',
        niveles=[dict(h='4.0', ejes='todos', seccion='H:30x1.6x40x1', artic='', viga='H:20x1x45x0.8', cargas='3.0', W=''),
                 dict(h='3.5', ejes='todos', seccion='H:30x1.6x40x1', artic='', viga='H:20x1x45x0.8', cargas='3.0', W=''),
                 dict(h='3.5', ejes='todos', seccion='H:25x1.2x35x0.9', artic='', viga='H:20x1x40x0.8', cargas='2.5', W='')],
        E='', fc='', tipo='Grupo B', Q='3', k1='1.0', a0='224', c='975', Ta='0.8', Tb='1.7', k='0.445',
        Ts='1.0', factor_Fu='1.1', mult_V='1', zeta='0.03'),
    'Ejemplo 6 (hoja 19: K y M dadas, ocupación inmediata)': dict(
        modo='matrices', unidades='t, cm', alturas='5 5 5 5', pesos='',
        K='300 -150 0 0\n-150 225 -75 0\n0 -75 150 -75\n0 0 -75 75', M='3 0 0 0\n0 3 0 0\n0 0 1.5 0\n0 0 0 1.5',
        subgrupo='A2', intensidad='Base de diseño', material='Concreto', Q='2', k1='0.8',
        a0='80', c='500', Ta='0.4', Tb='1.6', k='0.605', Ts='0.53', factor_Fu='1.1', mult_V='1'),
    'Ejemplo 7 (hoja 21: método estático, estructura irregular)': dict(
        modo='matrices', unidades='t, cm', alturas='3 3 3 3 3', pesos='400 400 400 400 300',
        K='3000 -1500 0 0 0\n-1500 3000 -1500 0 0\n0 -1500 3000 -1500 0\n0 0 -1500 3000 -1500\n0 0 0 -1500 1500',
        M='0.40775 0.40775 0.40775 0.40775 0.30581',
        subgrupo='B1', intensidad='Base de diseño', material='Concreto', Q='2', k1='0.8', zona='III',
        regularidad='Irregular', a0='224', c='975', Ta='0.8', Tb='1.7', k='0.445', Ts='1.05', factor_Fu='1.1',
        mult_V='1'),
    # Examen del 29-oct-2025, problema 1: marco de concreto con apoyos mixtos (empotrado, articulado, empotrado y
    # rodillos), secciones distintas por columna, E en T/cm², grupo A y sismo máximo considerado (MC) -> seguridad de vida.
    # Los parámetros del espectro SASID son ilustrativos: el examen los pide de las coordenadas.
    'Ejemplo 8 (examen: apoyos mixtos con rodillo, grupo A, MC)': dict(
        crujias='4 4 2', vigas='rigidas', material='Concreto', direccion='X', unidad_E='t/cm²',
        niveles=[dict(h='3', ejes='todos', seccion='20x40 30x30 20x40 30x60', artic='2', rod='4', cargas='2', W=''),
                 dict(h='3', ejes='todos', seccion='20x40 30x30 20x40 30x60', artic='', rod='', cargas='1', W='')],
        E='158', fc='', subgrupo='A2', intensidad='Infrecuente', Q='2', k1='0.8', zeta='0.05',
        a0='300', c='1300', Ta='0.8', Tb='1.7', k='0.445', Ts='1.0', factor_Fu='1.1', mult_V='1'),
    # Examen del 29-oct-2025, problema 2: nave industrial. Se mide T = 0.17 s y se pide el peso total (2a); luego la
    # fuerza sísmica estática para operación continua / limitación de daño (2b). Perfil IR 305x21.1 con el eje fuerte en X.
    'Ejemplo 9 (examen: nave con periodo medido, operación continua)': dict(
        crujias='5 7 5', vigas='rigidas', material='Acero', direccion='X', T_medido='0.17',
        niveles=[dict(h='5.3', ejes='todos', seccion='I:3688/A:26.8/I2:98', artic='', rod='', cargas='', W='')],
        E='2100000', fc='', subgrupo='B1', intensidad='Frecuente', Q='1', k1='1.0', zeta='0.05',
        a0='224', c='975', Ta='0.8', Tb='1.7', k='0.445', Ts='1.0', Ks='0.25', factor_Fu='1.1', mult_V='1'),
    'Ejemplo 10 (examen: marco mixto, concreto abajo y acero IR arriba)': dict(
        crujias='5 4 4.5', vigas='flexibles', material='Concreto', direccion='X', unidad_E='t/cm²', E='221.35944',
        niveles=[dict(h='4.5', ejes='todos', seccion='50x50', viga='25x50', artic='2 3', cargas='1.4', E_col=''),
                 dict(h='3.5', ejes='todos', seccion='50x50', viga='25x50', cargas='2', E_col=''),
                 dict(h='3.5', ejes='todos', seccion='50x50', viga='25x50', cargas='2', E_col=''),
                 dict(h='3.5', ejes='todos', seccion='I:3688/A:26.8/I2:98', viga='25x50', cargas='1.8', E_col='2100')],
        fc='', subgrupo='B1', intensidad='Base de diseño', Q='2', k1='1.0', zeta='0.05',
        a0='114', c='320', Ta='0.9', Tb='2.2', k='0.390', Ts='0.5', factor_Fu='1.1', mult_V='1'),
    # Pizarrón del profesor: volados de 3 m (izq.) y 2 m (der.), rodillo en el eje 1, empotrado en el 2 y articulado en el 3.
    # Pesos: W1 = 4(3)+2(6)+2(3) = 30 t, W2 = 4(6)+2(3) = 30 t, W3 = 3(9) = 27 t. Secciones y espectro: ilustrativos.
    'Ejemplo 11 (pizarrón: volados 3 y 2 m, apoyos mixtos, 3 niveles)': dict(
        crujias='6 3', volado_izq='3', volado_der='2', vigas='rigidas', material='Concreto', direccion='X',
        unidad_E='t/cm²', E='158',
        niveles=[dict(h='3', ejes='todos', seccion='40x40', rod='1', artic='3', cargas='2 2 4 0'),
                 dict(h='4', ejes='todos', seccion='40x40', cargas='0 4 2 0'),
                 dict(h='3.5', ejes='1 3', seccion='40x40', cargas='0 3 3 0')],
        fc='', subgrupo='B1', intensidad='Base de diseño', Q='2', k1='1.0', zeta='0.05',
        a0='114', c='320', Ta='0.9', Tb='2.2', k='0.390', Ts='0.5', factor_Fu='1.1', mult_V='1'),
    # Problema de tarea: 4 niveles; C1 50x50 de concreto (E = 221.35944 t/cm²) y C2 = IR 305x21.1 de acero (E = 2.1e6 kg/cm²).
    # En el piso del nivel 3 no hay losa en la crujía 3, así que la columna del eje 4 es una sola de 7 m (del piso 2 al 4).
    # Vigas: 25x50 de concreto en los pisos 1 y 2; IR 305x21.1 de acero en los pisos 3 y 4 (la ampliación).
    # Espectro: valores de la captura del problema (grupo A, ocupación inmediata, sismo base de diseño).
    'Ejemplo 12 (tarea: columnas mixtas y piso sin losa en una crujía)': dict(
        crujias='5 4 4.5', vigas='rigidas', material='Concreto', direccion='X', unidad_E='t/cm²', E='221.35944',
        niveles=[dict(h='4.5', ejes='todos', seccion='50x50', viga='25x50', artic='2 3', cargas='1.4'),
                 dict(h='3.5', ejes='todos', seccion='50x50', viga='25x50', cargas='2'),
                 dict(h='3.5', ejes='todos', seccion='50x50 50x50 50x50 I:3688/A:26.8', E_col='221.35944 221.35944 221.35944 2100',
                      viga='I:3688/A:26.8', E_viga='2100', sin_losa='3', cargas='2'),
                 dict(h='3.5', ejes='todos', seccion='I:3688/A:26.8', E_col='2100', viga='I:3688/A:26.8', E_viga='2100',
                      cargas='1.8')],
        fc='', subgrupo='A2', intensidad='Base de diseño', Q='1', k1='0.8', zeta='0.05',
        a0='242', c='1062', Ta='0.8', Tb='1.7', k='0.430', Ts='1.0', factor_Fu='1.1', mult_V='1'),
}


def main():
    args = sys.argv[1:]
    if not args:
        from ame_app import gui
        gui()
        return
    entrada = Path(args[0])
    salida = Path(args[args.index('-o') + 1]) if '-o' in args else entrada.with_suffix('.xlsx')
    datos = json.loads(entrada.read_text(encoding='utf-8'))
    datos = convertir_legacy(datos)
    if datos.get('modo') == 'matrices':
        p = preparar_datos(datos)
    elif 'crujias' in datos:
        p = preparar_modelo(datos)
    elif isinstance(datos.get('niveles'), list):
        p = preparar_estructura(datos)
    else:
        p = preparar_datos(datos)
    res = calcular(p)
    if hasattr(sys.stdout, 'reconfigure'):  # la consola de Windows no imprime λ, γ, Σ con su codificación por defecto
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    print(reporte_texto(res))
    exportar_excel(res, salida)
    print(f'\nExportado: {salida}')
    if '--pdf' in args:
        from ame_pdf import memoria_pdf
        pdf = salida.with_suffix('.pdf')
        memoria_pdf(res, str(pdf), proyecto=datos.get('proyecto', ''), autor=datos.get('autor', ''))
        print(f'Memoria de cálculo: {pdf}')


if __name__ == '__main__':
    main()
